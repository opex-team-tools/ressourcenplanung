"""Excel-Export der Planung im gewohnten Format (Rolling Forecast).
Blatt 'Tage': Mitarbeiter x Projekt x Arbeitstag, Kopf mit Monat / KW / Datum.
Blatt 'Wochen': gleiche Zeilen, Summe je KW plus Auslastung.
Blatt 'Projekte': Budget, gebucht, geplant, offen je Projekt und Berater.
Blatt 'Übersicht' (vorne): Gesamtstand, wer hat welche Projekte, welche Projekte sind wie verplant."""
from __future__ import annotations

import datetime as dt
import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

import logic

NAVY, MID, LIGHT, LILAC, GREY, XL = "000055", "6688BB", "BBCCEE", "E4E4F6", "999999", "EBEBEB"
ALERT = "B3261E"
F_HEAD = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
F_BOLD = Font(name="Calibri", bold=True, color=NAVY, size=10)
F_TXT = Font(name="Calibri", color=NAVY, size=10)
F_RED = Font(name="Calibri", bold=True, color=ALERT, size=10)
FILL = {k: PatternFill("solid", fgColor=v) for k, v in
        {"navy": NAVY, "mid": MID, "light": LIGHT, "lilac": LILAC, "xl": XL, "hol": "D9D9D9", "white": "FFFFFF"}.items()}
THIN = Side(style="thin", color="D2D2D2")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober",
          "November", "Dezember"]
H = logic.TAG_H


def _rows(cons: pd.DataFrame, projn: pd.DataFrame, ent: pd.DataFrame, mem: pd.DataFrame) -> list[dict]:
    """Zeilenstruktur: je Mitarbeiter seine Projekte (Budget oder Einträge), dann Intern / Abwesenheiten."""
    out = []
    for c in cons.itertuples():
        e = ent[ent["consultant_id"] == c.id] if not ent.empty else ent
        pids = set(mem.loc[mem["consultant_id"] == c.id, "project_id"]) if not mem.empty else set()
        if not e.empty:
            pids |= set(e.loc[e["project_id"].notna(), "project_id"].astype(int))
        pids = [p for p in pids if p in projn.index]
        pids.sort(key=lambda p: projn.at[p, "label"].lower())
        rows = [{"cid": c.id, "pid": p, "art": None, "label": projn.at[p, "label"]} for p in pids]
        arts = [a for a in list(logic.INTERN_ARTEN) + list(logic.ABWESEND_ARTEN)
                if not e.empty and (e["art"] == a).any()]
        for a in ["Intern", "Urlaub"] + [x for x in arts if x not in ("Intern", "Urlaub")]:
            rows.append({"cid": c.id, "pid": None, "art": a, "label": a})
        out.append({"cid": c.id, "kz": c.name, "name": (c.vollname if isinstance(c.vollname, str) and c.vollname
                                                         else c.name), "rows": rows})
    return out


def _sheet_uebersicht(wb, cons: pd.DataFrame, projn: pd.DataFrame, ent: pd.DataFrame, mem: pd.DataFrame,
                      today: dt.date, titel: str) -> None:
    """Gesamtstand auf einen Blick: wer hat welche Projekte, welche Projekte sind wie verplant."""
    ws = wb.create_sheet("Übersicht")
    wb.move_sheet(ws, offset=-(len(wb.sheetnames) - 1))
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{titel} | Gesamtstand"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    ws["A2"] = (f"Stand {today:%d.%m.%Y} · KW {today.isocalendar().week} · Werte in Beratertagen (1 = 8 h) · "
                "gebucht = bis heute, geplant = ab morgen")
    ws["A2"].font = Font(name="Calibri", italic=True, size=10, color=MID)
    m0 = logic.monday(today)
    h4, h12 = m0 + dt.timedelta(days=27), m0 + dt.timedelta(weeks=12, days=-3)
    f12 = logic.build_fact(m0, h12, cons, ent)
    ms = logic.member_stats(ent, mem, today)
    stats = logic.project_stats(projn.reset_index(), ent, mem, today)
    fut = ent[(ent["tag"] > today) & ent["project_id"].notna()] if not ent.empty else ent
    fut_pids = set(fut["project_id"].astype(int)) if not fut.empty else set()

    def laufend(s) -> bool:
        return bool(s.aktiv) and ((pd.notna(s.ende) and s.ende >= today) or s.id in fut_pids or s.offen > 0.01)

    lauf = [s for s in stats.sort_values(["kunde", "name"]).itertuples() if laufend(s)]
    lauf_ids = {s.id for s in lauf}

    def head(row: int, cols: list[str], start_col: int = 1) -> None:
        for i, h in enumerate(cols, start=start_col):
            c = ws.cell(row, i, h)
            c.font, c.fill, c.border = F_HEAD, FILL["navy"], BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 30

    # ---------------- Teil 1: Mitarbeiter
    r = 4
    ws.cell(r, 1, "1 | Mitarbeiter: Auslastung und Projekte").font = Font(name="Calibri", bold=True, size=12,
                                                                          color=MID, italic=True)
    r += 1
    head(r, ["Mitarbeiter", "Laufende Projekte (geplante Tage ab morgen / offene BT, Ort)", "", "",
             "Auslastung 4 Wochen", "Auslastung 12 Wochen", "Freie Tage 12 Wochen"])
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    r += 1
    for c in cons.itertuples():
        f = f12[f12["consultant_id"] == c.id] if not f12.empty else f12
        u4 = logic.kpi(f[f["tag"] <= h4]) if not f.empty else None
        u12 = logic.kpi(f) if not f.empty else None
        frei = float(f["frei"].sum()) / H if not f.empty else 0.0
        m = ms[(ms["consultant_id"] == c.id) & ms["project_id"].isin(lauf_ids)] if not ms.empty else ms
        teile = []
        for x in m.sort_values("geplant", ascending=False).itertuples():
            fe = fut[(fut["consultant_id"] == c.id) & (fut["project_id"] == x.project_id)] if not fut.empty else fut
            n_vo = int((fe["art"] == "Vor Ort").sum()) if not fe.empty else 0
            n_re = int((fe["art"] == "Remote").sum()) if not fe.empty else 0
            ort = "Vor Ort" if n_vo and not n_re else "Remote" if n_re and not n_vo else \
                f"{n_vo} Vor Ort / {n_re} Remote" if n_vo else "noch ohne Tage"
            z = f"{round(x.geplant, 1):g} / {round(x.offen, 1):g}".replace(".", ",")
            teile.append(f"{projn.at[x.project_id, 'label']}: {z} BT, {ort}")
        name = c.vollname if isinstance(c.vollname, str) and c.vollname else c.name
        vals = [f"{name} ({c.name})", "\n".join(teile) or "keine laufenden Projekte", None, None, u4, u12,
                round(frei, 1)]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(r, i, v)
            cell.font, cell.border = F_TXT, BOX
            cell.alignment = Alignment(vertical="top", wrap_text=(i == 2))
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        ws.cell(r, 1).font = F_BOLD
        for col in (5, 6):
            ws.cell(r, col).number_format = "0%"
            u = ws.cell(r, col).value
            if u is not None and u > 1.0001:
                ws.cell(r, col).font = F_RED
            elif u is not None and u < 0.6:
                ws.cell(r, col).fill = FILL["lilac"]
        ws.cell(r, 7).number_format = "0.0"
        ws.row_dimensions[r].height = max(16, 14 * max(len(teile), 1))
        r += 1

    # ---------------- Teil 2: Projekte x Mitarbeiter
    r += 2
    ws.cell(r, 1, "2 | Laufende Projekte: Budget, Verplanung und Team (geplante Tage ab morgen je Mitarbeiter)").font = \
        Font(name="Calibri", bold=True, size=12, color=MID, italic=True)
    r += 1
    kz = cons["name"].tolist()
    head(r, ["Kunde", "Projekt", "Budget BT", "Gebucht", "Geplant", "Offen BT", "Status"] + kz)
    r0 = r + 1
    r += 1
    ms_i = ms.set_index(["project_id", "consultant_id"]) if not ms.empty else ms
    for s in lauf:
        status = ("über Budget" if s.offen < -0.01 else "voll verplant" if s.offen <= 0.01 else
                  f"{s.offen:g} BT offen".replace(".", ","))
        if s.budget <= 0.01:
            status = "kein Budget"
        vals = [s.kunde, s.name if s.name != s.kunde else "", s.budget, s.gebucht, s.geplant, s.offen, status]
        for i, v in enumerate(vals, start=1):
            cell = ws.cell(r, i, v)
            cell.font, cell.border = F_TXT, BOX
        ws.cell(r, 1).font = F_BOLD
        if s.offen < -0.01:
            ws.cell(r, 6).font = ws.cell(r, 7).font = F_RED
        for j, c in enumerate(cons.itertuples()):
            cell = ws.cell(r, 8 + j)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center")
            if not ms.empty and (s.id, c.id) in ms_i.index:
                g = float(ms_i.loc[(s.id, c.id), "geplant"])
                cell.value = round(g, 1)
                cell.font = F_BOLD if g > 0 else F_TXT
                cell.fill = FILL["lilac"] if g > 0 else FILL["white"]
        r += 1
    for rr in range(r0, r):
        for col in range(3, 7):
            ws.cell(rr, col).number_format = "0.0"
        for col in range(8, 8 + len(kz)):
            ws.cell(rr, col).number_format = "0.0"
    ws.cell(r + 1, 1, "Leere Zelle = nicht im Projekt · 0 = im Projekt, aber keine Tage mehr geplant · "
                      "Auslastung lila = unter 60 %, rot = über 100 %").font = \
        Font(name="Calibri", italic=True, size=9, color=MID)
    for col, w in zip("ABCDEFG", (30, 30, 12, 12, 12, 12, 16)):
        ws.column_dimensions[col].width = w
    for j in range(len(kz)):
        ws.column_dimensions[L(8 + j)].width = 7
    ws.freeze_panes = "B4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _sheet_guete(wb, cons: pd.DataFrame, projn: pd.DataFrame, ent: pd.DataFrame, val: pd.DataFrame,
                 today: dt.date, titel: str) -> None:
    """Planungsgüte: geplante Projekttage vs. gebuchte VAL-Tage je Monat (VAL ÷ Plan, 100 % = exakt)."""
    ws = wb.create_sheet("Planungsgüte")
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{titel} | Planungsgüte (VAL ÷ Plan)"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    ws["A2"] = ("100 % = exakt geplant · unter 100 % = mehr geplant als gebucht · über 100 % = mehr gebucht als "
                "geplant · offen = VAL noch nicht gemeldet")
    ws["A2"].font = Font(name="Calibri", italic=True, size=10, color=MID)
    # ab dem ersten gemeldeten Monat (davor gab es keine VAL-Erfassung), mindestens der Vormonat
    von = min(val["monat"]) if not val.empty else logic.month_start(today, -1)
    monate, m = [], von.replace(day=1)
    while m <= today:
        monate.append(m)
        m = logic.month_start(m, 1)
    e = ent[(ent["tag"] >= monate[0]) & (ent["tag"] <= logic.month_end(today))] if not ent.empty else ent
    pv = logic.plan_val(e, val)
    gm = logic.guete_matrix(pv)
    lab = [f"{MONATE[x.month - 1][:3]} {x:%y}" for x in monate]
    LILA_OK = PatternFill("solid", fgColor=LILAC)

    def head(row, cols):
        for i, h in enumerate(cols, start=1):
            c = ws.cell(row, i, h)
            c.font, c.fill, c.border = F_HEAD, FILL["navy"], BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def guete_cell(c, g):
        c.value = g
        c.number_format = "0%"
        c.alignment = Alignment(horizontal="center")
        if g is None:
            return
        d = abs(g - 1)
        c.font = F_RED if d > 0.25 else F_BOLD
        if d <= 0.1:
            c.fill = LILA_OK

    # Teil 1: Matrix Güte
    r = 4
    ws.cell(r, 1, "1 | Planungsgüte je Mitarbeiter und Monat").font = Font(name="Calibri", bold=True, size=12,
                                                                         color=MID, italic=True)
    r += 1
    head(r, ["Mitarbeiter"] + lab + [f"Σ {today.year}"])
    r += 1
    rows_team = {m: [0.0, 0.0] for m in monate}
    for c in cons.itertuples():
        ws.cell(r, 1, c.vollname if isinstance(c.vollname, str) and c.vollname else c.name).font = F_BOLD
        ws.cell(r, 1).border = BOX
        sp, sv = 0.0, 0.0
        for j, m in enumerate(monate):
            cell = ws.cell(r, 2 + j)
            cell.border = BOX
            z = gm[(gm["consultant_id"] == c.id) & (gm["monat"] == m)] if not gm.empty else gm
            if z.empty:
                continue
            if not bool(z["gemeldet"].iat[0]):
                cell.value, cell.font = "offen", Font(name="Calibri", italic=True, size=10, color=GREY)
                cell.alignment = Alignment(horizontal="center")
                continue
            p_, v_ = float(z["plan"].iat[0]), float(z["val"].iat[0])
            guete_cell(cell, logic.guete(p_, v_))
            rows_team[m][0] += p_
            rows_team[m][1] += v_
            if m.year == today.year:
                sp, sv = sp + p_, sv + v_
        cell = ws.cell(r, 2 + len(monate))
        cell.border = BOX
        guete_cell(cell, logic.guete(sp, sv))
        r += 1
    ws.cell(r, 1, "Team").font = F_BOLD
    ws.cell(r, 1).fill = FILL["lilac"]
    tp = tv = 0.0
    for j, m in enumerate(monate):
        cell = ws.cell(r, 2 + j)
        cell.border = BOX
        guete_cell(cell, logic.guete(*rows_team[m]))
        if m.year == today.year:
            tp, tv = tp + rows_team[m][0], tv + rows_team[m][1]
    cell = ws.cell(r, 2 + len(monate))
    cell.border = BOX
    guete_cell(cell, logic.guete(tp, tv))

    ws.column_dimensions["A"].width = 24
    for j in range(2, 3 + len(monate)):
        ws.column_dimensions[L(j)].width = 9
    ws.freeze_panes = "B6"

    # Blatt 2: Details
    w2 = wb.create_sheet("VAL Details")
    w2.sheet_view.showGridLines = False
    w2["A1"] = f"{titel} | Plan vs. VAL je Monat, Mitarbeiter und Projekt (Beratertage)"
    w2["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    heads = ["Monat", "Mitarbeiter", "Projekt", "Plan BT", "VAL BT", "Abweichung BT", "Planungsgüte", "Status"]
    for i, h in enumerate(heads, start=1):
        c = w2.cell(3, i, h)
        c.font, c.fill, c.border = F_HEAD, FILL["navy"], BOX
        c.alignment = Alignment(horizontal="center", vertical="center")
    r = 4
    cmap = {c.id: (c.vollname if isinstance(c.vollname, str) and c.vollname else c.name) for c in cons.itertuples()}
    if not pv.empty:
        pv = pv[pv["consultant_id"].isin(cmap)].sort_values(["monat", "consultant_id", "project_id"])
        for x in pv.itertuples():
            vals = [f"{MONATE[x.monat.month - 1]} {x.monat.year}", cmap[x.consultant_id],
                    projn.at[x.project_id, "label"] if x.project_id in projn.index else "",
                    round(x.plan, 2), round(x.val, 2) if x.gemeldet else None,
                    round(x.val - x.plan, 2) if x.gemeldet else None]
            for i, v in enumerate(vals, start=1):
                c = w2.cell(r, i, v)
                c.font, c.border = F_TXT, BOX
            for col in (4, 5, 6):
                w2.cell(r, col).number_format = "0.00"
            g = logic.guete(x.plan, x.val) if x.gemeldet else None
            c = w2.cell(r, 7)
            c.border = BOX
            c.value, c.number_format = g, "0%"
            if g is not None:
                c.font = F_RED if abs(g - 1) > 0.25 else F_BOLD
                if abs(g - 1) <= 0.1:
                    c.fill = LILA_OK
            w2.cell(r, 8, "gemeldet" if x.gemeldet else "offen").font = F_TXT
            w2.cell(r, 8).border = BOX
            r += 1
    for col, w in zip("ABCDEFGH", (16, 22, 38, 10, 10, 14, 13, 11)):
        w2.column_dimensions[col].width = w
    w2.freeze_panes = "A4"
    w2.auto_filter.ref = f"A3:H{max(r - 1, 3)}"


def build_excel(cons: pd.DataFrame, projn: pd.DataFrame, ent: pd.DataFrame, mem: pd.DataFrame,
                start: dt.date, end: dt.date, today: dt.date, titel: str = "AUSLASTUNG",
                val: pd.DataFrame | None = None) -> bytes:
    days = [d for d in pd.date_range(start, end).date if d.weekday() < 5]
    weeks = sorted({logic.monday(d) for d in days})
    ms = logic.member_stats(ent, mem, today).set_index(["project_id", "consultant_id"]) if not mem.empty or not ent.empty \
        else pd.DataFrame()
    e = ent.copy()
    if not e.empty:
        e["key"] = [f"p{int(p)}" if pd.notna(p) else a for p, a in zip(e["project_id"], e["art"])]
    cell_h = e.groupby(["consultant_id", "key", "tag"])["stunden"].sum().to_dict() if not e.empty else {}
    cell_art = e.groupby(["consultant_id", "key", "tag"])["art"].first().to_dict() if not e.empty else {}
    people = _rows(cons, projn, ent, mem)

    wb = Workbook()
    # ------------------------------------------------------------------ Blatt Tage
    ws = wb.active
    ws.title = "Tage"
    first = 6  # erste Datenspalte (A Mitarbeiter, B Projekt, C Budget, D Offen, E Summe)
    ws["A1"] = f"{titel} | Rolling Forecast"
    ws["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    ws["A2"] = f"Stand {today:%d.%m.%Y} · Aktuelle KW {today.isocalendar().week} · Werte in Beratertagen (1 = 8 h)"
    ws["A2"].font = Font(name="Calibri", italic=True, size=10, color=MID)
    hdr = ["Mitarbeiter", "Projekt / Tätigkeit", "Budget BT", "Offen BT", "Σ Zeitraum"]
    for i, h in enumerate(hdr, start=1):
        ws.merge_cells(start_row=4, start_column=i, end_row=6, end_column=i)
        c = ws.cell(4, i, h)
        c.font, c.fill, c.alignment = F_HEAD, FILL["navy"], Alignment(horizontal="center", vertical="center", wrap_text=True)
    col_of = {}
    for j, d in enumerate(days):
        col = first + j
        col_of[d] = col
        c = ws.cell(6, col, d)
        c.number_format = "DD.MM."
        c.font, c.fill = F_HEAD, FILL["mid"]
        c.alignment = Alignment(horizontal="center", text_rotation=90)
        ws.column_dimensions[L(col)].width = 4.2
    # Monats- und KW-Köpfe (verbunden)
    for key_fn, row, fill in ((lambda d: (d.year, d.month), 4, "navy"),
                              (lambda d: logic.monday(d), 5, "mid")):
        j = 0
        while j < len(days):
            k = key_fn(days[j])
            k2 = j
            while k2 + 1 < len(days) and key_fn(days[k2 + 1]) == k:
                k2 += 1
            if k2 > j:
                ws.merge_cells(start_row=row, start_column=first + j, end_row=row, end_column=first + k2)
            txt = f"{MONATE[k[1] - 1]} {k[0]}" if row == 4 else f"KW {k.isocalendar().week}"
            c = ws.cell(row, first + j, txt)
            c.font, c.fill, c.alignment = F_HEAD, FILL[fill], Alignment(horizontal="center")
            j = k2 + 1
    r = 7
    for p in people:
        # Kopfzeile Mitarbeiter: Summe gebuchter Tage je Tag (rot bei Überbuchung)
        ws.cell(r, 1, p["name"]).font = F_BOLD
        ws.cell(r, 2, f"Summe ({p['kz']})").font = F_BOLD
        tot_period = 0.0
        for d, col in col_of.items():
            if logic.feiertag_name(d):
                c = ws.cell(r, col, "x")
                c.fill = FILL["hol"]
                continue
            h = sum(v for (cid, key, tag), v in cell_h.items() if cid == p["cid"] and tag == d)
            if h:
                c = ws.cell(r, col, round(h / H, 2))
                c.font = F_RED if h > H + 1e-9 else F_BOLD
                tot_period += h / H
        ws.cell(r, 5, round(tot_period, 2)).font = F_BOLD
        for col in range(1, first + len(days)):
            ws.cell(r, col).fill = FILL["lilac"] if ws.cell(r, col).value != "x" else FILL["hol"]
            ws.cell(r, col).border = BOX
        r += 1
        for row in p["rows"]:
            key = f"p{row['pid']}" if row["pid"] is not None else row["art"]
            ws.cell(r, 2, row["label"]).font = F_TXT
            if row["pid"] is not None and not ms.empty and (row["pid"], p["cid"]) in ms.index:
                b = ms.loc[(row["pid"], p["cid"])]
                ws.cell(r, 3, round(float(b["budget"]), 1)).font = F_TXT
                c = ws.cell(r, 4, round(float(b["offen"]), 1))
                c.font = F_RED if b["offen"] < -0.01 else F_TXT
            tot = 0.0
            for d, col in col_of.items():
                c = ws.cell(r, col)
                c.border = BOX
                if logic.feiertag_name(d):
                    c.value, c.fill = "x", FILL["hol"]
                    continue
                h = cell_h.get((p["cid"], key, d), 0.0)
                if h:
                    c.value = round(h / H, 2)
                    art = cell_art.get((p["cid"], key, d))
                    c.fill = FILL["lilac"] if art == "Vor Ort" else FILL["white"] if art == "Remote" else FILL["xl"]
                    c.font = F_TXT
                    tot += h / H
            ws.cell(r, 5, round(tot, 2)).font = F_TXT
            for col in range(1, 6):
                ws.cell(r, col).border = BOX
            r += 1
        r += 1
    for col, w in zip("ABCDE", (22, 42, 10, 10, 10)):
        ws.column_dimensions[col].width = w
    for col in (3, 4, 5):
        for rr in range(7, r):
            ws.cell(rr, col).number_format = "0.0"
    ws.freeze_panes = ws.cell(7, first)
    ws.cell(r + 1, 1, "Legende:").font = F_BOLD
    for i, (t, f) in enumerate((("Vor Ort", "lilac"), ("Remote", "white"), ("Intern / Abwesend", "xl"),
                                ("Feiertag (x)", "hol"))):
        c = ws.cell(r + 1, 2 + i, t)
        c.fill, c.border, c.font = FILL[f], BOX, F_TXT
    ws.cell(r + 2, 1, "Rote Werte = überbucht (mehr als 1 Tag bzw. 8 h) · Offen negativ = mehr geplant als budgetiert").font = \
        Font(name="Calibri", italic=True, size=9, color=MID)

    # ------------------------------------------------------------------ Blatt Wochen
    w2 = wb.create_sheet("Wochen")
    w2["A1"] = f"{titel} | Wochenübersicht (Beratertage)"
    w2["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    heads = ["Mitarbeiter", "Projekt / Tätigkeit"] + [f"KW {w.isocalendar().week}\n{w:%d.%m.}" for w in weeks]
    for i, h in enumerate(heads, start=1):
        c = w2.cell(3, i, h)
        c.font, c.fill = F_HEAD, FILL["navy"]
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if i > 2:
            w2.column_dimensions[L(i)].width = 8
    w2.row_dimensions[3].height = 30
    fact = logic.build_fact(start, end, cons, ent)
    r = 4
    for p in people:
        w2.cell(r, 1, p["name"]).font = F_BOLD
        w2.cell(r, 2, "Auslastung").font = F_BOLD
        f = fact[fact["consultant_id"] == p["cid"]]
        for i, w in enumerate(weeks):
            fw = f[f["tag"].map(logic.monday) == w]
            u = logic.kpi(fw)
            c = w2.cell(r, 3 + i, u if u is not None else "")
            c.number_format = "0%"
            c.font = F_RED if (u or 0) > 1.0001 else F_BOLD
            c.fill, c.border = FILL["lilac"], BOX
        w2.cell(r, 1).fill = w2.cell(r, 2).fill = FILL["lilac"]
        r += 1
        for row in p["rows"]:
            key = f"p{row['pid']}" if row["pid"] is not None else row["art"]
            w2.cell(r, 2, row["label"]).font = F_TXT
            for i, w in enumerate(weeks):
                h = sum(cell_h.get((p["cid"], key, w + dt.timedelta(days=k)), 0.0) for k in range(5))
                c = w2.cell(r, 3 + i, round(h / H, 2) if h else None)
                c.number_format, c.border, c.font = "0.0", BOX, F_TXT
            r += 1
        r += 1
    w2.column_dimensions["A"].width, w2.column_dimensions["B"].width = 22, 42
    w2.freeze_panes = "C4"

    # ------------------------------------------------------------------ Blatt Projekte
    w3 = wb.create_sheet("Projekte")
    w3["A1"] = f"{titel} | Projekte und Budgets (Beratertage)"
    w3["A1"].font = Font(name="Calibri", bold=True, size=14, color=NAVY)
    heads = ["Kunde", "Projekt", "Berater", "Budget BT", "Gebucht bis heute", "Geplant ab morgen", "Offen BT",
             "Start", "Ende"]
    for i, h in enumerate(heads, start=1):
        c = w3.cell(3, i, h)
        c.font, c.fill = F_HEAD, FILL["navy"]
    stats = logic.project_stats(projn.reset_index(), ent, mem, today)
    cmap = dict(zip(cons["id"], cons["name"]))
    r = 4
    for s in stats.sort_values(["kunde", "name"]).itertuples():
        vals = [s.kunde, s.name if s.name != s.kunde else "", "Gesamt", s.budget, s.gebucht, s.geplant, s.offen,
                s.start if pd.notna(s.start) else None, s.ende if pd.notna(s.ende) else None]
        for i, v in enumerate(vals, start=1):
            c = w3.cell(r, i, v)
            c.font, c.fill, c.border = F_BOLD, FILL["lilac"], BOX
        w3.cell(r, 7).font = F_RED if s.offen < -0.01 else F_BOLD
        r += 1
        if not ms.empty:
            sub = ms.reset_index()
            for m in sub[sub["project_id"] == s.id].itertuples():
                vals = ["", "", cmap.get(m.consultant_id, ""), m.budget, m.gebucht, m.geplant, m.offen]
                for i, v in enumerate(vals, start=1):
                    c = w3.cell(r, i, v)
                    c.font, c.border = F_TXT, BOX
                w3.cell(r, 7).font = F_RED if m.offen < -0.01 else F_TXT
                r += 1
    for col, w in zip("ABCDEFGHI", (26, 34, 10, 11, 16, 16, 10, 12, 12)):
        w3.column_dimensions[col].width = w
    for rr in range(4, r):
        for col in (4, 5, 6, 7):
            w3.cell(rr, col).number_format = "0.0"
        for col in (8, 9):
            w3.cell(rr, col).number_format = "DD.MM.YYYY"
    w3.freeze_panes = "A4"

    _sheet_uebersicht(wb, cons, projn, ent, mem, today, titel)
    if val is not None:
        _sheet_guete(wb, cons, projn, ent, val, today, titel)
        wb.move_sheet("Planungsgüte", offset=-(len(wb.sheetnames) - 3))
        wb.move_sheet("VAL Details", offset=-(len(wb.sheetnames) - 3))
    wb.active = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
