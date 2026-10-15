"""Ressourcenplanung | Operations Team
Start: streamlit run app.py
"""
from __future__ import annotations

import base64
import datetime as dt
import hmac
import html
import os
import re
import time

import pandas as pd
import plotly.graph_objects as go
from plotly.subplots import make_subplots
import streamlit as st

import db
import export
import logic

st.set_page_config(page_title="Ressourcenplanung", page_icon="📅", layout="wide")

BASE = os.path.dirname(os.path.abspath(__file__))
LOGO_PATH = os.path.join(BASE, "assets", "wp_logo.png")

# Farbcode (Corporate Design)
WP = dict(navy="#000055", mid="#6688BB", light="#BBCCEE", lilac="#E4E4F6",
          grey="#999999", lgrey="#D2D2D2", xlgrey="#EBEBEB", dark="#00295F", alert="#B3261E")
SERIES = ["#000055", "#6688BB", "#BBCCEE", "#999999", "#00295F", "#63686A",
          "#D2D2D2", "#8FA6CF", "#4D4D8F", "#B8B8B8"]
FONT = "Calibri, Carlito, 'Segoe UI', Arial, sans-serif"


def secret(section: str, key: str, default=None):
    try:
        return st.secrets[section][key]
    except Exception:
        return default


ZIEL = float(secret("app", "zielauslastung", 0.80))
LOGO_TOKEN = secret("app", "logo_dev_token", None) or None
# Markenname nur in den Secrets, nicht im (ggf. öffentlichen) Code
MARKE = secret("app", "marke", "Ressourcenplanung")
EXPORT_TITEL = secret("app", "export_titel", "AUSLASTUNG")


def _list_secret(section: str, key: str) -> list[str]:
    v = secret(section, key, None)
    if isinstance(v, str):
        v = [v]
    return [str(x).strip() for x in (v or []) if str(x).strip()]


TEAM_PW = secret("auth", "team_passwort", None)
SELBSTREG = str(secret("auth", "selbstregistrierung", "false")).lower() in ("1", "true", "ja", "yes")
ADMINS = {a.lower() for a in _list_secret("auth", "admins")}
# Zusätzlich berechtigt, VAL für alle zu buchen (Kürzel, in den Secrets unter [auth] val_alle)
VAL_ALLE = {a.lower() for a in _list_secret("auth", "val_alle")}
TEAM_SEED = _list_secret("app", "team")
ROLLEN = ["Managing Partner", "Partner", "Mitglied der Geschäftsleitung", "Senior Manager", "Manager",
          "Senior Consultant", "Consultant", "Werkstudent/in"]


def is_admin() -> bool:
    """Ohne konfigurierte Admins (Demo) darf jeder alles."""
    if not ADMINS:
        return True
    return str(st.session_state.get("user", "")).lower() in ADMINS


def _s(v) -> str:
    return "" if v is None or (not isinstance(v, str) and pd.isna(v)) else str(v)


@st.cache_data(ttl=7 * 24 * 3600, show_spinner=False)
def _fetch_logo(url: str, name: str) -> str:
    """Logo serverseitig laden und als Data URI cachen (schnell, kein Tracking pro Aufruf).
    Fallback: Initialen-Badge im Mittelblau."""
    if url.startswith("data:"):
        return url
    try:
        import urllib.request
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=4) as r:
            data = r.read()
            mime = r.headers.get_content_type() or "image/png"
        if len(data) < 200 or not mime.startswith("image/"):
            raise ValueError("kein brauchbares Logo")
        return f"data:{mime};base64,{base64.b64encode(data).decode()}"
    except Exception:
        return db.initials_svg(name)


def logo(name, domain, logo_blob) -> str:
    return _fetch_logo(db.logo_src(name, domain, logo_blob, LOGO_TOKEN), name)


@st.cache_resource(show_spinner="Datenbank wird vorbereitet ...")
def init_once() -> bool:
    """Schema und ggf. Demo-Daten genau einmal pro Serverprozess."""
    db.init_db(team=TEAM_SEED or None)
    return True


def heute() -> dt.date:
    return dt.date.today()


# ====================================================================== Styling
def inject_css() -> None:
    st.markdown(f"""
<style>
html, body, [class*="css"], .stMarkdown, .stDataFrame {{ font-family: {FONT}; }}
h1, h2, h3, h4 {{ color: {WP['navy']}; font-family: {FONT}; }}
.wp-title {{ font-size: 1.55rem; color: {WP['navy']}; margin: 0 0 .15rem 0; }}
.wp-title b {{ font-weight: 700; }}
.wp-line {{ border: 0; border-top: 2px solid {WP['navy']}; margin: .2rem 0 1rem 0; }}
.wp-sub {{ font-style: italic; color: {WP['mid']}; border-bottom: 1px solid {WP['mid']};
          padding-bottom: 2px; margin: 1.2rem 0 .6rem 0; font-size: 1.05rem; }}
.wp-wordmark {{ font-family: 'Goudy Old Style', Garamond, Georgia, serif; color: {WP['navy']};
               font-size: 1.25rem; line-height: 1.15; padding: .2rem 0 .8rem 0; }}
.wp-wordmark small {{ display:block; font-family:{FONT}; font-size:.75rem; color:{WP['mid']};
                     letter-spacing:.06em; text-transform:uppercase; margin-top:.25rem; }}
div[data-testid="stMetric"] {{ background: #F5F6FB; border-left: 4px solid {WP['navy']};
                               padding: .6rem .9rem; border-radius: 2px; }}
div[data-testid="stMetricLabel"] p {{ color: {WP['mid']}; font-size: .85rem; }}
div[data-testid="stMetricValue"] {{ color: {WP['navy']}; }}

table.wpm {{ width: 100%; border-collapse: separate; border-spacing: 4px; table-layout: fixed; }}
table.wpm th {{ color: {WP['navy']}; font-weight: 600; font-size: .85rem; text-align: center;
               padding: 4px 2px; border-bottom: 2px solid {WP['navy']}; }}
table.wpm th.today {{ background: {WP['navy']}; color: #fff; border-radius: 3px 3px 0 0; }}
table.wpm th.name, table.wpm td.name {{ text-align: left; width: 13%; }}
table.wpm th.util, table.wpm td.util {{ width: 11%; }}
table.wpm td {{ height: 62px; vertical-align: middle; text-align: center; border-radius: 4px;
               font-size: .78rem; color: {WP['navy']}; padding: 3px; overflow: hidden; }}
table.wpm td.name {{ font-weight: 600; font-size: .92rem; background: transparent; }}
table.wpm td.name span {{ display:block; font-weight:400; font-size:.72rem; color:{WP['grey']}; }}
td.c-onsite {{ background: {WP['lilac']}; border: 1px solid {WP['light']}; }}
td.c-remote {{ background: #FFFFFF; border: 1px dashed {WP['mid']}; }}
td.c-intern {{ background: {WP['xlgrey']}; color: #555 !important; }}
td.c-absent {{ background: repeating-linear-gradient(45deg,#F2F2F2,#F2F2F2 6px,#E6E6E6 6px,#E6E6E6 12px);
              color: {WP['grey']} !important; }}
td.c-free {{ background: #FFFFFF; border: 1px solid {WP['xlgrey']}; color: {WP['lgrey']} !important; }}
td.c-hol {{ background: {WP['xlgrey']}; color: {WP['grey']} !important; font-style: italic; }}
td.today-col {{ box-shadow: inset 0 0 0 2px {WP['navy']}; }}
td img.lg {{ height: 26px; max-width: 70%; object-fit: contain; display:block; margin: 0 auto 2px auto; }}
td .cl {{ display:block; font-weight:600; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }}
td .pj {{ display:block; font-size:.68rem; color:{WP['mid']}; white-space:nowrap; overflow:hidden;
         text-overflow:ellipsis; }}
.ubar {{ background:{WP['xlgrey']}; height: 8px; border-radius: 4px; overflow:hidden; margin-top:4px; }}
.ubar > div {{ height: 100%; background:{WP['navy']}; }}
.legend span {{ display:inline-block; padding: 2px 10px; margin-right: 8px; border-radius: 3px;
               font-size: .78rem; color:{WP['navy']}; }}
</style>""", unsafe_allow_html=True)


def header(titel: str, kernsatz: str) -> None:
    st.markdown(f"<div class='wp-title'><b>{html.escape(titel)}</b> | {html.escape(kernsatz)}</div>"
                f"<hr class='wp-line'>", unsafe_allow_html=True)


def sub(text: str) -> None:
    st.markdown(f"<div class='wp-sub'>{html.escape(text)}</div>", unsafe_allow_html=True)


def branding() -> None:
    if os.path.exists(LOGO_PATH):
        st.logo(LOGO_PATH, size="large")
    else:
        st.sidebar.markdown(f"<div class='wp-wordmark'>{html.escape(MARKE)}"
                            "<small>Ressourcenplanung Operations</small></div>", unsafe_allow_html=True)


def plotly_layout(fig: go.Figure, height: int = 380, pct: bool = False) -> go.Figure:
    fig.update_layout(
        height=height, margin=dict(l=10, r=10, t=30, b=10), plot_bgcolor="#FFFFFF",
        paper_bgcolor="#FFFFFF", font=dict(family=FONT, color=WP["navy"], size=13),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, x=0, title=None),
        hoverlabel=dict(font_family=FONT))
    fig.update_xaxes(showgrid=False, linecolor=WP["lgrey"])
    fig.update_yaxes(gridcolor=WP["xlgrey"], zeroline=False)
    if pct:
        fig.update_yaxes(tickformat=".0%")
    return fig


def fmt_pct(v) -> str:
    return "n/a" if v is None or pd.isna(v) else f"{v:.0%}"


# ====================================================================== Login
def users() -> dict:
    try:
        return dict(st.secrets["users"])
    except Exception:
        return {}


def login_gate() -> bool:
    if st.session_state.get("user"):
        return True
    # Vor der Anmeldung keine Seitenleiste und keine Hinweise auf Team oder Inhalt
    st.markdown("<style>[data-testid=stSidebar],[data-testid=stSidebarCollapsedControl]{display:none}</style>",
                unsafe_allow_html=True)
    header("Ressourcenplanung", "Anmeldung")
    names = db.consultants()["name"].tolist()
    if TEAM_PW:
        with st.form("login"):
            vn = st.text_input("Vor- und Nachname", placeholder="z. B. Max Mustermann")
            pw = st.text_input("Team-Passwort", type="password")
            if st.form_submit_button("Anmelden", type="primary"):
                ok = bool(vn.strip()) and hmac.compare_digest(str(TEAM_PW), pw)
                person = db.find_consultant(vn) if ok else None
                if ok and person is None and SELBSTREG and len(vn.split()) >= 2:
                    kz = db.suggest_kuerzel(vn)
                    db.create_consultant(kz, vn.strip())
                    person = db.find_consultant(vn)
                if person is not None:
                    st.session_state.user = person["name"]
                    st.session_state.user_label = _s(person.get("vollname")) or person["name"]
                    st.rerun()
                time.sleep(1.5)  # bremst Durchprobieren
                if ok:
                    st.error("Name nicht im Team gefunden. Bitte Schreibweise prüfen oder den Admin bitten, "
                             "dich anzulegen.")
                else:
                    st.error("Name oder Passwort falsch.")
        return False
    u = users()
    if not u:
        st.info("Demo-Modus: kein Passwort in den Secrets hinterlegt. Name wählen genügt.")
        name = st.selectbox("Wer bist du?", names, index=None)
        if st.button("Weiter", type="primary", disabled=not name):
            st.session_state.user = name
            st.rerun()
        return False
    with st.form("login"):
        eingabe = st.text_input("Benutzername")
        pw = st.text_input("Passwort", type="password")
        if st.form_submit_button("Anmelden", type="primary"):
            # Groß/Kleinschreibung und Leerzeichen beim Benutzernamen egal
            name = next((k for k in u if k.strip().lower() == eingabe.strip().lower()), None)
            kz, soll = login_eintrag(name, u.get(name)) if name else (None, None)
            if kz and soll is not None and hmac.compare_digest(str(soll), pw):
                st.session_state.user = kz
                st.session_state.login_name = name
                st.rerun()
            time.sleep(1.5)
            st.error("Benutzername oder Passwort falsch.")
    return False


def login_eintrag(name: str, wert) -> tuple[str | None, str | None]:
    """Secrets [users]: entweder  kürzel = "passwort"  (alt)  oder
    "vorname.nachname" = { kuerzel = "abc", passwort = "..." }  (neu, eigener Login je Person)."""
    if hasattr(wert, "get") and not isinstance(wert, str):
        return str(wert.get("kuerzel") or name).strip(), wert.get("passwort")
    return name, wert


# ====================================================================== Gemeinsame Bausteine
SONST = ["Urlaub", "Intern", "Akquise", "Krank"]
ORT_ICON = {"Vor Ort": "🏢", "Remote": "🏠"}
ART_ICON = {"Urlaub": "🌴", "Intern": "🏛️", "Akquise": "🎯", "Krank": "🤒", "Elternzeit": "👶", "Löschen": "🧽"}
H = logic.TAG_H


def extra_css() -> None:
    st.markdown(f"""
<style>
.dh {{ display:flex; justify-content:space-between; align-items:center; font-weight:700; color:{WP['navy']};
      font-size:.88rem; margin-bottom:4px; }}
.dh.today > span:first-child {{ background:{WP['navy']}; color:#fff; padding:1px 7px; border-radius:4px; }}
.hb {{ font-size:.72rem; padding:1px 7px; border-radius:10px; font-weight:600; white-space:nowrap; }}
.hb.ok {{ background:{WP['lilac']}; color:{WP['navy']}; }}
.hb.low {{ background:#F2F2F2; color:{WP['grey']}; }}
.hb.over {{ background:#FBE3E1; color:{WP['alert']}; }}
.ent2 {{ display:flex; gap:6px; align-items:center; font-size:.76rem; line-height:1.2; color:{WP['navy']};
        padding:4px 5px; border-radius:4px; overflow:hidden; }}
.ent2 > div {{ overflow:hidden; min-width:0; }}
.ent2 b, .ent2 small, .ent2 i {{ display:block; white-space:nowrap; overflow:hidden; text-overflow:ellipsis; }}
.ent2 small {{ color:{WP['navy']}; font-size:.72rem; }}
.ent2 i {{ color:{WP['mid']}; font-size:.68rem; font-style:normal; }}
.ent2 img {{ height:20px; width:20px; object-fit:contain; flex:0 0 20px; }}
.ent2.onsite {{ background:{WP['lilac']}; border-left:4px solid {WP['navy']}; }}
.ent2.remote {{ background:#FFFFFF; border:1px dashed {WP['mid']}; border-left:4px dashed {WP['mid']}; }}
.ent2.intern {{ background:{WP['xlgrey']}; }}
.ent2.absent {{ background:repeating-linear-gradient(45deg,#F4F4F4,#F4F4F4 6px,#EAEAEA 6px,#EAEAEA 12px); color:#777; }}
.brush {{ background:{WP['lilac']}; border:1px solid {WP['light']}; border-radius:6px; padding:8px 12px;
         color:{WP['navy']}; margin:4px 0 8px 0; }}
.brush b {{ font-size:1.02rem; }}
.wk {{ color:{WP['mid']}; font-size:.85rem; }}
.kpis {{ display:flex; gap:22px; flex-wrap:wrap; }}
.kpis div b {{ display:block; font-size:1.3rem; color:{WP['navy']}; line-height:1.1; }}
.kpis div span {{ font-size:.74rem; color:{WP['mid']}; }}
.kpis div b.neg {{ color:{WP['alert']}; }}
.bar {{ height:9px; background:{WP['xlgrey']}; border-radius:5px; overflow:hidden; display:flex; margin-top:8px; }}
.bar .g {{ background:{WP['navy']}; }} .bar .p {{ background:{WP['mid']}; }} .bar .o {{ background:{WP['alert']}; }}
.ptitle b {{ font-size:1.05rem; color:{WP['navy']}; }} .ptitle span {{ color:{WP['mid']}; font-size:.85rem; }}
td .ent {{ display:flex; align-items:center; justify-content:center; gap:4px; font-size:.7rem; white-space:nowrap;
          overflow:hidden; text-overflow:ellipsis; line-height:1.25; }}
td .ent img {{ height:16px; width:16px; object-fit:contain; }}
td.over {{ box-shadow: inset 0 0 0 2px {WP['alert']} !important; }}
.legend2 span {{ display:inline-block; padding:2px 10px; margin-right:8px; border-radius:3px; font-size:.78rem;
               color:{WP['navy']}; }}
/* Mein Bereich: Projektkarten */
.pc-h {{ display:flex; gap:10px; align-items:center; }}
.pc-h img {{ width:42px; height:42px; object-fit:contain; flex:0 0 42px; }}
.pc-h > div {{ min-width:0; flex:1; }}
.pc-h b {{ display:block; font-size:1.02rem; color:{WP['navy']}; white-space:nowrap; overflow:hidden;
          text-overflow:ellipsis; }}
.pc-h span {{ display:block; font-size:.78rem; color:{WP['mid']}; white-space:nowrap; overflow:hidden;
             text-overflow:ellipsis; min-height:1.1em; }}
.ort {{ font-style:normal; font-size:.72rem; padding:2px 9px; border-radius:10px; white-space:nowrap;
       color:{WP['navy']}; font-weight:600; }}
.ort.onsite {{ background:{WP['lilac']}; }} .ort.remote {{ border:1px dashed {WP['mid']}; }}
.ort.mixed {{ background:#F5F6FB; border:1px solid {WP['light']}; }}
.ort.none {{ background:#F2F2F2; color:{WP['grey']}; font-weight:400; }}
.pc-n {{ display:flex; gap:14px; flex-wrap:wrap; font-size:.74rem; color:{WP['mid']}; margin-top:7px;
        align-items:baseline; }}
.pc-n b {{ color:{WP['navy']}; font-size:1rem; }} .pc-n b.neg {{ color:{WP['alert']}; }}
.pc-n .von {{ margin-left:auto; color:{WP['grey']}; }}
.pc-f {{ font-size:.75rem; color:{WP['grey']}; margin:5px 0 2px 0; }} .pc-f b {{ color:{WP['navy']}; }}
.pc-add {{ min-height:118px; display:flex; flex-direction:column; justify-content:center; gap:4px;
          color:{WP['mid']}; font-size:.8rem; }}
.pc-add > b {{ color:{WP['navy']}; font-size:1.02rem; }}
/* Akquise-Board */
.aq-col {{ display:flex; justify-content:space-between; align-items:center; background:#F5F6FB; padding:8px 12px;
          border-radius:4px; color:{WP['navy']}; margin-bottom:6px; }}
.aq-col b {{ font-size:1.02rem; }} .aq-col span {{ background:#fff; border-radius:10px; padding:0 9px; font-size:.8rem; }}
.aq-h {{ display:flex; gap:10px; align-items:center; }}
.aq-h img {{ width:36px; height:36px; object-fit:contain; flex:0 0 36px; }}
.aq-h b {{ display:block; color:{WP['navy']}; font-size:1rem; }}
.aq-h span {{ display:block; color:{WP['mid']}; font-size:.8rem; }}
.aq-n {{ display:flex; gap:14px; margin-top:8px; font-size:.74rem; color:{WP['mid']}; }}
.aq-n b {{ color:{WP['navy']}; font-size:.95rem; }}
.aq-w {{ height:6px; background:{WP['xlgrey']}; border-radius:3px; margin-top:8px; overflow:hidden; }}
.aq-w > div {{ height:100%; background:{WP['navy']}; }}
.aq-wl {{ font-size:.7rem; color:{WP['grey']}; margin-top:3px; }}
.aq-t {{ margin-top:8px; display:flex; flex-wrap:wrap; gap:4px; }}
.aq-t i {{ font-size:.72rem; color:{WP['grey']}; }}
.aq-p {{ background:{WP['lilac']}; color:{WP['navy']}; border-radius:10px; padding:1px 8px; font-size:.72rem;
        font-weight:600; }}
/* Team-Board */
.regel {{ border-left:3px solid {WP['navy']}; background:#F5F6FB; padding:7px 12px; border-radius:3px;
         color:{WP['navy']}; font-size:.95rem; }}
.mt-t {{ color:{WP['navy']}; font-size:.95rem; white-space:pre-wrap; }}
.mt-m {{ color:{WP['grey']}; font-size:.75rem; margin-top:2px; }}
/* Gesamtübersicht: Kurzfrist 3 Wochen */
.kfw {{ overflow-x:auto; }}
table.kf {{ width:100%; border-collapse:separate; border-spacing:2px; table-layout:fixed; min-width:980px; }}
table.kf th {{ font-size:.68rem; color:{WP['navy']}; font-weight:600; text-align:center; padding:2px 0;
              border-bottom:1px solid {WP['light']}; white-space:nowrap; overflow:hidden; }}
table.kf th.wk {{ background:{WP['navy']}; color:#fff; border-radius:3px; font-size:.72rem; border:0; }}
table.kf th.td {{ background:{WP['lilac']}; border-radius:3px; }}
table.kf th.nm, table.kf td.nm {{ width:96px; text-align:left; }}
table.kf td.nm b {{ font-size:.76rem; color:{WP['navy']}; white-space:nowrap; overflow:hidden; text-overflow:ellipsis;
                   display:block; }}
table.kf td {{ height:40px; vertical-align:middle; text-align:center; border-radius:3px; padding:1px;
              background:#fff; border:1px solid {WP['xlgrey']}; overflow:hidden; }}
table.kf td.nm {{ border:0; background:transparent; }}
table.kf tr.grp td {{ height:auto; text-align:left; font-size:.72rem; font-weight:700; letter-spacing:.04em;
              text-transform:uppercase; background:{WP['lilac']}; border:0; border-radius:2px; padding:3px 8px; }}
table.kf.gross tr.grp td {{ height:auto; }}
table.kf td.fr {{ background:#fff; }}
table.kf td.ho {{ background:{WP['xlgrey']}; color:{WP['grey']}; font-size:.6rem; font-style:italic; }}
table.kf td.td {{ box-shadow: inset 0 0 0 1.5px {WP['navy']}; }}
table.kf td.ov {{ box-shadow: inset 0 0 0 2px {WP['alert']}; }}
table.kf th.au, table.kf td.au {{ width:58px; font-size:.66rem; font-weight:700; border:0; }}
table.kf .k {{ display:inline-flex; align-items:center; gap:3px; font-size:.72rem; color:{WP['navy']};
              border-radius:3px; padding:1px 3px; margin:1px 0; line-height:26px; white-space:nowrap; font-weight:600; }}
table.kf .k i {{ width:26px; height:26px; display:inline-block; background-size:contain;
                background-repeat:no-repeat; background-position:center; }}
table.kf.gross td {{ height:50px; }}
table.kf.gross .k {{ font-size:.8rem; line-height:32px; }}
table.kf.gross .k i {{ width:34px; height:32px; }}
table.kf .k.bl {{ opacity:.35; }}
table.kf th.ex, table.kf td.ex {{ width:62px; border:0; background:transparent; font-size:.95rem; }}
table.kf td.ex .vn {{ font-size:.62rem; color:{WP['grey']}; font-weight:400; }}
table.kf .k.on {{ background:{WP['lilac']}; border-left:2px solid {WP['navy']}; }}
table.kf .k.re {{ background:#fff; border:1px dashed {WP['mid']}; }}
table.kf .k.in {{ background:{WP['xlgrey']}; }}
table.kf .k.ab {{ background:repeating-linear-gradient(45deg,#F4F4F4,#F4F4F4 4px,#EAEAEA 4px,#EAEAEA 8px); color:#777; }}
/* Mein Bereich: 12-Wochen-Kalender */
.calh {{ font-size:.8rem; color:{WP['navy']}; font-weight:700; border-bottom:2px solid {WP['navy']};
        padding:0 4px 2px 4px; }}
.calkw b {{ display:block; color:{WP['navy']}; font-size:.92rem; }} .calkw span {{ font-size:.68rem; color:{WP['grey']}; }}
.calu {{ background:#F5F6FB; border-radius:6px; padding:6px 8px; font-size:.7rem; color:{WP['mid']}; }}
.calu b {{ font-size:1rem; }}
.cday {{ min-height:52px; }}
.cday .dn {{ font-size:.64rem; color:{WP['grey']}; margin-bottom:2px; }}
.cday .frei {{ font-size:.72rem; color:{WP['lgrey']}; }}
div[class*="st-key-calday_"] {{ position:relative; padding:4px 6px !important; gap:0 !important;
                               transition: background .12s, border-color .12s; }}
div[class*="st-key-calday_"]:hover {{ background:#F5F6FB; border-color:{WP['mid']} !important; cursor:pointer; }}
div[class*="st-key-calday_"][class*="_today"] {{ box-shadow: inset 0 0 0 2px {WP['navy']}; }}
div[class*="st-key-calday_"][class*="_over"] {{ box-shadow: inset 0 0 0 2px {WP['alert']}; }}
div[class*="st-key-calday_"][class*="_past"] .cday {{ opacity:.55; }}
div[class*="st-key-calday_"][class*="_hol"] {{ background:{WP['xlgrey']}; }}
div[class*="st-key-calbtn_"] {{ position:absolute !important; inset:0; z-index:3; margin:0 !important; }}
div[class*="st-key-calbtn_"] div {{ height:100%; width:100%; }}
div[class*="st-key-calbtn_"] button {{ width:100%; height:100%; opacity:0; cursor:pointer; }}
table.cal {{ width:100%; border-collapse:separate; border-spacing:4px; table-layout:fixed; }}
table.cal th {{ font-size:.8rem; color:{WP['navy']}; font-weight:700; text-align:left; padding:2px 6px;
               border-bottom:2px solid {WP['navy']}; }}
table.cal th.kw, table.cal td.kw {{ width:96px; }}
table.cal th.u, table.cal td.u {{ width:104px; }}
table.cal td {{ vertical-align:top; border-radius:6px; padding:3px 4px; background:#fff;
               border:1px solid {WP['xlgrey']}; height:50px; overflow:hidden; }}
table.cal td.kw {{ border:0; background:transparent; vertical-align:middle; }}
table.cal td.kw b {{ display:block; color:{WP['navy']}; font-size:.92rem; }}
table.cal td.kw span {{ font-size:.68rem; color:{WP['grey']}; }}
table.cal td.hol {{ background:{WP['xlgrey']}; color:{WP['grey']}; font-style:italic; font-size:.72rem; }}
table.cal td.today {{ box-shadow: inset 0 0 0 2px {WP['navy']}; }}
table.cal td.over {{ box-shadow: inset 0 0 0 2px {WP['alert']}; }}
table.cal td.past {{ opacity:.55; }}
table.cal td.u {{ border:0; background:#F5F6FB; vertical-align:middle; font-size:.7rem; color:{WP['mid']}; }}
table.cal td.u b {{ font-size:1rem; }}
table.cal .dn {{ font-size:.64rem; color:{WP['grey']}; margin-bottom:2px; }}
table.cal .frei {{ font-size:.72rem; color:{WP['lgrey']}; }}
.chip {{ display:flex; align-items:center; gap:5px; border-radius:4px; padding:2px 5px; margin-bottom:2px;
        font-size:.74rem; color:{WP['navy']}; overflow:hidden; white-space:nowrap; }}
.chip i {{ width:16px; height:16px; flex:0 0 16px; background-size:contain; background-repeat:no-repeat;
          background-position:center; }}
.chip b {{ overflow:hidden; text-overflow:ellipsis; font-weight:600; }}
.chip small {{ color:{WP['mid']}; font-size:.64rem; margin-left:auto; padding-left:3px; }}
.chip.onsite {{ background:{WP['lilac']}; border-left:3px solid {WP['navy']}; }}
.chip.remote {{ background:#fff; border:1px dashed {WP['mid']}; }}
.chip.intern {{ background:{WP['xlgrey']}; }}
.chip.absent {{ background:repeating-linear-gradient(45deg,#F4F4F4,#F4F4F4 6px,#EAEAEA 6px,#EAEAEA 12px); color:#666; }}
</style>""", unsafe_allow_html=True)


def bt(v) -> str:
    """Beratertage hübsch: 12 / 12,5"""
    v = round(float(v or 0), 1)
    return f"{v:g}".replace(".", ",")


def hh(v) -> str:
    return f"{float(v):g}".replace(".", ",") + " h"


def client_logo_map(proj_all: pd.DataFrame) -> dict:
    return {r.id: logo(r.kunde, r.domain, r.logo) for r in proj_all.itertuples()}


def week_days(ws: dt.date) -> list[dt.date]:
    return [ws + dt.timedelta(days=i) for i in range(5)]


def person_label(cons: pd.DataFrame) -> dict:
    return {r.name: (f"{_s(r.vollname)} ({r.name})" if _s(r.vollname) else r.name) for r in cons.itertuples()}


def art_label(e) -> str:
    """Anzeige für Nicht-Projekt-Einträge: Team-Termine mit ihrem Namen."""
    n = getattr(e, "notiz", None)
    if e.art == "Intern" and isinstance(n, str) and n.strip():
        return f"📌 {n.strip()}"
    return f"{ART_ICON.get(e.art, '')} {e.art}"


def entry_html(e, projn: pd.DataFrame, logos: dict) -> str:
    h = hh(e.stunden)
    if pd.notna(e.project_id) and int(e.project_id) in projn.index:
        pid = int(e.project_id)
        p = projn.loc[pid]
        cls = "onsite" if e.art == "Vor Ort" else "remote"
        return (f"<div class='ent2 {cls}' title='{html.escape(p['label'])} · {e.art} · {h}'>"
                f"<img src='{logos[pid]}'><div><b>{html.escape(p['kunde'])}</b>"
                f"<small>{ORT_ICON.get(e.art, '')} {e.art} · {h}</small><i>{html.escape(p['name'])}</i></div></div>")
    cls = "absent" if e.art in logic.ABWESEND_ARTEN else "intern"
    return (f"<div class='ent2 {cls}'><div><b>{html.escape(art_label(e))}</b>"
            f"<small>{h}</small></div></div>")


def day_header_html(d: dt.date, tot: float, today: dt.date, abw: float = 0.0) -> str:
    t = " today" if d == today else ""
    if logic.feiertag_name(d):
        badge = ""
    elif abw > 0 and tot - abw <= 0:
        badge = "<span class='hb low'>abwesend</span>"
    elif tot <= 0:
        badge = "<span class='hb low'>frei</span>"
    elif tot > H:
        badge = f"<span class='hb over'>⚠ {hh(tot)}</span>"
    elif tot < H:
        badge = f"<span class='hb low'>{hh(tot)} · {hh(H - tot)} frei</span>"
    else:
        badge = "<span class='hb ok'>8 h ✓</span>"
    return f"<div class='dh{t}'><span>{logic.WOTAG[d.weekday()]} {d:%d.%m.}</span>{badge}</div>"


# ====================================================================== Seite: Mein Kalender
def current_brush():
    """('p', project_id) | ('a', 'Urlaub' ...) | ('a', 'Löschen') | None"""
    ss = st.session_state
    other = ss.get("k_other")
    if other and other in ss.get("k_other_map", {}):
        return ("p", int(ss["k_other_map"][other]))
    b = ss.get("k_brush")
    if not b:
        return None
    return ("p", int(b[1:])) if b.startswith("p") else ("a", b)


def _valid_days(days) -> list[dt.date]:
    return [d for d in days if d.weekday() < 5 and not logic.feiertag_name(d)]


def apply_brush(cid: int, days: list[dt.date]) -> None:
    ss = st.session_state
    br = current_brush()
    days = _valid_days(days)
    if br is None:
        st.toast("Bitte zuerst oben auswählen, was eingetragen werden soll.")
        return
    if not days:
        return
    if br == ("a", "Löschen"):
        n = db.clear_days(cid, days)
        st.toast(f"🧽 {n} Einträge entfernt")
        return
    h = float(ss.get("k_h") or H)
    if br[0] == "p":
        ort = ss.get("k_ort") or "Vor Ort"
        rows = [(cid, d, ort, br[1], h) for d in days]
        db.add_member(br[1], cid)  # taucht damit unter "Meine Projekte" auf
    else:
        rows = [(cid, d, br[1], None, h) for d in days]
    db.upsert_entries(rows, ss.get("user", ""))
    e = db.entries(min(days), max(days), consultant_id=cid)
    tot = e.groupby("tag")["stunden"].sum() if not e.empty else pd.Series(dtype=float)
    over = [d for d in days if tot.get(d, 0) > H]
    msg = f"✓ {len(rows)} Tag{'e' if len(rows) != 1 else ''} eingetragen"
    if over:
        msg += " · ⚠ überbucht am " + ", ".join(f"{d:%d.%m.}" for d in over[:5])
    st.toast(msg)


def del_entry_cb(entry_id: int) -> None:
    db.delete_entry(entry_id)
    st.toast("Eintrag entfernt")


def clear_range_cb(cid: int, von: dt.date, bis: dt.date) -> None:
    n = db.clear_range([cid], von, bis)
    st.toast(f"🧽 {n} Einträge im Zeitraum entfernt")


def _shift_k(n: int) -> None:
    st.session_state.k_off = 0 if n == 0 else st.session_state.get("k_off", 0) + n


def _clear_other() -> None:
    st.session_state.k_other = None


def kalender_body(cid: int, cons: pd.DataFrame) -> None:
    """Tagesansicht: Pinsel wählen, Tage anklicken (wird in 'Mein Bereich' eingebettet)."""
    today = heute()
    ss = st.session_state
    proj = db.projects()
    pmap = proj.set_index("id")
    mem = db.members()
    ms = logic.member_stats(db.entries(consultant_id=cid, project_only=True),
                            mem[mem["consultant_id"] == cid], today).set_index("project_id")
    mine = [pid for pid in pmap.index if pid in ms.index]
    others = [pid for pid in pmap.index if pid not in ms.index]

    # ---------------- 1 | Pinsel
    sub("1 | Was trage ich ein?")

    def lab(o: str) -> str:
        if o.startswith("p"):
            pid = int(o[1:])
            t = pmap.at[pid, "label"]
            if pid in ms.index:
                b = ms.loc[pid]
                t += f"  ({bt(b['offen'])} von {bt(b['budget'])} BT offen)" if b["budget"] > 0 else "  (kein Budget)"
            return t
        return f"{ART_ICON.get(o, '')} {o}"

    opts = [f"p{p}" for p in mine] + SONST + ["Löschen"]
    if ss.get("k_brush") not in opts:
        ss.k_brush = opts[0]
    st.pills("Meine Projekte und Abwesenheiten", opts, format_func=lab, key="k_brush",
             selection_mode="single", on_change=_clear_other)
    other_map = {pmap.at[p, "label"]: p for p in others}
    ss["k_other_map"] = other_map
    cc = st.columns([3, 2, 2])
    cc[0].selectbox("… oder ein anderes Projekt", list(other_map), index=None, key="k_other",
                    placeholder="Projekt suchen, bei dem ich noch nicht eingeplant bin")
    br = current_brush()
    is_p = br is not None and br[0] == "p"
    cc[1].segmented_control("Ort", ["Vor Ort", "Remote"], default="Vor Ort", key="k_ort",
                            format_func=lambda x: f"{ORT_ICON[x]} {x}", disabled=not is_p)
    cc[2].segmented_control("Stunden pro Tag", [2, 4, 6, 8], default=8, key="k_h",
                            format_func=lambda x: f"{x} h", disabled=br == ("a", "Löschen"))
    h_sel = float(ss.get("k_h") or H)
    if br is None:
        st.markdown("<div class='brush'>Bitte oben auswählen, was eingetragen werden soll.</div>",
                    unsafe_allow_html=True)
        short = "eintragen"
    elif br == ("a", "Löschen"):
        st.markdown("<div class='brush'>🧽 <b>Radiergummi</b>: angeklickte Tage werden geleert.</div>",
                    unsafe_allow_html=True)
        short = "🧽 leeren"
    elif is_p:
        ort = ss.get("k_ort") or "Vor Ort"
        p = pmap.loc[br[1]]
        st.markdown(f"<div class='brush'>Pinsel: <b>{ORT_ICON[ort]} {html.escape(p['kunde'])}</b> · "
                    f"{html.escape(p['name'])} · {ort} · {hh(h_sel)} pro Tag</div>", unsafe_allow_html=True)
        short = f"＋ {p['kunde'][:14]}"
    else:
        st.markdown(f"<div class='brush'>Pinsel: <b>{ART_ICON.get(br[1], '')} {br[1]}</b> · {hh(h_sel)} pro Tag</div>",
                    unsafe_allow_html=True)
        short = f"＋ {br[1]}"

    # ---------------- 2 | Kalender
    sub("2 | Tage anklicken")
    ss.setdefault("k_off", 0)
    nb = st.columns([1.2, 1, 1.2, 2, 3])
    nb[0].button("◀ Früher", on_click=_shift_k, args=(-1,), width="stretch")
    nb[1].button("Heute", on_click=_shift_k, args=(0,), width="stretch")
    nb[2].button("Später ▶", on_click=_shift_k, args=(1,), width="stretch")
    n_w = nb[3].selectbox("Wochen", [2, 3, 4, 6], index=1, key="k_nw", label_visibility="collapsed",
                          format_func=lambda x: f"{x} Wochen anzeigen")
    start = logic.monday(today) + dt.timedelta(weeks=ss.k_off)
    weeks = [start + dt.timedelta(weeks=i) for i in range(n_w)]
    end = weeks[-1] + dt.timedelta(days=4)
    ent = db.entries(start, end, consultant_id=cid)
    projn = db.projects(active_only=False).set_index("id")
    logos = client_logo_map(projn.reset_index())
    by_day = {d: g for d, g in ent.groupby("tag")} if not ent.empty else {}
    st.markdown("<div class='legend2'><span style='background:#E4E4F6;border-left:4px solid #000055'>🏢 Vor Ort</span>"
                "<span style='border:1px dashed #6688BB'>🏠 Remote</span>"
                "<span style='background:#FBE3E1;color:#B3261E'>⚠ mehr als 8 h = überbucht</span></div>",
                unsafe_allow_html=True)

    for ws in weeks:
        wd = week_days(ws)
        we = ent[ent["tag"].isin(wd)] if not ent.empty else ent
        f = logic.build_fact(ws, wd[-1], cons[cons["id"] == cid], we)
        hc = st.columns([4, 1.3, 1.3], vertical_alignment="bottom")
        hc[0].markdown(f"**KW {ws.isocalendar().week:02d}** &nbsp; {ws:%d.%m.} bis {wd[-1]:%d.%m.} &nbsp; "
                       f"<span class='wk'>{hh(f['gebucht'].sum())} geplant von {hh(f['kapazitaet'].sum())}</span>",
                       unsafe_allow_html=True)
        hc[1].button("Mo bis Do", key=f"w4_{cid}_{ws}", on_click=apply_brush, args=(cid, wd[:4]),
                     disabled=br is None, width="stretch", help="Auswahl für Mo bis Do dieser Woche eintragen")
        hc[2].button("Mo bis Fr", key=f"w5_{cid}_{ws}", on_click=apply_brush, args=(cid, wd),
                     disabled=br is None, width="stretch", help="Auswahl für die ganze Woche eintragen")
        cols = st.columns(5, gap="small")
        for i, d in enumerate(wd):
            with cols[i].container(border=True):
                es = by_day.get(d)
                tot = float(es["stunden"].sum()) if es is not None else 0.0
                abw = float(es.loc[es["art"].isin(logic.ABWESEND_ARTEN), "stunden"].sum()) if es is not None else 0.0
                st.markdown(day_header_html(d, tot, today, abw), unsafe_allow_html=True)
                hol = logic.feiertag_name(d)
                if hol:
                    st.caption(f"Feiertag: {hol}")
                    continue
                if es is not None:
                    for e in es.sort_values(["project_id", "art"]).itertuples():
                        ec = st.columns([6, 1], vertical_alignment="center", gap="small")
                        ec[0].markdown(entry_html(e, projn, logos), unsafe_allow_html=True)
                        ec[1].button("✕", key=f"del_{e.id}", on_click=del_entry_cb, args=(int(e.id),),
                                     type="tertiary", help="Eintrag entfernen")
                st.button(short, key=f"add_{cid}_{d}", on_click=apply_brush, args=(cid, [d]),
                          disabled=br is None, width="stretch")
        st.write("")

    with st.expander("Serie eintragen (z. B. Mo bis Do jede Woche bis Projektende) oder Zeitraum leeren"):
        s1, s2, s3, s4 = st.columns(4)
        von = s1.date_input("Von", today, format="DD.MM.YYYY", key="k_s_von")
        bis = s2.date_input("Bis", today + dt.timedelta(weeks=8), format="DD.MM.YYYY", key="k_s_bis")
        wt = s3.multiselect("Wochentage", logic.WOTAG[:5], default=logic.WOTAG[:4], key="k_s_wt")
        rh = s4.selectbox("Rhythmus", list(logic.RHYTHMEN), key="k_s_rh")
        sdays = logic.plan_days(von, bis, wt, logic.RHYTHMEN[rh])
        st.caption(f"{len(sdays)} Tage mit der aktuellen Auswahl oben ({short.replace('＋ ', '')}, {hh(h_sel)}). "
                   "Bestehende Einträge anderer Projekte bleiben erhalten, Feiertage werden übersprungen.")
        b1, b2, _ = st.columns([1, 1, 3])
        b1.button("Serie eintragen", type="primary", on_click=apply_brush, args=(cid, sdays),
                  disabled=not sdays or br is None, key="k_s_go")
        b2.button("Zeitraum leeren", on_click=clear_range_cb, args=(cid, von, bis), disabled=von > bis,
                  key="k_s_clear", help="Löscht alle Einträge dieser Person zwischen Von und Bis")



def member_projects_table(cid: int) -> None:
    """Projekte eines Beraters mit Budget, gebucht, geplant und offenen Beratertagen."""
    today = heute()
    cons = db.consultants(active_only=False)
    mem = db.members()
    ms = logic.member_stats(db.entries(consultant_id=cid, project_only=True),
                            mem[mem["consultant_id"] == cid], today).set_index("project_id")
    if ms.empty:
        st.caption("Noch keinem Projekt zugeordnet.")
        return
    projn = db.projects(active_only=False).set_index("id")
    logos = client_logo_map(projn.reset_index())
    lmap = dict(zip(cons["id"], cons["name"]))
    rows = []
    for pid, b in ms.iterrows():
        if pid not in projn.index or not projn.at[pid, "aktiv"]:
            continue
        p = projn.loc[pid]
        rows.append({"logo": logos[pid], "Projekt": p["label"], "Budget BT": b["budget"], "Gebucht": b["gebucht"],
                     "Geplant": b["geplant"], "Offen": b["offen"], "Ende": p["ende"],
                     "Leitung": lmap.get(int(p["leiter_id"]), "") if pd.notna(p.get("leiter_id")) else ""})
    if not rows:
        st.caption("Keine aktiven Projekte.")
        return
    df = pd.DataFrame(rows).sort_values("Offen")
    st.dataframe(df.style.map(lambda v: f"color:{WP['alert']};font-weight:700" if v < -0.01 else "",
                              subset=["Offen"]),
                 hide_index=True, width="stretch", column_config={
                     "logo": st.column_config.ImageColumn("", width="small"),
                     "Budget BT": st.column_config.NumberColumn(format="%.1f"),
                     "Gebucht": st.column_config.NumberColumn("Ist", format="%.1f",
                                                              help="VAL, wo gemeldet; sonst geplante Tage bis heute"),
                     "Geplant": st.column_config.NumberColumn("Geplant (ab morgen)", format="%.1f"),
                     "Offen": st.column_config.NumberColumn("Offen BT", format="%.1f"),
                     "Ende": st.column_config.DateColumn(format="DD.MM.YYYY")})
    st.caption("Offen = Budget minus gebucht minus geplant. Negativ (rot) = mehr geplant als budgetiert. "
               "Budget anpassen unter **Projekte > Team & Budget**.")


def week_readonly(cid: int, ws: dt.date) -> None:
    """Woche eines Beraters zum Ansehen (Meeting): je Tag Logos, Ort und Stunden."""
    today = heute()
    wd = week_days(ws)
    ent = db.entries(wd[0], wd[-1], consultant_id=cid)
    projn = db.projects(active_only=False).set_index("id")
    logos = client_logo_map(projn.reset_index())
    by_day = {d: g for d, g in ent.groupby("tag")} if not ent.empty else {}
    cols = st.columns(5, gap="small")
    for i, d in enumerate(wd):
        with cols[i].container(border=True):
            es = by_day.get(d)
            tot = float(es["stunden"].sum()) if es is not None else 0.0
            abw = float(es.loc[es["art"].isin(logic.ABWESEND_ARTEN), "stunden"].sum()) if es is not None else 0.0
            st.markdown(day_header_html(d, tot, today, abw), unsafe_allow_html=True)
            hol = logic.feiertag_name(d)
            if hol:
                st.caption(f"Feiertag: {hol}")
            elif es is None:
                st.caption("nichts geplant")
            else:
                st.markdown("".join(entry_html(e, projn, logos) + "<div style='height:4px'></div>"
                                    for e in es.sort_values(["project_id", "art"]).itertuples()),
                            unsafe_allow_html=True)


# ====================================================================== Seite: Team (Wochenmatrix)
def render_matrix(ws: dt.date) -> str:
    today = heute()
    days = week_days(ws)
    cons = db.consultants()
    ent = db.entries(days[0], days[-1])
    projn = db.projects(active_only=False).set_index("id")
    logos = client_logo_map(projn.reset_index())
    by = {k: g for k, g in ent.groupby(["consultant_id", "tag"])} if not ent.empty else {}
    fact = logic.build_fact(days[0], days[-1], cons, ent)

    head = "<tr><th class='name'>Berater</th>"
    for d in days:
        head += f"<th{' class=today' if d == today else ''}>{logic.day_label(d)}</th>"
    head += "<th class='util'>Auslastung KW</th></tr>"
    rows = ""
    for c in cons.itertuples():
        rows += (f"<tr><td class='name'>{html.escape(c.name)}"
                 f"<span>{html.escape(_s(c.vollname) or _s(c.rolle))}</span></td>")
        for d in days:
            tc = " today-col" if d == today else ""
            hol = logic.feiertag_name(d)
            g = by.get((c.id, d))
            if g is None:
                rows += (f"<td class='c-hol{tc}'>{html.escape(hol)}</td>" if hol
                         else f"<td class='c-free{tc}'>frei</td>")
                continue
            tot = float(g["stunden"].sum())
            arts = set(g["art"])
            cls = ("c-onsite" if "Vor Ort" in arts else "c-remote" if "Remote" in arts
                   else "c-absent" if arts <= set(logic.ABWESEND_ARTEN) else "c-intern")
            over = " over" if tot > H else ""
            inner = ""
            for e in g.sort_values("stunden", ascending=False).head(2).itertuples():
                if pd.notna(e.project_id) and int(e.project_id) in projn.index:
                    p = projn.loc[int(e.project_id)]
                    inner += (f"<div class='ent' title='{html.escape(p['label'])} · {e.art} · {hh(e.stunden)}'>"
                              f"<img src='{logos[int(e.project_id)]}'>{ORT_ICON.get(e.art, '')} "
                              f"<b>{html.escape(p['kunde'][:16])}</b>&nbsp;{hh(e.stunden)}</div>")
                else:
                    inner += f"<div class='ent'>{html.escape(art_label(e))} {hh(e.stunden)}</div>"
            if len(g) > 2:
                inner += f"<div class='ent'>+ {len(g) - 2} weitere</div>"
            if tot > H:
                inner += f"<div class='ent' style='color:{WP['alert']};font-weight:700'>⚠ {hh(tot)}</div>"
            rows += f"<td class='{cls}{tc}{over}'>{inner}</td>"
        fc = fact[fact["consultant_id"] == c.id]
        u = logic.kpi(fc)
        w = 0 if u is None else min(u, 1) * 100
        rows += (f"<td class='util'><b>{fmt_pct(u)}</b><div class='ubar'><div style='width:{w:.0f}%'></div></div>"
                 "</td></tr>")
    legend = (f"<div class='legend'><span style='background:{WP['lilac']};border:1px solid {WP['light']}'>🏢 Vor Ort</span>"
              f"<span style='border:1px dashed {WP['mid']}'>🏠 Remote</span>"
              f"<span style='background:{WP['xlgrey']}'>Intern / Akquise</span>"
              f"<span style='background:#EEE'>Urlaub / Krank</span>"
              f"<span style='box-shadow: inset 0 0 0 2px {WP['alert']}'>⚠ überbucht</span></div>")
    return f"<table class='wpm'>{head}{rows}</table>{legend}"


@st.fragment(run_every=60)
def matrix_live(ws: dt.date) -> None:
    st.markdown(render_matrix(ws), unsafe_allow_html=True)
    st.caption(f"Live-Ansicht, aktualisiert {dt.datetime.now():%H:%M:%S} (automatisch jede Minute)")


def shift_week(n: int) -> None:
    st.session_state.kw_offset = 0 if n == 0 else st.session_state.get("kw_offset", 0) + n


def team_week() -> None:
    today = heute()
    st.session_state.setdefault("kw_offset", 0)
    ws = logic.monday(today) + dt.timedelta(weeks=st.session_state.kw_offset)
    we = ws + dt.timedelta(days=4)
    cons = db.consultants()
    ent_w = db.entries(ws, we)
    f_week = logic.build_fact(ws, we, cons, ent_w)
    m0 = logic.monday(today)
    e4 = m0 + dt.timedelta(days=27)
    f_4w = logic.build_fact(m0, e4, cons, db.entries(m0, e4))
    e_today = db.entries(today, today)
    vor_ort_heute = e_today.loc[e_today["art"] == "Vor Ort", "consultant_id"].nunique() if not e_today.empty else 0
    u_w, u_4 = logic.kpi(f_week), logic.kpi(f_4w)
    over_days = int((f_4w.assign(t=f_4w["gebucht"] + (H - f_4w["kapazitaet"]))["t"] > H).sum()) if not f_4w.empty else 0

    zt = team_ziel(cons)
    k = st.columns(5)
    k[0].metric("Auslastung KW", fmt_pct(u_w),
                None if u_w is None else f"{(u_w - zt) * 100:+.0f} pp vs. Ziel {zt:.0%}",
                help="Ziel = Durchschnitt der Zielauslastungen der Positionen im Team")
    k[1].metric("Heute vor Ort", f"{vor_ort_heute} / {len(cons)}")
    k[2].metric("Auslastung 4 Wochen", fmt_pct(u_4),
                None if u_4 is None else f"{(u_4 - zt) * 100:+.0f} pp vs. Ziel")
    k[3].metric("Frei 4 Wochen (BT)", bt(f_4w["frei"].sum() / H) if not f_4w.empty else "0")
    k[4].metric("Überbucht 4 Wochen", f"{over_days} Tage")

    n = st.columns([1, 1, 1, 7])
    n[0].button("◀ Woche", on_click=shift_week, args=(-1,), width="stretch")
    n[1].button("Heute", on_click=shift_week, args=(0,), width="stretch")
    n[2].button("Woche ▶", on_click=shift_week, args=(1,), width="stretch")
    n[3].markdown(f"**KW {ws.isocalendar().week:02d}** &nbsp; {ws:%d.%m.} bis {we:%d.%m.%Y}")
    matrix_live(ws)

    c1, c2 = st.columns([3, 2])
    with c1:
        sub("Kundenpräsenz dieser Woche (Beratertage)")
        pe = ent_w[ent_w["art"].isin(logic.PROJEKT_ARTEN) & ent_w["project_id"].notna()] if not ent_w.empty else ent_w
        if pe.empty:
            st.caption("Keine Projekttage in dieser Woche.")
        else:
            projn = db.projects(active_only=False).set_index("id")
            pe = pe.assign(kunde=pe["project_id"].astype(int).map(projn["kunde"]),
                           vo=pe["stunden"].where(pe["art"] == "Vor Ort", 0) / H,
                           re=pe["stunden"].where(pe["art"] == "Remote", 0) / H)
            g = pe.groupby("kunde")[["vo", "re"]].sum().reset_index()
            g = g.assign(t=g["vo"] + g["re"]).sort_values("t")
            fig = go.Figure([
                go.Bar(y=g["kunde"], x=g["vo"], name="🏢 Vor Ort", orientation="h", marker_color=WP["navy"]),
                go.Bar(y=g["kunde"], x=g["re"], name="🏠 Remote", orientation="h", marker_color=WP["light"])])
            fig.update_layout(barmode="stack")
            st.plotly_chart(plotly_layout(fig, height=60 + 32 * len(g)), width="stretch")
    with c2:
        sub("Letzte Änderungen")
        ch = db.recent_changes(12)
        if not ch.empty:
            ch["zeitpunkt"] = pd.to_datetime(ch["zeitpunkt"]).dt.strftime("%d.%m. %H:%M")
            ch["eintrag"] = ch["kunde"].fillna(ch["art"]) + " · " + ch["stunden"].map(lambda v: hh(v))
            st.dataframe(ch[["zeitpunkt", "von", "berater", "tag", "eintrag"]], hide_index=True,
                         width="stretch", height=300,
                         column_config={"tag": st.column_config.DateColumn("Tag", format="DD.MM.")})


# ====================================================================== Seite: Projekte
def kpi_html(budget: float, gebucht: float, geplant: float, offen: float) -> str:
    neg = " class='neg'" if offen < -0.01 else ""
    tot = max(budget, gebucht + geplant, 0.01)
    over = max(gebucht + geplant - budget, 0)
    return (f"<div class='kpis'><div><b>{bt(budget)}</b><span>Budget BT</span></div>"
            f"<div><b>{bt(gebucht)}</b><span>Ist (VAL bzw. bis heute)</span></div>"
            f"<div><b>{bt(geplant)}</b><span>geplant ab morgen</span></div>"
            f"<div><b{neg}>{bt(offen)}</b><span>{'offen' if offen >= -0.01 else 'über Budget'}</span></div></div>"
            f"<div class='bar'><div class='g' style='width:{min(gebucht, budget or gebucht) / tot * 100:.1f}%'></div>"
            f"<div class='p' style='width:{max(min(geplant, budget - gebucht), 0) / tot * 100:.1f}%'></div>"
            f"<div class='o' style='width:{over / tot * 100:.1f}%'></div></div>")


def auswertung_projekte() -> None:
    today = heute()
    proj = db.projects()
    stats = logic.project_stats(proj, db.entries(project_only=True), db.members(), today)
    if stats.empty:
        st.info("Noch keine Projekte angelegt.")
        return
    k = st.columns(4)
    k[0].metric("Budget gesamt", f"{bt(stats['budget'].sum())} BT")
    k[1].metric("Ist", f"{bt(stats['gebucht'].sum())} BT",
                help="Gemeldete VAL-Tage; für Monate ohne Meldung die geplanten Tage bis heute")
    k[2].metric("Geplant ab morgen", f"{bt(stats['geplant'].sum())} BT")
    k[3].metric("Offen", f"{bt(stats['offen'].clip(lower=0).sum())} BT",
                f"{bt(-stats['offen'].clip(upper=0).sum())} BT über Budget", delta_color="inverse")
    s = stats.sort_values("offen")
    fig = go.Figure([
        go.Bar(y=s["label"], x=s["gebucht"], name="Ist", orientation="h", marker_color=WP["navy"]),
        go.Bar(y=s["label"], x=s["geplant"], name="Geplant", orientation="h", marker_color=WP["mid"]),
        go.Bar(y=s["label"], x=s["offen"].clip(lower=0), name="Offen", orientation="h", marker_color=WP["light"]),
        go.Bar(y=s["label"], x=-s["offen"].clip(upper=0), name="Über Budget", orientation="h",
               marker_color=WP["alert"])])
    fig.update_layout(barmode="stack", xaxis_title="Beratertage")
    st.plotly_chart(plotly_layout(fig, 80 + 38 * len(s)), width="stretch")

    sub("Staffing je Projekt und Woche (Beratertage, nächste 12 Wochen)")
    m0 = logic.monday(today)
    e = db.entries(m0, m0 + dt.timedelta(weeks=12) - dt.timedelta(days=1), project_only=True)
    if not e.empty:
        e["woche"] = [logic.monday(d) for d in e["tag"]]
        e["projekt"] = e["project_id"].astype(int).map(db.projects(active_only=False).set_index("id")["label"])
        e["bt"] = e["stunden"] / H
        pv = e.pivot_table(index="projekt", columns="woche", values="bt", aggfunc="sum").fillna(0)
        pv.columns = [logic.week_label(c) for c in pv.columns]
        mx = max(pv.values.max(), 1)
        st.dataframe(pv.style.format("{:.1f}").map(lambda v: util_css(v / mx * 1.1) if v else ""),
                     width="stretch")


# ====================================================================== Wiederverwendete Bausteine
def client_picker(prefix: str):
    """Kunde tippen: bestehende Kunden werden erkannt, neue mit Logo-Vorschau angelegt.
    Rückgabe: (client_id | None, neuer_name | None, domain | None)."""
    cl = db.clients()
    kunde = st.selectbox("Kunde", cl["name"].tolist(), index=None, accept_new_options=True,
                         placeholder="Kundenname tippen", key=f"{prefix}_kunde")
    if not kunde:
        return None, None, None
    exact = cl[cl["name"] == kunde]
    if not exact.empty:
        r = exact.iloc[0]
    else:
        r = db.match_client(kunde)
    if r is not None:
        c1, c2 = st.columns([1, 12])
        c1.image(logo(r["name"], r["domain"], r["logo"]), width=36)
        if r["name"] != kunde:
            c2.success(f"Erkannt als **{r['name']}**")
        if not r["aktiv"]:
            db.save_table("clients", ["aktiv"], pd.DataFrame([{"id": r["id"], "aktiv": 1}]), pd.DataFrame({"id": []}))
        return int(r["id"]), None, None
    st.warning(f"Neuer Kunde: **{kunde}**. Domain prüfen, das Logo wird daraus gezogen.")
    c1, c2 = st.columns([1, 8])
    dom = c2.text_input("Website / Domain", db.guess_domain(kunde), key=f"{prefix}_dom")
    c1.image(logo(kunde, dom, None), width=36)
    return None, kunde, dom
STOPS = [(0.0, (255, 255, 255)), (0.5, (228, 228, 246)), (0.8, (187, 204, 238)),
         (1.0, (102, 136, 187)), (1.2, (0, 0, 85))]
def styled_pct(df: pd.DataFrame, na_text: str = "abwesend"):
    """Prozent-Tabelle als Text (leere Werte = na_text), Farben aus den Zahlen."""
    disp = df.apply(lambda col: [fmt_pct(v) if v is not None and v == v else na_text for v in col])
    return disp.style.apply(lambda col: [util_css(df.at[i, col.name]) for i in col.index], axis=0)


def util_css(v) -> str:
    if v is None or pd.isna(v):
        return f"background-color:#F7F7F7;color:{WP['grey']}"
    v = max(0.0, min(float(v), 1.2))
    for (a, ca), (b, cb) in zip(STOPS, STOPS[1:]):
        if v <= b:
            t = (v - a) / (b - a) if b > a else 0
            rgb = tuple(int(ca[i] + (cb[i] - ca[i]) * t) for i in range(3))
            break
    txt = "#FFFFFF" if v >= 0.95 else WP["navy"]
    return f"background-color:rgb{rgb};color:{txt}"
def _aw_pick(key: str, order: list[str]) -> None:
    rows = st.session_state[key].selection.rows
    st.session_state.aw_sel = order[rows[0]] if rows else None


def _aw_step(order: list[str], step: int) -> None:
    cur = st.session_state.get("aw_sel")
    i = order.index(cur) if cur in order else -1
    st.session_state.aw_sel = order[(i + step) % len(order)] if i >= 0 or step > 0 else order[-1]
    st.session_state.aw_ver = st.session_state.get("aw_ver", 0) + 1


def _aw_all() -> None:
    st.session_state.aw_sel = None
    st.session_state.aw_ver = st.session_state.get("aw_ver", 0) + 1


def ziel_von(cons: pd.DataFrame) -> dict:
    """Zielauslastung je Mitarbeiter (Kürzel) aus seiner Position, sonst Standard aus den Secrets."""
    z = pd.to_numeric(cons["ziel_position"], errors="coerce") if "ziel_position" in cons else pd.Series(dtype=float)
    return {n: (float(v) if pd.notna(v) and v > 0 else ZIEL) for n, v in zip(cons["name"], z)} if len(z) \
        else {n: ZIEL for n in cons["name"]}


def team_ziel(cons: pd.DataFrame) -> float:
    z = list(ziel_von(cons).values())
    return sum(z) / len(z) if z else ZIEL


BER_FARBE = {db.BEREICHE[0]: WP["navy"], db.BEREICHE[1]: WP["mid"]}
BER_KURZ = {db.BEREICHE[0]: "OpEx", db.BEREICHE[1]: "ServiceEx"}


def bereiche_von(cons: pd.DataFrame) -> list[str]:
    """Vorhandene Bereiche in fester Reihenfolge (Operational Excellence zuerst)."""
    da = list(dict.fromkeys(cons["bereich"])) if "bereich" in cons and not cons.empty else []
    return [b for b in db.BEREICHE if b in da] + [b for b in da if b not in db.BEREICHE]


def ber_farbe(b: str) -> str:
    return BER_FARBE.get(b, WP["grey"])


def _kurzname(cons_row) -> str:
    v = _s(cons_row.vollname)
    return v if v else cons_row.name


def team_small_multiples(pers: pd.DataFrame, order: list[str], names: dict, periods: list[str],
                         avg: dict, ziele: dict) -> None:
    """Auslastung je Mitarbeiter über die Zeit: eine Kachel pro Person, gleiche Achsen, nach Position."""
    ncol = 3
    nrow = -(-len(order) // ncol)
    titles = [f"<b>{html.escape(names.get(n, n))}</b>  Ø {fmt_pct(avg[n])} · Ziel {ziele.get(n, ZIEL):.0%}"
              for n in order]
    fig = make_subplots(rows=nrow, cols=ncol, shared_xaxes=True, shared_yaxes=True, subplot_titles=titles,
                        vertical_spacing=0.32 / max(nrow, 1) + 0.03, horizontal_spacing=0.035)
    ymax = 1.2
    for i, nm in enumerate(order):
        d = pers[pers["name"] == nm].set_index("p_label").reindex(periods)
        y = d["auslastung"]
        if y.notna().any():
            ymax = max(ymax, float(y.max()) + 0.1)
        r, c = i // ncol + 1, i % ncol + 1
        fig.add_trace(go.Scatter(
            x=periods, y=y, mode="lines+markers", showlegend=False, connectgaps=False,
            line=dict(color=WP["navy"], width=2), fill="tozeroy", fillcolor="rgba(187,204,238,0.35)",
            marker=dict(size=7, color=[WP["alert"] if (v == v and v > 1.0001) else WP["navy"] for v in y],
                        line=dict(color="#FFFFFF", width=1.5)),
            hovertemplate=f"<b>{html.escape(names.get(nm, nm))}</b><br>%{{x}}: %{{y:.0%}}<extra></extra>"),
            row=r, col=c)
        fig.add_hline(y=1.0, line_color=WP["alert"], line_width=1, line_dash="dot", row=r, col=c)
        fig.add_hline(y=ziele.get(nm, ZIEL), line_color=WP["grey"], line_width=1, line_dash="dash", row=r, col=c)
    fig.update_yaxes(tickformat=".0%", range=[0, ymax], gridcolor=WP["xlgrey"], zeroline=False,
                     tickfont=dict(size=10), nticks=4)
    fig.update_xaxes(showgrid=False, linecolor=WP["lgrey"], tickfont=dict(size=10), tickangle=0, nticks=6)
    fig.update_annotations(font=dict(size=13, color=WP["navy"], family=FONT))
    for a in fig.layout.annotations:
        a.update(x=a.x - (1 / ncol) / 2 + 0.01, xanchor="left")
    fig.update_layout(height=60 + 205 * nrow, margin=dict(l=10, r=10, t=40, b=10), plot_bgcolor="#FFFFFF",
                      paper_bgcolor="#FFFFFF", font=dict(family=FONT, color=WP["navy"], size=12),
                      hoverlabel=dict(font_family=FONT))
    st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})
    st.caption("Eine Kachel je Mitarbeiter, sortiert nach Position, gleiche Skala. Gestrichelt grau = "
               "Zielauslastung der Position, gepunktet rot = 100 %, rote Punkte = überbucht. Zeile in der Tabelle unten anklicken "
               "für die Einzelansicht.")


def umsatz_auswertung(nur_team: bool = False) -> None:
    """Umsatz je Mitarbeiter und Monat = Projekttage × Tagessatz. Ist = gemeldete VAL, sonst Plan."""
    today = heute()
    cons = db.consultants()
    if pd.to_numeric(cons["tagessatz_eff"], errors="coerce").fillna(0).sum() <= 0:
        st.info("Noch keine Tagessätze hinterlegt. Im Reiter **Team** unter Positionen und Tagessätze eintragen.")
        return
    c1, c2 = st.columns([2, 3])
    bis_m = c1.slider("Monate voraus", 1, 12, 3, key="um_mon")
    start = dt.date(today.year, 1, 1)
    end = logic.month_end(today, bis_m - 1)
    ts = dict(zip(cons["id"], pd.to_numeric(cons["tagessatz_eff"], errors="coerce").fillna(0.0)))
    akt = logic.month_start(today)
    # Ist = gemeldete VAL; Monate ohne Meldung: Plan (Vergangenheit = "VAL offen")
    v = db.val(start, end)
    v = v[v["consultant_id"].isin(ts)] if not v.empty else v
    gemeldet = (set(zip(v["consultant_id"], v["monat"])) if not v.empty else set()) | logic.GEMELDET
    e = db.entries(start, end, project_only=True)
    e = e[e["art"].isin(logic.PROJEKT_ARTEN) & e["consultant_id"].isin(ts)] if not e.empty else e
    teile = []
    if not v.empty:
        teile.append(pd.DataFrame({"consultant_id": v["consultant_id"], "monat": v["monat"], "status": "Ist (VAL)",
                                   "umsatz": v["tage"] * v["consultant_id"].map(ts)}))
    if not e.empty:
        e = e.assign(monat=e["tag"].map(logic.month_start))
        e = e[[(c, m_) not in gemeldet for c, m_ in zip(e["consultant_id"], e["monat"])]]
        teile.append(pd.DataFrame({"consultant_id": e["consultant_id"], "monat": e["monat"],
                                   "status": ["VAL offen" if m_ < akt else "Plan" for m_ in e["monat"]],
                                   "umsatz": e["stunden"] / H * e["consultant_id"].map(ts)}))
    if not teile:
        st.info("Keine Projekttage im Zeitraum.")
        return
    e = pd.concat(teile, ignore_index=True)
    months = []
    m = start
    while m <= end:
        months.append(m)
        m = logic.month_start(m, 1)
    mkurz = [logic.MONATE[x.month - 1] + ("" if x.year == start.year else f" {x:%y}") for x in months]
    names = {r.id: _kurzname(r) for r in cons.itertuples()}
    farbe = {"Ist (VAL)": WP["navy"], "VAL offen": WP["light"], "Plan": WP["mid"]}
    stati = ("Ist (VAL)", "VAL offen", "Plan")

    # Team gesamt, bei mehreren Bereichen je Monat ein Balken je Bereich (Operational Excellence zuerst)
    bers = bereiche_von(cons)
    gruppiert = len(bers) > 1
    e = e.assign(bereich=e["consultant_id"].map(dict(zip(cons["id"], cons["bereich"]))) if gruppiert else "")
    idx = pd.MultiIndex.from_product([months, bers if gruppiert else [""]], names=["monat", "bereich"])
    g = e.groupby(["monat", "bereich", "status"])["umsatz"].sum().unstack(fill_value=0.0).reindex(idx, fill_value=0.0)
    if gruppiert:
        xs = [[logic.month_label(m_) for m_, _ in g.index], [BER_KURZ.get(b, b) for _, b in g.index]]
    else:
        xs = [logic.month_label(m_) for m_, _ in g.index]
    fig = go.Figure()
    for stt in stati:
        if stt in g:
            fig.add_trace(go.Bar(x=xs, y=g[stt], name=stt, marker_color=farbe[stt],
                                 marker_pattern_shape="/" if stt == "VAL offen" else "",
                                 hovertemplate=f"{stt} %{{x}}: %{{y:,.0f}} €<extra></extra>"))
    tot = g.sum(axis=1)
    fig.add_trace(go.Scatter(x=xs, y=tot, mode="text", text=[f"{v / 1000:,.0f} T€" if v else "" for v in tot],
                             textposition="top center", showlegend=False, hoverinfo="skip",
                             textfont=dict(color=WP["navy"], size=11 if gruppiert else 12)))
    if gruppiert:
        teil = " · ".join(f"{BER_KURZ.get(b, b)} {tot.xs(b, level='bereich').sum() / 1000:,.0f} T€" for b in bers)
        titel = f"Umsatz je Monat und Bereich · {teil}".replace(",", ".")
    else:
        titel = f"Umsatz Team je Monat · Σ {tot.sum() / 1000:,.0f} T€".replace(",", ".")
    fig.update_layout(barmode="stack", bargap=0.25 if gruppiert else 0.35, title=dict(text=titel, x=0, font=dict(size=15)))
    fig.update_yaxes(tickformat=",.0f", ticksuffix=" €", rangemode="tozero")
    fig = plotly_layout(fig, 340)
    fig.update_layout(separators=",.", legend=dict(x=1, xanchor="right", y=1.08, traceorder="normal"))
    st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})

    if nur_team:
        st.caption("Umsatz = Projekttage × Tagessatz der Position. **Ist (VAL)** = gemeldete VAL, **VAL offen** = "
                   "vergangener Monat ohne Meldung (mit Plan bewertet), **Plan** = laufender und künftige Monate. "
                   "Interne Tage, Akquise und Abwesenheit erzeugen keinen Umsatz.")
        return
    # Je Mitarbeiter (Kacheln)
    per = e.groupby(["consultant_id", "monat", "status"])["umsatz"].sum()
    summe = e.groupby("consultant_id")["umsatz"].sum().reindex(list(names), fill_value=0.0)
    order = [c for c in cons["id"] if c in summe.index]  # Reihenfolge nach Position
    ncol = 3
    nrow = -(-len(order) // ncol)
    titles = [f"<b>{html.escape(names[c])}</b>  Σ {summe[c] / 1000:,.0f} T€".replace(",", ".") for c in order]
    sf = make_subplots(rows=nrow, cols=ncol, shared_xaxes=True, shared_yaxes=True, subplot_titles=titles,
                       vertical_spacing=0.32 / max(nrow, 1) + 0.03, horizontal_spacing=0.035)
    for i, cid in enumerate(order):
        r, c = i // ncol + 1, i % ncol + 1
        for stt in stati:
            y = [float(per.get((cid, mo, stt), 0.0)) for mo in months]
            sf.add_trace(go.Bar(x=mkurz, y=y, name=stt, marker_color=farbe[stt], showlegend=(i == 0),
                                marker_pattern_shape="/" if stt == "VAL offen" else "",
                                legendgroup=stt, hovertemplate=f"<b>{html.escape(names[cid])}</b><br>{stt} "
                                                               "%{x}: %{y:,.0f} €<extra></extra>"), row=r, col=c)
    sf.update_layout(barmode="stack", bargap=0.3, height=80 + 200 * nrow, margin=dict(l=10, r=10, t=60, b=10),
                     plot_bgcolor="#FFFFFF", paper_bgcolor="#FFFFFF",
                     font=dict(family=FONT, color=WP["navy"], size=12),
                     legend=dict(orientation="h", yanchor="bottom", y=1.03, x=1, xanchor="right",
                                 traceorder="normal"), separators=",.",
                     title=dict(text="Umsatz je Mitarbeiter und Monat", x=0, font=dict(size=15)))
    sf.update_yaxes(tickformat="~s", ticksuffix=" €", gridcolor=WP["xlgrey"], zeroline=False,
                    tickfont=dict(size=10), nticks=4)
    sf.update_xaxes(showgrid=False, linecolor=WP["lgrey"], tickfont=dict(size=10), tickangle=0)
    sf.update_annotations(font=dict(size=13, color=WP["navy"], family=FONT))
    for a in sf.layout.annotations:
        a.update(x=a.x - (1 / ncol) / 2 + 0.01, xanchor="left")
    st.plotly_chart(sf, width="stretch", config={"displayModeBar": False})
    st.caption("Umsatz = Projekttage × Tagessatz. **Ist (VAL)** = gemeldete VAL-Tage (Reiter VAL buchen). "
               "**VAL offen** = vergangener Monat ohne VAL-Meldung, vorläufig mit Plan bewertet. **Plan** = laufender "
               "und künftige Monate. Interne Tage, Urlaub usw. erzeugen keinen Umsatz. Nur für Admins sichtbar.")


def auswertung_auslastung() -> None:
    """Meeting-Modus: das Team in Reihenfolge der Position durchgehen."""
    today = heute()
    ss = st.session_state
    c1, c2 = st.columns([3, 2])
    gran = c1.radio("Granularität", list(logic.GRANULARITAETEN), format_func=logic.GRANULARITAETEN.get,
                    horizontal=True, key="aw_gran")
    monate = c2.slider("Horizont (Monate)", 1, 6, 3, key="aw_mon")
    start = logic.monday(today) if gran != "Monat" else logic.month_start(today)
    end = logic.month_end(today, monate - 1)
    cons = db.consultants()
    fact = logic.build_fact(start, end, cons, db.entries(start, end))
    if fact.empty:
        st.info("Keine Daten im Zeitraum.")
        return
    team = logic.aggregate(fact, gran, today, by_consultant=False)
    pers = logic.aggregate(fact, gran, today, by_consultant=True)
    periods = team["p_label"].tolist()
    plab = person_label(cons)
    avg = {n: logic.kpi(fact[fact["name"] == n]) for n in cons["name"]}
    order = list(cons["name"])  # Reihenfolge nach Position (Hierarchie)
    sel = ss.get("aw_sel") if ss.get("aw_sel") in order else None

    # ---------------- Navigation durch das Team
    n = st.columns([1.3, 1.3, 1.3, 5], vertical_alignment="center")
    n[0].button("◀ Vorherige/r", on_click=_aw_step, args=(order, -1), width="stretch")
    n[1].button("Nächste/r ▶", on_click=_aw_step, args=(order, 1), type="primary", width="stretch")
    n[2].button("Ganzes Team", on_click=_aw_all, disabled=sel is None, width="stretch")
    if sel:
        n[3].markdown(f"**{html.escape(plab.get(sel, sel))}** · Platz {order.index(sel) + 1} von {len(order)} · "
                      f"Ø {fmt_pct(avg[sel])}")
    else:
        n[3].markdown("**Team gesamt** · Zeile in der Tabelle anklicken oder mit **Nächste/r** der Hierarchie "
                      "nach durchgehen")

    # ---------------- Liniendiagramm: ganzes Team als Kacheln, sonst die ausgewählte Person
    if not sel:
        team_small_multiples(pers, order, {r.name: _kurzname(r) for r in cons.itertuples()}, periods, avg,
                             ziel_von(cons))
    fig = go.Figure()
    if sel:
        d = pers[pers["name"] == sel]
        y, name = d["auslastung"], plab.get(sel, sel)
        x = d["p_label"]
    else:
        y, name, x = team["auslastung"], "Team", team["p_label"]
    fig.add_trace(go.Scatter(x=x, y=y, name=name, mode="lines+markers+text",
                             text=[fmt_pct(v) for v in y], textposition="top center",
                             line=dict(color=WP["navy"], width=4), marker=dict(size=9),
                             textfont=dict(color=WP["navy"], size=12)))
    fig.add_hline(y=1.0, line_color=WP["alert"], line_width=1, line_dash="dot",
                  annotation_text="100 %", annotation_font_color=WP["alert"])
    z_ = ziel_von(cons).get(sel, ZIEL) if sel else team_ziel(cons)
    fig.add_hline(y=z_, line_dash="dash", line_color=WP["grey"],
                  annotation_text=f"Ziel {z_:.0%}", annotation_font_color=WP["grey"])
    fig.update_xaxes(categoryorder="array", categoryarray=periods)
    fig.update_yaxes(rangemode="tozero")
    fig.update_layout(showlegend=False, title=dict(text=f"Auslastung {name}", x=0, font=dict(size=15)))
    if sel:
        st.plotly_chart(plotly_layout(fig, 360, pct=True), width="stretch")

    # ---------------- Matrix (anklickbar)
    sub("Auslastung je Berater, nach Position (Zeile anklicken)")
    pv = pers.pivot(index="name", columns="p_label", values="auslastung").reindex(index=order, columns=periods)
    pv.insert(0, "Ø", [avg[x] for x in order])
    pv.index = [plab.get(x, x) for x in order]
    pv.index.name = "Berater"
    sel_lab = plab.get(sel, sel) if sel else None
    sty = styled_pct(pv).apply(
        lambda r: ["font-weight:800; text-decoration:underline" if r.name == sel_lab else ""] * len(r), axis=1)
    key = f"aw_df_{ss.get('aw_ver', 0)}"
    st.dataframe(sty, width="stretch", height=38 + 35 * len(pv), key=key, on_select=lambda: _aw_pick(key, order),
                 selection_mode="single-row")

    # ---------------- Detail des ausgewählten Beraters
    if not sel:
        return
    cid = int(cons.loc[cons["name"] == sel, "id"].iloc[0])
    sub(f"{plab.get(sel, sel)}: Was steht an?")
    wk = st.segmented_control("Woche", ["Diese Woche", "Nächste Woche", "Übernächste Woche"],
                              default="Diese Woche", key="aw_wk", label_visibility="collapsed")
    off = {"Nächste Woche": 1, "Übernächste Woche": 2}.get(wk or "", 0)
    ws = logic.monday(today) + dt.timedelta(weeks=off)
    st.caption(f"KW {ws.isocalendar().week:02d} · {ws:%d.%m.} bis {ws + dt.timedelta(days=4):%d.%m.%Y}")
    week_readonly(cid, ws)
    sub(f"{plab.get(sel, sel)}: Projekte und offene Beratertage")
    member_projects_table(cid)


def kunden_logos() -> None:
    cl = db.clients(active_only=False)
    cl.insert(1, "vorschau", [logo(r.name, r.domain, r.logo) for r in cl.itertuples()])
    cl["aktiv"] = cl["aktiv"].astype(bool)
    ed = st.data_editor(cl.drop(columns=["logo"]), hide_index=True, width="stretch",
                        num_rows="dynamic", key="ed_clients", column_config={
                            "id": None,
                            "vorschau": st.column_config.ImageColumn("Logo", width="small"),
                            "name": st.column_config.TextColumn("Kunde", required=True),
                            "domain": st.column_config.TextColumn("Domain (für Logo)"),
                            "aliases": st.column_config.TextColumn("Aliasse (Komma getrennt)"),
                            "aktiv": st.column_config.CheckboxColumn("Aktiv")},
                        disabled=["vorschau"])
    if st.button("Kunden speichern", type="primary"):
        db.save_table("clients", ["name", "domain", "aliases", "aktiv"],
                      ed.assign(aktiv=ed["aktiv"].fillna(True).astype(int)), cl)
        st.session_state.pop("ed_clients", None)
        st.rerun()

    sub("Eigenes Logo hochladen (überschreibt automatisches Logo)")
    c1, c2 = st.columns(2)
    wahl = c1.selectbox("Kunde", cl["name"].tolist(), index=None, key="logo_kunde")
    up = c2.file_uploader("PNG, JPG oder SVG", type=["png", "jpg", "jpeg", "svg"])
    b1, b2 = st.columns([1, 5])
    if b1.button("Logo setzen", disabled=not (wahl and up)):
        mime = "image/svg+xml" if up.name.lower().endswith(".svg") else up.type
        db.set_logo(int(cl[cl["name"] == wahl]["id"].iloc[0]),
                    f"data:{mime};base64,{base64.b64encode(up.getvalue()).decode()}")
        st.rerun()
    if b2.button("Automatisches Logo wiederherstellen", disabled=not wahl):
        db.set_logo(int(cl[cl["name"] == wahl]["id"].iloc[0]), None)
        st.rerun()


def anlegen_mitarbeiter() -> None:
    cons_all = db.consultants(active_only=False)
    if not is_admin():
        admins = ", ".join(sorted(ADMINS)) or "Admin"
        st.info(f"Neue Mitarbeiter legt der Admin an ({admins}).")
        st.dataframe(db.consultants()[["name", "vollname", "rolle"]].map(_s), hide_index=True, width="stretch",
                     column_config={"name": "Kürzel", "vollname": "Name", "rolle": "Rolle"})
        return
    sub("Neuen Mitarbeiter anlegen")
    c1, c2, c3 = st.columns([1, 2, 2])
    kz = c1.text_input("Kürzel (Login)", key="nm_kz", max_chars=12, placeholder="z. B. abc")
    vn = c2.text_input("Vollständiger Name", key="nm_vn")
    ro = c3.selectbox("Rolle", ROLLEN, index=0, accept_new_options=True, key="nm_ro")
    exists = kz.strip().lower() in {str(x).lower() for x in cons_all["name"]}
    if exists:
        st.caption("Kürzel existiert bereits. Anlegen reaktiviert bzw. aktualisiert den Eintrag.")
    if st.button("Mitarbeiter anlegen", type="primary", disabled=not kz.strip(), key="nm_go"):
        _, neu = db.create_consultant(kz, vn.strip(), ro or "")
        for k in ("nm_kz", "nm_vn"):
            st.session_state.pop(k, None)
        st.session_state["nm_done"] = (f"**{kz.strip()}** {'angelegt' if neu else 'reaktiviert'}. "
                                       "Login ab sofort mit dem Team-Passwort möglich.")
        st.rerun()
    if st.session_state.get("nm_done"):
        st.success(st.session_state.pop("nm_done"))

    sub("Team bearbeiten")
    team_editor()

    with st.expander("Admin: Planungsdaten zurücksetzen (z. B. Demo-Daten vor dem Echtstart löschen)"):
        st.warning("Löscht **alle Buchungen, Projekte und Kunden** unwiderruflich. Das Team bleibt erhalten.")
        ok = st.text_input("Zur Bestätigung LÖSCHEN eintippen", key="reset_ok")
        if st.button("Alles zurücksetzen", disabled=ok != "LÖSCHEN", key="reset_go"):
            db.reset_planning()
            st.session_state.pop("reset_ok", None)
            st.success("Planungsdaten gelöscht.")
def team_editor() -> None:
    c = db.consultants(active_only=False)
    c["aktiv"] = c["aktiv"].astype(bool)
    c = c[["id", "name", "vollname", "rolle", "sort", "aktiv"]]
    for col in ("vollname", "rolle"):
        c[col] = c[col].map(_s)
    ed = st.data_editor(c, hide_index=True, width="stretch", num_rows="dynamic", key="ed_team",
                        column_config={"id": None,
                                       "name": st.column_config.TextColumn("Kürzel (Login)", required=True),
                                       "vollname": "Name",
                                       "rolle": st.column_config.SelectboxColumn("Rolle", options=ROLLEN),
                                       "sort": st.column_config.NumberColumn("Reihenfolge", step=1),
                                       "aktiv": st.column_config.CheckboxColumn("Aktiv")})
    if st.button("Team speichern", type="primary", key="team_save"):
        db.save_table("consultants", ["name", "vollname", "rolle", "sort", "aktiv"],
                      ed.assign(aktiv=ed["aktiv"].fillna(True).astype(int)), c)
        st.session_state.pop("ed_team", None)
        st.rerun()


def _eur(v) -> str:
    return f"{float(v):,.0f} €".replace(",", ".")


def _pos_liste() -> list[str]:
    return db.positionen()["name"].tolist()


@st.dialog("Mitarbeiter", width="medium")
def ma_dialog(cid: int | None) -> None:
    """Anlegen (cid None) oder bearbeiten. Tagessatz kommt aus der Position, optional individuell."""
    ss = st.session_state
    pos = db.positionen()
    satz = dict(zip(pos["name"], pos["tagessatz"]))
    zmap = dict(zip(pos["name"], pos["zielauslastung"]))
    plist = pos["name"].tolist()
    r = None
    if cid is not None:
        cons = db.consultants(active_only=False)
        r = cons[cons["id"] == cid].iloc[0]
    k = f"{cid or 'neu'}"
    vn = st.text_input("Vor- und Nachname", _s(r["vollname"]) if r is not None else "", key=f"ma_vn_{k}")
    vorschlag = db.suggest_kuerzel(vn) if (r is None and vn.strip()) else ""
    kz = st.text_input("Kürzel (Login)", r["name"] if r is not None else vorschlag, key=f"ma_kz_{k}_{vorschlag}",
                       max_chars=12, help="Mit diesem Kürzel meldet sich die Person an.")
    rolle_alt = _s(r["rolle"]) if r is not None else "Consultant"
    if rolle_alt and rolle_alt not in plist:
        plist = plist + [rolle_alt]
    rolle = st.selectbox("Position", plist, index=plist.index(rolle_alt) if rolle_alt in plist else 0,
                         key=f"ma_ro_{k}",
                         format_func=lambda p: p + (f"  ·  {_eur(satz[p])}" if pd.notna(satz.get(p)) else "")
                         + (f"  ·  Ziel {zmap[p]:.0%}" if pd.notna(zmap.get(p)) else ""))
    ber_alt = _s(r.get("bereich")) if r is not None else ""
    ber_alt = ber_alt or db.BEREICHE[0]
    blist = db.BEREICHE + ([ber_alt] if ber_alt not in db.BEREICHE else [])
    bereich = st.segmented_control("Bereich", blist, default=ber_alt, key=f"ma_be_{k}") or ber_alt
    ind_alt = float(r["tagessatz"]) if r is not None and pd.notna(r.get("tagessatz")) and r["tagessatz"] else None
    abw = st.toggle("Abweichender Tagessatz", value=ind_alt is not None, key=f"ma_abw_{k}")
    ind = None
    if abw:
        ind = st.number_input("Individueller Tagessatz (€)", min_value=0, step=50,
                              value=int(ind_alt or satz.get(rolle) or 0), key=f"ma_ts_{k}")
    else:
        s_ = satz.get(rolle)
        st.caption(f"Tagessatz laut Position: **{_eur(s_)}**" if pd.notna(s_) else
                   "Für diese Position ist noch kein Tagessatz hinterlegt (unten unter Positionen).")
    st.write("")
    b1, b2 = st.columns(2)
    if b1.button("Speichern", type="primary", width="stretch", disabled=not kz.strip(), key=f"ma_save_{k}"):
        alle = db.consultants(active_only=False)
        dup = alle[(alle["name"].str.lower() == kz.strip().lower()) & (alle["id"] != (cid or -1))]
        if not dup.empty:
            st.error(f"Kürzel **{kz.strip()}** ist schon vergeben.")
            return
        if cid is None:
            neu_id, _ = db.create_consultant(kz, vn.strip(), rolle, None)
            db.update_consultant(neu_id, kz, vn, rolle, ind, True, bereich)
            hinweis = "" if TEAM_PW else (f" Login: in den Streamlit-Secrets unter [users] die Zeile "
                                          f"`{kz.strip()} = \"Passwort\"` ergänzen.")
            ss["team_msg"] = f"✓ {vn.strip() or kz.strip()} angelegt.{hinweis}"
        else:
            db.update_consultant(cid, kz, vn, rolle, ind, True, bereich)
            ss["team_msg"] = f"✓ {vn.strip() or kz.strip()} gespeichert."
        st.rerun()
    if b2.button("Abbrechen", width="stretch", key=f"ma_cancel_{k}"):
        st.rerun()
    if cid is not None:
        st.divider()
        if st.button("Ausgeschieden: Mitarbeiter deaktivieren", type="tertiary", key=f"ma_deact_{k}",
                     help="Verschwindet aus Auswahl und Auswertungen, die Historie bleibt erhalten."):
            db.update_consultant(cid, r["name"], _s(r["vollname"]), _s(r["rolle"]), ind_alt, False)
            ss["team_msg"] = f"{_s(r['vollname']) or r['name']} deaktiviert."
            st.rerun()


def page_team() -> None:
    """Team pflegen: einfache Karten, Bearbeiten per Klick, Tagessätze je Position."""
    header("Team", "Mitarbeiter, Positionen und Tagessätze")
    ss = st.session_state
    if ss.get("team_msg"):
        st.success(ss.pop("team_msg"))
    cons = db.consultants()
    sub(f"Aktives Team ({len(cons)})")
    open_id = "x"
    gruppen = [(b, list(cons[cons["bereich"] == b].itertuples())) for b in bereiche_von(cons)] or [("", [])]
    gruppen[-1][1].append("neu")
    for b, karten in gruppen:
        if b:
            st.markdown(f"<div style='margin:.6rem 0 .3rem;font-weight:700;color:{ber_farbe(b)}'>{html.escape(b)} "
                        f"<span style='color:{WP['grey']};font-weight:400'>· "
                        f"{sum(1 for x in karten if x != 'neu')}</span></div>", unsafe_allow_html=True)
        for i in range(0, len(karten), 3):
            cols = st.columns(3)
            for col, r in zip(cols, karten[i:i + 3]):
                with col.container(border=True):
                    if r == "neu":
                        st.markdown("<div class='pc-add' style='min-height:74px'><b>＋ Neuer Mitarbeiter</b>"
                                    "<span>Name, Position, fertig.</span></div>", unsafe_allow_html=True)
                        if st.button("Anlegen", key="ma_new", width="stretch", type="primary"):
                            open_id = None
                        continue
                    ind = pd.notna(r.tagessatz) and r.tagessatz
                    satz = (f"{_eur(r.tagessatz_eff)} <span style='color:{WP['grey']}'>"
                            f"({'individuell' if ind else 'Position'})</span>" if pd.notna(r.tagessatz_eff)
                            else f"<span style='color:{WP['grey']}'>kein Tagessatz</span>")
                    zl = r.ziel_position if pd.notna(r.ziel_position) and r.ziel_position else ZIEL
                    st.markdown(f"<div class='pc-h'><div><b>{html.escape(_s(r.vollname) or r.name)}</b>"
                                f"<span>{html.escape(r.name)} · {html.escape(_s(r.rolle) or 'ohne Position')}</span>"
                                f"</div></div><div class='pc-f' style='font-size:.85rem'>Tagessatz {satz} · "
                                f"Ziel {zl:.0%}</div>", unsafe_allow_html=True)
                    if st.button("✏️ Bearbeiten", key=f"ma_ed_{r.id}", width="stretch"):
                        open_id = int(r.id)
    if open_id != "x":
        ma_dialog(open_id)

    # ---------------- Positionen & Tagessätze
    sub("Positionen: Tagessatz und Zielauslastung")
    st.caption("Tagessatz und Zielauslastung hängen an der Position. Wer eine abweichende Vereinbarung beim "
               "Tagessatz hat, bekommt beim Mitarbeiter einen individuellen Satz. Die Zielauslastung erscheint als "
               f"gestrichelte Linie in den Auslastungsdiagrammen (ohne Eintrag: Standard {ZIEL:.0%}).")
    pos = db.positionen()
    n_per = {p: int((cons["rolle"] == p).sum()) for p in pos["name"]}
    with st.form("pos_form", border=True):
        h = st.columns([2.2, 1.4, 1.4])
        h[0].markdown("**Position**")
        h[1].markdown("**Tagessatz (€)**")
        h[2].markdown("**Zielauslastung (%)**")
        werte = {}
        for p in pos.itertuples():
            c = st.columns([2.2, 1.4, 1.4], vertical_alignment="center")
            c[0].markdown(f"{html.escape(p.name)} <span style='color:{WP['grey']}'>· {n_per.get(p.name, 0)} "
                          f"{'Person' if n_per.get(p.name, 0) == 1 else 'Personen'}</span>", unsafe_allow_html=True)
            ts = c[1].number_input("Tagessatz", min_value=0, step=50, label_visibility="collapsed",
                                   value=int(p.tagessatz) if pd.notna(p.tagessatz) else 0, key=f"pos_{p.name}")
            zl = c[2].number_input("Ziel", min_value=0, max_value=150, step=5, label_visibility="collapsed",
                                   value=int(round(p.zielauslastung * 100)) if pd.notna(p.zielauslastung)
                                   else int(round(ZIEL * 100)), key=f"posz_{p.name}")
            werte[p.name] = (ts, zl)
        st.write("")
        c = st.columns([2.2, 1.4, 1.4])
        neu_n = c[0].text_input("Neue Position (optional)", key="pos_new_n", placeholder="z. B. Principal")
        neu_s = c[1].number_input("Tagessatz", min_value=0, step=50, value=0, key="pos_new_s")
        neu_z = c[2].number_input("Zielauslastung", min_value=0, max_value=150, step=5,
                                  value=int(round(ZIEL * 100)), key="pos_new_z")
        if st.form_submit_button("Speichern", type="primary"):
            for p in pos.itertuples():
                ts, zl = werte[p.name]
                db.save_position(p.name, float(ts) or None, int(p.sort), (zl / 100) or None)
            if neu_n.strip():
                db.save_position(neu_n, float(neu_s) or None, None, (neu_z / 100) or None)
            ss["team_msg"] = "✓ Positionen gespeichert."
            st.rerun()

    # ---------------- Ehemalige
    alle = db.consultants(active_only=False)
    weg = alle[alle["aktiv"].astype(int) == 0]
    if not weg.empty:
        with st.expander(f"Ehemalige Mitarbeiter ({len(weg)})"):
            for r in weg.itertuples():
                c1, c2 = st.columns([4, 1], vertical_alignment="center")
                c1.markdown(f"{html.escape(_s(r.vollname) or r.name)} · {html.escape(r.name)} · "
                            f"{html.escape(_s(r.rolle))}")
                if c2.button("Reaktivieren", key=f"ma_re_{r.id}", width="stretch"):
                    db.update_consultant(int(r.id), r.name, _s(r.vollname), _s(r.rolle),
                                         r.tagessatz if pd.notna(r.tagessatz) else None, True)
                    st.rerun()


MONATE_LANG = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober",
               "November", "Dezember"]


def _monat_lang(m: dt.date) -> str:
    return f"{MONATE_LANG[m.month - 1]} {m.year}"


def _val_keep() -> None:
    ss = st.session_state
    if not ss.get("val_wer"):
        ss.val_wer = ss.get("val_last")


def page_val() -> None:
    """VAL (verrechenbare Arbeitsleistung) je Projekt und Monat für ein ganzes Jahr, wie im CRM."""
    today = heute()
    ss = st.session_state
    header("VAL buchen", "Gebuchte verrechenbare Tage je Projekt und Monat")
    cons = db.consultants()
    if cons.empty:
        st.info("Noch keine Mitarbeiter angelegt.")
        return
    names = cons["name"].tolist()
    lab = {r.name: _vorname_label(r) for r in cons.itertuples()}
    c1, c2 = st.columns([5, 1], vertical_alignment="bottom")
    if is_admin() or str(ss.get("user", "")).lower() in VAL_ALLE:
        if ss.get("val_wer") not in names:
            ss.val_wer = ss.user if ss.user in names else names[0]
        with c1:
            st.pills("Mitarbeiter", names, format_func=lambda n: lab.get(n, n), key="val_wer",
                     selection_mode="single", on_change=_val_keep)
    else:
        if ss.user not in names:
            st.warning(f"Dein Kürzel **{ss.user}** ist im Team nicht angelegt. Bitte den Admin ansprechen.")
            return
        ss.val_wer = ss.user
        c1.markdown(f"**{html.escape(lab.get(ss.user, ss.user))}**")
    jahr = c2.selectbox("Jahr", [today.year - 1, today.year, today.year + 1], index=1, key="val_jahr")
    wer = ss.val_wer
    ss.val_last = wer
    cid = int(cons.loc[cons["name"] == wer, "id"].iloc[0])
    monate = [dt.date(jahr, m, 1) for m in range(1, 13)]
    mcol = {m: logic.MONATE[m.month - 1] for m in monate}

    projn = db.projects(active_only=False).set_index("id")
    vdf = db.val(monate[0], monate[-1], consultant_id=cid)
    ent = db.entries(monate[0], logic.month_end(monate[-1]), consultant_id=cid, project_only=True)
    mem = db.members()
    mine = set(mem.loc[mem["consultant_id"] == cid, "project_id"].astype(int))
    im_jahr = {p for p in mine if p in projn.index
               and (pd.isna(projn.at[p, "start"]) or projn.at[p, "start"] <= monate[-1].replace(month=12, day=31))
               and (pd.isna(projn.at[p, "ende"]) or projn.at[p, "ende"] >= monate[0])}
    extra = set(ss.get(f"val_extra_{cid}_{jahr}", []))
    pids = (im_jahr | set(vdf["project_id"].astype(int)) | set(ent["project_id"].dropna().astype(int)) | extra) \
        if not ent.empty or not vdf.empty else (im_jahr | extra)
    pids = sorted([int(p) for p in pids if p in projn.index], key=lambda p: projn.at[p, "label"].lower())
    werte = {(int(r.project_id), r.monat): float(r.tage) for r in vdf.itertuples()} if not vdf.empty else {}
    sdf = db.val_sonst(monate[0], monate[-1], consultant_id=cid)
    werte.update({(r.art, r.monat): float(r.tage) for r in sdf.itertuples()} if not sdf.empty else {})
    gemeldet = (set(vdf["monat"]) if not vdf.empty else set()) | (set(sdf["monat"]) if not sdf.empty else set())
    SONST_ICON = {"Allgemeine Tätigkeit": "🏛️", "Akquise": "🎯"}

    rows = []
    for p in list(pids) + db.VAL_SONST:
        r = {"pid": str(p), "Projekt": projn.at[p, "label"] if isinstance(p, int) else f"{SONST_ICON.get(p, '')} {p}"}
        for m in monate:
            v = werte.get((p, m))
            r[mcol[m]] = bt(v) if v is not None and m in gemeldet else LEER
        rows.append(r)
    original = pd.DataFrame(rows, columns=["pid", "Projekt"] + list(mcol.values()))
    cfg = {"pid": None, "Projekt": st.column_config.TextColumn("Projekt", disabled=True, width="medium", pinned=True)}
    for m, c in mcol.items():
        cfg[c] = st.column_config.TextColumn(c, width=54, help=f"VAL-Tage {MONATE_LANG[m.month - 1]} {jahr}")
    st.caption("Je Projekt die gebuchten Tage pro Monat eintragen, z. B. 9 oder 9,5. Listen aus dem CRM lassen sich "
               "direkt einfügen: erste Zelle anklicken, **Strg+V**. Ein Monat gilt als gemeldet, sobald dort "
               "mindestens ein Wert steht (0 ist erlaubt). Leere Monate rechnen weiter mit der Planung. "
               "Allgemeine Tätigkeit und Akquise erzeugen keinen Umsatz und zählen nicht als Auslastung.")
    key = f"val_ed2_{cid}_{jahr}_{len(pids)}"
    left, right = st.columns([12, 1.1], gap="small")
    with left:
        ed = st.data_editor(original, hide_index=True, width="stretch", num_rows="fixed", key=key,
                            height=38 + 35 * max(len(original), 1), column_config=cfg)
    summen = [sum(_num(v) for v in ed.loc[i, list(mcol.values())]) for i in ed.index] if not ed.empty else []
    with right:
        st.dataframe(pd.DataFrame({"Σ Jahr": [bt(x) for x in summen]}), hide_index=True, width="stretch",
                     height=38 + 35 * max(len(original), 1),
                     column_config={"Σ Jahr": st.column_config.TextColumn("Σ Jahr")})
    if not ed.empty:
        msum = {c: sum(_num(v) for v in ed[c]) for c in mcol.values()}
        tot = pd.DataFrame([{"Projekt": "Σ Monat", **{c: bt(v) if v else "" for c, v in msum.items()},
                             "Σ Jahr": bt(sum(msum.values()))}])
        st.dataframe(tot, hide_index=True, width="stretch",
                     column_config={"Projekt": st.column_config.TextColumn("", width="medium", pinned=True),
                                    **{c: st.column_config.TextColumn(c, width=54) for c in mcol.values()}})
    b1, b2, _ = st.columns([1.3, 1.8, 3])
    if b1.button("Speichern", type="primary", width="stretch", key="val_save", disabled=ed.empty):
        n_m = 0
        for m, c in mcol.items():
            belegt = [i for i in ed.index if str(ed.at[i, c] or "").strip()]
            proj = [(int(ed.at[i, "pid"]), max(_num(ed.at[i, c]), 0.0)) for i in belegt if ed.at[i, "pid"].isdigit()]
            sonst = [(ed.at[i, "pid"], max(_num(ed.at[i, c]), 0.0)) for i in belegt if not ed.at[i, "pid"].isdigit()]
            if belegt or m in gemeldet:
                db.save_val(cid, m, proj, ss.user)
                db.save_val_sonst(cid, m, sonst, ss.user)
                n_m += 1 if belegt else 0
        ss.pop(key, None)
        ss["val_msg"] = f"✓ Gespeichert: {n_m} Monat{'e' if n_m != 1 else ''} mit VAL für {lab.get(wer, wer)}."
        st.rerun()
    with b2.popover("＋ Projekt hinzufügen", width="stretch"):
        art = st.segmented_control("Projekt", ["Bestehendes", "Neues"], default="Bestehendes", key=f"val_art_{cid}",
                                   label_visibility="collapsed")
        if art == "Neues":
            kl = db.clients()["name"].tolist()
            kunde = st.selectbox("Kunde", kl, index=None, accept_new_options=True, key=f"val_nk_{cid}",
                                 placeholder="Kunde wählen oder neu eintippen")
            pname = st.text_input("Projektname (optional)", key=f"val_np_{cid}",
                                  placeholder="leer = Projekt heißt wie der Kunde")
            if st.button("Anlegen und hinzufügen", disabled=not kunde, key=f"val_new_go_{cid}", type="primary"):
                hit = db.clients(active_only=False)
                hit = hit[hit["name"].str.lower() == str(kunde).strip().lower()]
                kid = int(hit["id"].iloc[0]) if not hit.empty else db.create_client(str(kunde).strip(), "")
                pid_neu = db.ensure_project(kid, pname.strip() or str(kunde).strip())
                db.add_member(pid_neu, cid)
                ss[f"val_extra_{cid}_{jahr}"] = list(extra) + [pid_neu]
                ss["val_msg"] = f"✓ Projekt angelegt. Logo und Budget bei Bedarf unter **Projekt anlegen** ergänzen."
                st.rerun()
        else:
            rest = db.projects()
            rest = rest[~rest["id"].isin(pids)]
            wahl = st.selectbox("Projekt", rest["label"].tolist(), index=None, key=f"val_add_{cid}",
                                placeholder="Projekt suchen")
            if st.button("Hinzufügen", disabled=not wahl, key=f"val_add_go_{cid}", type="primary"):
                ss[f"val_extra_{cid}_{jahr}"] = list(extra) + [int(rest.loc[rest["label"] == wahl, "id"].iloc[0])]
                st.rerun()
    if ss.get("val_msg"):
        st.toast(ss.pop("val_msg"))


def _guete_css(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return f"color:{WP['grey']}"
    if v >= 0.9:
        return f"background-color:{WP['lilac']};color:{WP['navy']};font-weight:600"
    if v >= 0.75:
        return f"color:{WP['navy']}"
    return f"color:{WP['alert']};font-weight:700"


def planungsguete_view() -> None:
    """Gesamtübersicht: Wie gut war die Planung? VAL ÷ Plan je Monat, Team und Mitarbeiter."""
    today = heute()
    cons = db.consultants()
    c1, _ = st.columns([2, 3])
    n_m = c1.slider("Monate zurück", 3, 12, 6, key="pg_n")
    von = logic.month_start(today, -n_m)
    bis = logic.month_end(today)
    pv = logic.plan_val(db.entries(von, bis, project_only=True), db.val(von, bis))
    gm = logic.guete_matrix(pv)
    vdf_all = db.val()
    erster = min(vdf_all["monat"]) if not vdf_all.empty else logic.month_start(today, -1)
    monate = [m for m in (logic.month_start(today, -i) for i in range(n_m, -1, -1)) if m >= erster]
    if gm.empty:
        st.info("Noch keine Daten. VAL-Tage meldet jeder Mitarbeiter im Reiter **VAL buchen**.")
        return
    pvg = pv[pv["gemeldet"]]
    vor = logic.month_start(today, -1)
    jan = dt.date(today.year, 1, 1)
    bers = bereiche_von(cons)
    gruppiert = len(bers) > 1
    gruppen = [(b, set(cons.loc[cons["bereich"] == b, "id"])) for b in bers] if gruppiert \
        else [("Team", set(cons["id"]))]

    def _gesamt(df: pd.DataFrame):
        """Gesamt-Planungsgüte: nur Summe geplant gegen Summe gebucht, ohne Projektbezug."""
        if df is None or df.empty:
            return None, 0.0, 0.0
        p_, v_ = float(df["plan"].sum()), float(df["val"].sum())
        return logic.guete(p_, v_), p_, v_

    # ---------------- Kennzahlen je Bereich: projektgenau und gesamt
    for nm, ids in gruppen:
        pb = pvg[pvg["consultant_id"].isin(ids)]
        gb = gm[gm["consultant_id"].isin(ids)]
        if gruppiert:
            st.markdown(f"<div style='margin:.4rem 0 .2rem;font-weight:700;color:{ber_farbe(nm)}'>"
                        f"{html.escape(nm)}</div>", unsafe_allow_html=True)
        k = st.columns(5)
        m_p = logic.guete_agg(pb[pb["monat"] == vor])
        m_g, m_pl, m_va = _gesamt(pb[pb["monat"] == vor])
        j_p = logic.guete_agg(pb[pb["monat"] >= jan])
        j_g, j_pl, j_va = _gesamt(pb[pb["monat"] >= jan])
        k[0].metric(f"{_monat_lang(vor)} projektgenau", fmt_pct(m_p) if m_p is not None else "n/a",
                    help="Je Projekt gerechnet: Abweichungen zwischen Projekten heben sich nicht auf.")
        k[1].metric(f"{_monat_lang(vor)} gesamt", fmt_pct(m_g) if m_g is not None else "n/a",
                    f"{bt(m_pl)} geplant · {bt(m_va)} gebucht", delta_color="off", delta_arrow="off",
                    help="Nur Summe geplanter gegen Summe gebuchter Projekttage (VAL), egal auf welchem Projekt.")
        k[2].metric(f"{today.year} bisher projektgenau", fmt_pct(j_p) if j_p is not None else "n/a")
        k[3].metric(f"{today.year} bisher gesamt", fmt_pct(j_g) if j_g is not None else "n/a",
                    f"{bt(j_pl)} geplant · {bt(j_va)} gebucht", delta_color="off", delta_arrow="off")
        n_gem = int(gb[(gb["monat"] == vor) & gb["gemeldet"]]["consultant_id"].nunique())
        n_soll = int(gb[(gb["monat"] == vor) & ((gb["plan"] > 0) | gb["gemeldet"])]["consultant_id"].nunique())
        k[4].metric(f"VAL gemeldet {_monat_lang(vor)}", f"{n_gem} von {n_soll}",
                    help="Mitarbeiter mit Planung oder Meldung im Monat")

    # ---------------- Verlauf: je Bereich projektgenau (durchgezogen) und gesamt (gestrichelt)
    x = [logic.month_label(m) for m in monate]
    fig = go.Figure()
    fig.add_hrect(y0=0.9, y1=1.0, fillcolor=WP["lilac"], opacity=0.6, line_width=0, layer="below")
    fig.add_hline(y=1.0, line_color=WP["navy"], line_width=1, line_dash="dot")
    for i, (nm, ids) in enumerate(gruppen):
        pb = pvg[pvg["consultant_id"].isin(ids)]
        farbe = ber_farbe(nm) if gruppiert else WP["navy"]
        kurz = BER_KURZ.get(nm, nm)
        y_p = [logic.guete_agg(pb[pb["monat"] == m]) for m in monate]
        y_g = [_gesamt(pb[pb["monat"] == m])[0] for m in monate]
        fig.add_trace(go.Scatter(x=x, y=y_p, mode="lines+markers+text", name=f"{kurz} projektgenau",
                                 text=[fmt_pct(v) if v is not None and v == v else "" for v in y_p],
                                 textposition="bottom center", line=dict(color=farbe, width=3), marker=dict(size=9),
                                 textfont=dict(color=farbe, size=11),
                                 hovertemplate=f"{html.escape(kurz)} projektgenau %{{x}}: %{{y:.0%}}<extra></extra>"))
        fig.add_trace(go.Scatter(x=x, y=y_g, mode="lines+markers+text", name=f"{kurz} gesamt",
                                 text=[fmt_pct(v) if v is not None and v == v else "" for v in y_g],
                                 textposition="top center", line=dict(color=farbe, width=2, dash="dash"),
                                 marker=dict(size=7, symbol="circle-open"), textfont=dict(color=farbe, size=10),
                                 hovertemplate=f"{html.escape(kurz)} gesamt %{{x}}: %{{y:.0%}}<extra></extra>"))
    fig.update_yaxes(range=[0, 1.12])
    fig.update_layout(title=dict(text="Planungsgüte " + ("je Bereich" if gruppiert else "Team")
                                 + ": projektgenau und gesamt (100 % = exakt wie geplant)", x=0, font=dict(size=15)))
    fig = plotly_layout(fig, 370, pct=True)
    fig.update_layout(showlegend=True, legend=dict(orientation="h", x=0, xanchor="left", y=-0.14, yanchor="top"),
                      margin=dict(b=40))
    st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})
    st.caption("Planungsgüte = 1 minus Abweichung zwischen Plan und VAL, bezogen auf den größeren Wert (0 bis 100 %). "
               "Beispiel: 10 geplant, 9 gebucht = 90 %; 9 geplant, 10 gebucht = 90 %. Lila Band = 90 bis 100 %. "
               "Durchgezogen = projektgenau (Abweichungen je Projekt, heben sich nicht auf). Gestrichelt = gesamt "
               "(nur Summe geplant gegen Summe gebucht). Nur Monate, für die VAL gemeldet wurde.")

    sub("Planungsgüte je Mitarbeiter und Monat")
    rows = []
    for r in cons.itertuples():
        row = ({"Bereich": BER_KURZ.get(r.bereich, r.bereich)} if gruppiert else {}) | {"Mitarbeiter": _kurzname(r)}
        for m in monate:
            z = gm[(gm["consultant_id"] == r.id) & (gm["monat"] == m)]
            if z.empty:
                row[logic.month_label(m)] = None
            elif not bool(z["gemeldet"].iat[0]):
                row[logic.month_label(m)] = "offen"
            else:
                g_, d_ = z["guete"].iat[0], z["richtung"].iat[0]
                row[logic.month_label(m)] = (g_, d_)
        rows.append(row)
    raw = pd.DataFrame(rows).set_index(["Bereich", "Mitarbeiter"] if gruppiert else "Mitarbeiter")

    def _txt(v):
        if isinstance(v, str):
            return v
        if not isinstance(v, tuple) or v[0] is None or v[0] != v[0]:
            return ""
        pfeil = " ↑" if v[1] > 0.01 else " ↓" if v[1] < -0.01 else ""
        return f"{v[0]:.0%}{pfeil}"
    df = raw.map(_txt)
    raw = raw.map(lambda v: v[0] if isinstance(v, tuple) else v)
    css = raw.map(lambda v: f"color:{WP['grey']};font-style:italic" if v == "offen" else _guete_css(v))
    sty = df.style.apply(lambda _: css, axis=None)
    st.dataframe(sty, width="stretch", height=38 + 35 * len(df))
    st.caption("↑ = mehr gebucht als geplant, ↓ = weniger gebucht als geplant · offen = Planung vorhanden, VAL noch "
               "nicht gemeldet · leer = weder Plan noch VAL. Lila = 90 % und besser, rot = unter 75 %.")

    # ---------------- VAL nach Tätigkeit
    sub("Gebuchte Tage nach Tätigkeit (VAL)")
    v_p = db.val(von, bis)
    v_s = db.val_sonst(von, bis)
    teile = []
    if not v_p.empty:
        teile.append(v_p.assign(art="Projekte")[["consultant_id", "monat", "art", "tage"]])
    if not v_s.empty:
        teile.append(v_s[["consultant_id", "monat", "art", "tage"]])
    if not teile:
        st.caption("Noch keine VAL gemeldet.")
        return
    va = pd.concat(teile, ignore_index=True)
    arten = ["Projekte", "Allgemeine Tätigkeit", "Akquise"]
    farbe = {"Projekte": WP["navy"], "Allgemeine Tätigkeit": WP["light"], "Akquise": WP["mid"]}
    va = va.assign(bereich=va["consultant_id"].map(dict(zip(cons["id"], cons["bereich"]))) if gruppiert else "")
    va = va[va["bereich"].notna()]
    idx = pd.MultiIndex.from_product([monate, bers if gruppiert else [""]], names=["monat", "bereich"])
    g = va.groupby(["monat", "bereich", "art"])["tage"].sum().unstack(fill_value=0.0).reindex(idx, fill_value=0.0)
    xs = ([[logic.month_label(m_) for m_, _ in g.index], [BER_KURZ.get(b, b) for _, b in g.index]] if gruppiert
          else [logic.month_label(m_) for m_, _ in g.index])
    fig = go.Figure()
    for a in arten:
        if a in g:
            fig.add_trace(go.Bar(x=xs, y=g[a], name=a, marker_color=farbe[a],
                                 hovertemplate=f"{a} %{{x}}: %{{y:.1f}} BT<extra></extra>"))
    anteil = (g.get("Projekte", 0) / g.sum(axis=1).where(g.sum(axis=1) > 0)).tolist()
    fig.add_trace(go.Scatter(x=xs, y=g.sum(axis=1), mode="text",
                             text=[f"{a_:.0%} Projekt" if a_ == a_ else "" for a_ in anteil], textposition="top center",
                             showlegend=False, hoverinfo="skip", textfont=dict(color=WP["navy"], size=11)))
    fig.update_layout(barmode="stack", bargap=0.35, separators=",.",
                      title=dict(text="VAL-Tage je Monat nach Tätigkeit" + (" und Bereich" if gruppiert else ""),
                                 x=0, font=dict(size=15)))
    fig = plotly_layout(fig, 320)
    fig.update_layout(legend=dict(x=1, xanchor="right", y=1.1, traceorder="normal"))
    st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})
    pm = va.groupby(["consultant_id", "art"])["tage"].sum().unstack(fill_value=0.0)
    rows = []
    for r in cons.itertuples():
        if r.id not in pm.index:
            continue
        z = pm.loc[r.id]
        ges = float(z.sum())
        rows.append({**({"Bereich": BER_KURZ.get(r.bereich, r.bereich)} if gruppiert else {}),
                     "Mitarbeiter": _kurzname(r), **{a: bt(z.get(a, 0.0)) for a in arten}, "Gesamt": bt(ges),
                     "Anteil Projekte": fmt_pct(z.get("Projekte", 0.0) / ges) if ges else ""})
    if rows:
        st.dataframe(pd.DataFrame(rows), hide_index=True, width="stretch")
        st.caption("Summen im gewählten Zeitraum, nur gemeldete Monate. Allgemeine Tätigkeit und Akquise erzeugen "
                   "keinen Umsatz und zählen nicht als Auslastung.")


# ====================================================================== Reiter: Mein Bereich
def _mb_shift(n: int) -> None:
    st.session_state.mb_off = 0 if n == 0 else st.session_state.get("mb_off", 0) + n


def _num(v) -> float:
    """'1,5' / 1.5 / '' / None -> float"""
    if v is None or (not isinstance(v, str) and pd.isna(v)):
        return 0.0
    try:
        return float(str(v).replace(",", ".").strip() or 0)
    except ValueError:
        return 0.0


def _kw_col(ws: dt.date) -> str:
    return f"KW {ws.isocalendar().week:02d} · {ws:%d.%m.}"


LEER = " "  # leere Auswahl in Editoren (wird sonst als "None" angezeigt)
WOCHEN_HORIZONT = 12


def _mb_keep() -> None:
    """Pills lassen sich abwählen: dann bleibt die zuletzt gewählte Person aktiv."""
    ss = st.session_state
    if not ss.get("mb_wer"):
        ss.mb_wer = ss.get("mb_last")


def _vorname_label(r) -> str:
    v = _s(r.vollname)
    return f"{v.split()[0]} · {r.name}" if v else r.name


def page_mein_bereich() -> None:
    ss = st.session_state
    cons = db.consultants()
    header("Mein Bereich", "Meine Projekte und meine nächsten 12 Wochen")
    if cons.empty:
        st.info("Noch keine Mitarbeiter angelegt.")
        return
    names = cons["name"].tolist()
    lab = {r.name: _vorname_label(r) for r in cons.itertuples()}
    if not is_admin():
        # Jeder plant nur für sich selbst; nur Admins sehen und bearbeiten alle
        if ss.user not in names:
            st.warning(f"Dein Kürzel **{ss.user}** ist im Team nicht angelegt. Bitte den Admin ansprechen.")
            return
        st.markdown(f"**{html.escape(lab.get(ss.user, ss.user))}**")
        mein_bereich_body(int(cons.loc[cons["name"] == ss.user, "id"].iloc[0]), cons)
        return
    if ss.get("mb_wer") not in names:
        ss.mb_wer = ss.user if ss.user in names else names[0]
    st.pills("Wessen Planung? (nur Admin)", names, format_func=lambda n: lab.get(n, n), key="mb_wer",
             selection_mode="single", on_change=_mb_keep,
             help="Als Admin kannst du für jeden planen. Alle anderen sehen nur ihre eigene Planung.")
    wer = ss.mb_wer
    ss.mb_last = wer
    if wer != ss.user:
        st.caption(f"Du siehst und bearbeitest gerade die Planung von **{lab.get(wer, wer)}**.")
    cid = int(cons.loc[cons["name"] == wer, "id"].iloc[0])
    mein_bereich_body(cid, cons)


def _ort_badge(n_vo: int, n_re: int) -> str:
    if n_vo and n_re:
        return f"<em class='ort mixed'>🏢 {n_vo} · 🏠 {n_re}</em>"
    if n_vo:
        return "<em class='ort onsite'>🏢 Vor Ort</em>"
    if n_re:
        return "<em class='ort remote'>🏠 Remote</em>"
    return "<em class='ort none'>nicht eingeplant</em>"


def project_card_html(p, b, logo_src: str, n_vo: int, n_re: int, naechster, letzter) -> str:
    budget, geb, gep, offen = float(b["budget"]), float(b["gebucht"]), float(b["geplant"]), float(b["offen"])
    basis = max(budget, geb + gep, 0.01)
    w_g, w_p = 100 * geb / basis, 100 * gep / basis
    over = "<div class='o' style='width:%.1f%%'></div>" % (100 * -offen / basis) if offen < -0.01 else ""
    unter = _s(p["name"]) if _s(p["name"]) != _s(p["kunde"]) else ""
    if not unter and pd.notna(p.get("start")) and pd.notna(p.get("ende")):
        unter = f"{p['start']:%d.%m.%y} bis {p['ende']:%d.%m.%y}"
    if budget > 0.01:
        o_txt = (f"<span><b>{bt(offen)}</b> offen</span>" if offen > 0.01 else
                 "<span><b>0</b> offen</span>" if offen > -0.01 else
                 f"<span><b class='neg'>{bt(-offen)}</b> über Budget</span>")
        zahlen = (f"<span title='VAL, wo gemeldet; sonst geplante Tage bis heute'><b>{bt(geb)}</b> Ist</span>"
                  f"<span><b>{bt(gep)}</b> geplant</span>{o_txt}"
                  f"<span class='von'>von {bt(budget)} BT</span>")
    else:
        zahlen = (f"<span><b>{bt(geb)}</b> Ist</span><span><b>{bt(gep)}</b> geplant</span>"
                  "<span class='von'>kein Budget hinterlegt</span>")
    if naechster is not None:
        fuss = f"Nächster Einsatz <b>{logic.WOTAG[naechster.weekday()]} {naechster:%d.%m.}</b>"
        if letzter is not None:
            fuss += f" · geplant bis KW {letzter.isocalendar().week:02d}"
    else:
        fuss = "Keine Tage in den nächsten 12 Wochen"
    return (f"<div class='pc'><div class='pc-h'><img src='{logo_src}'><div><b title='{html.escape(p['label'])}'>"
            f"{html.escape(p['kunde'])}</b><span>{html.escape(unter)}</span></div>{_ort_badge(n_vo, n_re)}</div>"
            f"<div class='bar'><div class='g' style='width:{w_g:.1f}%'></div>"
            f"<div class='p' style='width:{w_p:.1f}%'></div>{over}</div>"
            f"<div class='pc-n'>{zahlen}</div><div class='pc-f'>{fuss}</div></div>")


def plan_popover(cid: int, pid: int, b, projn: pd.DataFrame, ort_jetzt: str) -> None:
    """Offene Tage pauschal verteilen und Ort umstellen."""
    today = heute()
    ss = st.session_state
    user = ss.user
    offen = float(b["offen"])
    k = f"{cid}_{pid}"
    st.markdown(f"**{html.escape(projn.at[pid, 'label'])}**")
    st.caption("Offene Tage pauschal auf die Wochen verteilen. Bereits geplante Tage bleiben erhalten.")
    a, b_ = st.columns(2)
    start = projn.at[pid, "start"]
    ab = a.date_input("Ab", max(today, start) if pd.notna(start) else today, format="DD.MM.YYYY", key=f"pp_ab_{k}")
    menge = b_.number_input("Tage", min_value=0.0, step=0.5, value=max(round(offen * 2) / 2, 0.0), key=f"pp_n_{k}")
    pw = st.segmented_control("Tage pro Woche", [1, 2, 3, 4, 5], default=4 if offen > 0 else 2, key=f"pp_pw_{k}")
    ort = st.segmented_control("Ort", ["Vor Ort", "Remote"], default=ort_jetzt, key=f"pp_ort_{k}",
                               format_func=lambda x: f"{ORT_ICON[x]} {x}")
    plan = db.distribute_open(cid, pid, menge, ab, float(pw or 4), ort or "Vor Ort", user, dry_run=True)
    if plan:
        st.caption("Vorschau: " + " · ".join(f"KW {w.isocalendar().week:02d} +{bt(p_)}" for w, _h, p_ in plan[:8])
                   + (f" … bis KW {plan[-1][0].isocalendar().week:02d}" if len(plan) > 8 else ""))
    if st.button(f"{bt(menge)} Tage verteilen", type="primary", disabled=not plan or menge <= 0, key=f"pp_go_{k}",
                 width="stretch"):
        db.distribute_open(cid, pid, menge, ab, float(pw or 4), ort or "Vor Ort", user)
        ss["mb_msg"] = f"✓ {bt(menge)} Tage für {projn.at[pid, 'label']} verteilt."
        st.rerun()
    st.divider()
    st.caption("Ort für alle künftigen Tage dieses Projekts umstellen")
    neu = st.segmented_control("Neuer Ort", ["Vor Ort", "Remote"], key=f"pp_ort2_{k}", label_visibility="collapsed",
                               format_func=lambda x: f"{ORT_ICON[x]} alle auf {x}")
    if st.button("Ort übernehmen", disabled=not neu, key=f"pp_ortgo_{k}", width="stretch"):
        n = db.set_ort(cid, pid, today, today + dt.timedelta(days=730), neu, user)
        ss["mb_msg"] = f"✓ {n} Tage auf {neu} umgestellt." if n else "Keine Tage umzustellen."
        st.rerun()
    st.divider()
    st.caption("Planung dieses Projekts löschen")
    c1, c2 = st.columns(2)
    von = c1.date_input("Ab", today, format="DD.MM.YYYY", key=f"pp_del_von_{k}")
    bis = c2.date_input("Bis", today + dt.timedelta(days=365), format="DD.MM.YYYY", key=f"pp_del_bis_{k}")
    n_del = len(db.entries(von, bis, consultant_id=cid, project_id=pid)) if von <= bis else 0
    ok = st.checkbox(f"Ja, {n_del} geplante Tage löschen", key=f"pp_del_ok_{k}", disabled=n_del == 0)
    if st.button("🗑 Tage löschen", disabled=not ok, key=f"pp_del_go_{k}", width="stretch"):
        n = db.delete_project_days(cid, pid, von, bis)
        ss["mb_msg"] = f"🗑 {n} Tage von {projn.at[pid, 'label']} gelöscht."
        st.rerun()


def mein_bereich_body(cid: int, cons: pd.DataFrame) -> None:
    today = heute()
    ss = st.session_state
    m0 = logic.monday(today)
    projn = db.projects(active_only=False).set_index("id")
    mem = db.members()
    ent_p = db.entries(consultant_id=cid, project_only=True)
    ms = logic.member_stats(ent_p, mem[mem["consultant_id"] == cid], today).set_index("project_id")
    h_end = m0 + dt.timedelta(weeks=WOCHEN_HORIZONT, days=-3)
    fut = ent_p[(ent_p["tag"] >= today) & (ent_p["tag"] <= h_end)] if not ent_p.empty else ent_p
    fut_any = ent_p[ent_p["tag"] >= today] if not ent_p.empty else ent_p
    pids_all = [p for p in ms.index if p in projn.index and projn.at[p, "aktiv"]]

    def relevant(p: int) -> bool:
        ende = projn.at[p, "ende"]
        return (ms.at[p, "offen"] > 0.01 or (pd.notna(ende) and ende >= today)
                or (not fut_any.empty and (fut_any["project_id"] == p).any()))

    nxt = fut.groupby("project_id")["tag"].min().to_dict() if not fut.empty else {}
    lst = fut_any.groupby("project_id")["tag"].max().to_dict() if not fut_any.empty else {}
    arts = fut.groupby(["project_id", "art"]).size().to_dict() if not fut.empty else {}
    zeige_alle = ss.get("mb_past", False)
    pids = [p for p in pids_all if zeige_alle or relevant(p)]
    pids.sort(key=lambda p: (nxt.get(p, dt.date.max), 0 if ms.at[p, "offen"] > 0.01 else 1,
                             projn.at[p, "label"].lower()))

    # ---------------- Kennzahlen
    f4 = logic.build_fact(m0, m0 + dt.timedelta(days=27), cons[cons["id"] == cid],
                          db.entries(m0, m0 + dt.timedelta(days=27), consultant_id=cid))
    f12 = logic.build_fact(today, h_end, cons[cons["id"] == cid], db.entries(today, h_end, consultant_id=cid))
    offen_sum = float(ms.loc[pids, "offen"].clip(lower=0).sum()) if pids else 0.0
    k = st.columns(4)
    k[0].metric("Laufende Projekte", len([p for p in pids_all if relevant(p)]))
    k[1].metric("Offene Budgettage", bt(offen_sum), help="Budget minus Ist (VAL) minus geplant, über alle Projekte")
    u4 = logic.kpi(f4)
    zp = ziel_von(cons[cons["id"] == cid]).get(cons.loc[cons["id"] == cid, "name"].iloc[0], ZIEL)
    k[2].metric("Auslastung 4 Wochen", fmt_pct(u4),
                None if u4 is None else f"{(u4 - zp) * 100:+.0f} pp vs. Ziel {zp:.0%}")
    k[3].metric("Freie Tage 12 Wochen", bt(f12["frei"].sum() / H) if not f12.empty else "0",
                help="Arbeitstage ohne Projekt, Intern oder Abwesenheit, Feiertage ausgenommen")

    # ---------------- Projektkarten
    sub("Meine Projekte")
    logos = client_logo_map(projn.reset_index())
    for i in range(0, len(pids), 3):
        cols = st.columns(3)
        for col, pid in zip(cols, pids[i:i + 3]):
            with col.container(border=True):
                b = ms.loc[pid]
                n_vo, n_re = arts.get((pid, "Vor Ort"), 0), arts.get((pid, "Remote"), 0)
                st.markdown(project_card_html(projn.loc[pid], b, logos[pid], n_vo, n_re, nxt.get(pid), lst.get(pid)),
                            unsafe_allow_html=True)
                lbl = "Offene Tage verteilen" if b["offen"] > 0.01 else "Tage planen · Ort"
                with st.popover(lbl, width="stretch"):
                    plan_popover(cid, pid, b, projn, "Remote" if n_re > n_vo else "Vor Ort")
    if not pids:
        st.caption("Aktuell keinem laufenden Projekt zugeordnet.")
    a1, a2, _ = st.columns([1.6, 2.2, 3], vertical_alignment="center")
    with a1.popover("＋ Projekt hinzufügen", width="stretch"):
        alle = db.projects()
        rest = alle[~alle["id"].isin(pids)]
        wahl = st.selectbox("Projekt", rest["label"].tolist(), index=None, key=f"mb_addp_{cid}",
                            placeholder="Projekt suchen")
        if st.button("Hinzufügen", disabled=not wahl, key=f"mb_addp_go_{cid}", type="primary"):
            db.add_member(int(rest.loc[rest["label"] == wahl, "id"].iloc[0]), cid)
            ss["mb_msg"] = f"✓ {wahl} hinzugefügt. Tage über die Karte oder unten eintragen."
            st.rerun()
        st.caption("Ganz neues Projekt: Reiter **Projekt anlegen**.")
    a2.toggle("Abgeschlossene Projekte zeigen", key="mb_past",
              help="Standard: nur Projekte mit künftigen Tagen, offenem Budget oder Ende in der Zukunft.")

    # ---------------- Planung 12 Wochen
    sub("Meine nächsten Wochen")
    ss.setdefault("mb_off", 0)
    nb = st.columns([2.2, 0.9, 0.8, 0.9, 1.6], vertical_alignment="center")
    view = nb[0].segmented_control("Ansicht", ["📆 Tage", "📅 Wochen"], default="📆 Tage", key="mb_view",
                                   label_visibility="collapsed")
    view = view or "📆 Tage"
    nb[1].button("◀ Früher", on_click=_mb_shift, args=(-1,), width="stretch", key="mb_prev")
    nb[2].button("Heute", on_click=_mb_shift, args=(0,), width="stretch", key="mb_today")
    nb[3].button("Später ▶", on_click=_mb_shift, args=(1,), width="stretch", key="mb_next")
    n_w = nb[4].selectbox("Wochen", [4, 8, 12], index=2, key="mb_nw", label_visibility="collapsed",
                          format_func=lambda x: f"{x} Wochen")
    start = m0 + dt.timedelta(weeks=ss.mb_off)
    weeks = [start + dt.timedelta(weeks=i) for i in range(n_w)]
    end = weeks[-1] + dt.timedelta(days=4)
    ent = db.entries(start, end, consultant_id=cid)
    # Projekte für Auswahllisten: meine laufenden plus alles, was im Zeitraum schon eingetragen ist
    plist = list(pids)
    if not ent.empty:
        plist += sorted({int(p) for p in ent["project_id"].dropna() if int(p) not in plist and int(p) in projn.index})

    if view == "📅 Wochen":
        wochen_editor(cid, cons, weeks, ent, plist, ms, projn)
    else:
        tabelle = st.toggle("▦ Schnellerfassung als Tabelle", key="mb_edit",
                            help="Aus: Kalender, Tag anklicken zum Bearbeiten. An: Tabelle Projekte × Tage für viele "
                                 "Einträge auf einmal.")
        if tabelle:
            tage_editor(cid, weeks, ent, plist, projn)
        else:
            fact = logic.build_fact(start, end, cons[cons["id"] == cid], ent)
            tage_kalender(cid, weeks, ent, projn, logos, today, fact, plist)
        zeitraum_leeren(cid, plist, projn)
    outlook_export(cid, cons, projn)
    if ss.get("mb_msg"):
        st.toast(ss.pop("mb_msg"))


def _ics_text(t: str) -> str:
    return t.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def _ics_fold(line: str) -> str:
    """Zeilen nach RFC 5545 auf 75 Byte umbrechen."""
    out, cur = [], b""
    for ch in line:
        enc = ch.encode()
        if len(cur) + len(enc) > (75 if not out else 74):
            out.append(cur.decode())
            cur = b""
        cur += enc
    out.append(cur.decode())
    return "\r\n ".join(out)


def build_ics(cid: int, kz: str, ent: pd.DataFrame, projn: pd.DataFrame) -> bytes:
    """Ganztägige Termine je Eintrag, als frei markiert. Feste UID je Person, Tag und Eintrag: Outlook erkennt
    beim erneuten Öffnen bekannte Termine und aktualisiert sie, statt sie doppelt anzulegen."""
    jetzt = dt.datetime.utcnow().strftime("%Y%m%dT%H%M%SZ")
    z = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Ressourcenplanung//DE", "CALSCALE:GREGORIAN",
         "METHOD:PUBLISH", "X-WR-CALNAME:Ressourcenplanung"]
    for e in ent.sort_values(["tag", "stunden"], ascending=[True, False]).itertuples():
        std = f"{e.stunden:g} h".replace(".", ",")
        if pd.notna(e.project_id) and int(e.project_id) in projn.index:
            titel = f"{projn.at[int(e.project_id), 'label']} · {e.art} · {std}"
            kat = "Projekt"
        else:
            n = getattr(e, "notiz", None)
            titel = f"{n.strip() if isinstance(n, str) and n.strip() else e.art} · {std}"
            kat = e.art
        try:
            seq = int(pd.Timestamp(e.updated_at).timestamp() // 60) % 2_000_000_000
        except Exception:
            seq = 0
        uid = f"rp-{cid}-{e.tag:%Y%m%d}-{re.sub(r'[^A-Za-z0-9]', '', str(e.slot))}@ressourcenplanung"
        z += ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{jetzt}", f"SEQUENCE:{seq}",
              f"DTSTART;VALUE=DATE:{e.tag:%Y%m%d}", f"DTEND;VALUE=DATE:{e.tag + dt.timedelta(days=1):%Y%m%d}",
              f"SUMMARY:{_ics_text(titel)}", f"CATEGORIES:{_ics_text(kat)}", "TRANSP:TRANSPARENT",
              "X-MICROSOFT-CDO-BUSYSTATUS:FREE", "X-MICROSOFT-CDO-ALLDAYEVENT:TRUE", "END:VEVENT"]
    z.append("END:VCALENDAR")
    return ("\r\n".join(_ics_fold(x) for x in z) + "\r\n").encode("utf-8")


def outlook_export(cid: int, cons: pd.DataFrame, projn: pd.DataFrame) -> None:
    """Eigene Planung als .ics für Outlook (ganztägig, als frei)."""
    today = heute()
    kz = str(cons.loc[cons["id"] == cid, "name"].iloc[0]) if (cons["id"] == cid).any() else str(cid)
    with st.expander("📅 In Outlook übernehmen"):
        c1, c2 = st.columns([1.4, 2], vertical_alignment="bottom")
        n = c1.segmented_control("Zeitraum", [3, 6], default=3, key="ol_n", format_func=lambda x: f"{x} Wochen") or 3
        start = logic.monday(today)
        bis = start + dt.timedelta(weeks=n, days=-3)
        ent = db.entries(max(today, start), bis, consultant_id=cid)
        c2.download_button(f"⬇️ Kalenderdatei ({len(ent)} Einträge)", build_ics(cid, kz, ent, projn),
                           file_name=f"Planung_{kz}_{today:%Y%m%d}.ics", mime="text/calendar", width="stretch",
                           disabled=ent.empty, key="ol_dl")
        st.caption(f"Ab heute bis {bis:%d.%m.%Y}. Je Eintrag ein ganztägiger Termin, als **frei** markiert, z. B. "
                   "„Kunde | Projekt · Vor Ort · 8 h“. Datei herunterladen und doppelklicken, Outlook übernimmt die "
                   "Termine. Jeder Termin hat eine feste Kennung, damit ein erneuter Download bekannte Termine "
                   "aktualisiert statt sie doppelt anzulegen. Gelöschte Planungen entfernt die Datei nicht aus Outlook.")


UMFANG = {"¼": 0.25, "½": 0.5, "¾": 0.75, "1": 1.0}


def _umfang_label(tage: float) -> str:
    return min(UMFANG, key=lambda k: abs(UMFANG[k] - tage))


def _chips_html(es, projn: pd.DataFrame) -> str:
    chips = []
    if es is not None:
        for e in es.sort_values(["project_id", "art"]).itertuples():
            teil = f"<small>{hh(e.stunden)}</small>" if abs(e.stunden - H) > 1e-9 else ""
            if pd.notna(e.project_id) and int(e.project_id) in projn.index:
                pid = int(e.project_id)
                c = "onsite" if e.art == "Vor Ort" else "remote"
                chips.append(f"<div class='chip {c}'><i class='lg{pid}'></i>"
                             f"<b>{html.escape(projn.at[pid, 'label'])}</b>{teil}</div>")
            else:
                c = "absent" if e.art in logic.ABWESEND_ARTEN else "intern"
                chips.append(f"<div class='chip {c}' title='{html.escape(art_label(e))}'>"
                             f"<b>{html.escape(art_label(e))}</b>{teil}</div>")
    return "".join(chips) or "<div class='frei'>frei</div>"


def tage_kalender(cid: int, weeks: list[dt.date], ent: pd.DataFrame, projn: pd.DataFrame, logos: dict,
                  today: dt.date, fact: pd.DataFrame, plist: list[int]) -> None:
    """Kalender: Zeilen = Wochen, Spalten = Mo bis Fr. Klick auf einen Tag öffnet die Bearbeitung."""
    used = sorted({int(p) for p in ent["project_id"].dropna()}) if not ent.empty else []
    css = "".join(f".lg{p}{{background-image:url('{logos[p]}')}}" for p in used if p in logos)
    st.markdown(f"<style>{css}</style><div class='legend2'>"
                "<span style='background:#E4E4F6;border-left:3px solid #000055'>🏢 Vor Ort</span>"
                "<span style='border:1px dashed #6688BB'>🏠 Remote</span>"
                "<span style='background:#EBEBEB'>🏛️ Intern</span>"
                "<span style='background:repeating-linear-gradient(45deg,#F4F4F4,#F4F4F4 6px,#EAEAEA 6px,#EAEAEA 12px)'>"
                "🌴 Abwesend</span><span style='box-shadow:inset 0 0 0 2px #B3261E'>mehr als 8 h</span>"
                "<span style='color:#6688BB'>👆 Tag anklicken zum Bearbeiten</span></div>", unsafe_allow_html=True)
    by_day = {d: g for d, g in ent.groupby("tag")} if not ent.empty else {}
    if not fact.empty:
        fact = fact.assign(ws=fact["tag"].map(logic.monday))
        g = fact.groupby("ws")[["kapazitaet", "projekt", "frei"]].sum()
    else:
        g = pd.DataFrame()
    widths = [1.05, 2, 2, 2, 2, 2, 1.15]
    hc = st.columns(widths, gap="small")
    for c, t in zip(hc, [""] + logic.WOTAG[:5] + ["Auslastung"]):
        c.markdown(f"<div class='calh'>{t}</div>", unsafe_allow_html=True)
    clicked = None
    for ws in weeks:
        wd = week_days(ws)
        rc = st.columns(widths, gap="small", vertical_alignment="center")
        rc[0].markdown(f"<div class='calkw'><b>KW {ws.isocalendar().week:02d}</b>"
                       f"<span>{wd[0]:%d.%m.} bis {wd[-1]:%d.%m.}</span></div>", unsafe_allow_html=True)
        for i, d in enumerate(wd):
            hol = logic.feiertag_name(d)
            es = by_day.get(d)
            tot = float(es["stunden"].sum()) if es is not None else 0.0
            state = "hol" if hol else "over" if tot > H + 1e-9 else "today" if d == today else \
                "past" if d < today else "std"
            with rc[1 + i].container(border=True, key=f"calday_{d:%Y%m%d}_{state}"):
                if hol:
                    st.markdown(f"<div class='cday'><div class='dn'>{d:%d.%m.}</div>"
                                f"<div class='frei'>{html.escape(hol)}</div></div>", unsafe_allow_html=True)
                    continue
                st.markdown(f"<div class='cday'><div class='dn'>{d:%d.%m.}</div>{_chips_html(es, projn)}</div>",
                            unsafe_allow_html=True)
                if st.button("bearbeiten", key=f"calbtn_{cid}_{d:%Y%m%d}"):
                    clicked = d
        if ws in g.index and g.at[ws, "kapazitaet"] > 0:
            u = g.at[ws, "projekt"] / g.at[ws, "kapazitaet"]
            frei = g.at[ws, "frei"] / H
            col = WP["alert"] if u > 1.0001 else WP["navy"]
            rc[6].markdown(f"<div class='calu'><b style='color:{col}'>{u:.0%}</b><div class='ubar'><div style='width:"
                           f"{min(u, 1) * 100:.0f}%;background:{col}'></div></div>"
                           f"<span>{bt(frei)} {'Tag' if abs(frei - 1) < 1e-9 else 'Tage'} frei</span></div>",
                           unsafe_allow_html=True)
        else:
            rc[6].markdown("<div class='calu'><span>abwesend</span></div>", unsafe_allow_html=True)
    if clicked is not None:
        st.session_state.pop("dd_for", None)
        tag_dialog(cid, clicked, plist, projn)


def _dd_remove(rid: int) -> None:
    ss = st.session_state
    ss.dd_rows = [r for r in ss.dd_rows if r["rid"] != rid]


def _dd_add(opt_map: dict) -> None:
    ss = st.session_state
    wahl = ss.get("dd_add")
    if not wahl:
        return
    art, pid = opt_map[wahl]
    slot = f"p{pid}" if pid is not None else art
    if any(r["slot"] == slot for r in ss.dd_rows):
        ss.dd_add = None
        return
    belegt = sum(UMFANG[ss.get(f"dd_u_{r['rid']}") or _umfang_label(r["tage"])] for r in ss.dd_rows)
    ss.dd_next += 1
    ss.dd_rows.append({"rid": ss.dd_next, "slot": slot, "pid": pid, "art": art,
                       "ort": art if pid is not None else None,
                       "tage": 0.5 if belegt > 0 else 1.0})
    ss.dd_add = None


@st.dialog("Tag bearbeiten", width="large")
def tag_dialog(cid: int, d: dt.date, plist: list[int], projn: pd.DataFrame) -> None:
    ss = st.session_state
    if ss.get("dd_for") != (cid, d):
        es = db.entries(d, d, consultant_id=cid)
        ss.dd_for = (cid, d)
        ss.dd_next = 0
        ss.dd_rows = []
        for e in es.itertuples():
            ss.dd_next += 1
            pid = int(e.project_id) if pd.notna(e.project_id) else None
            ss.dd_rows.append({"rid": ss.dd_next, "slot": e.slot, "pid": pid, "art": e.art,
                               "ort": e.art if pid is not None else None, "tage": e.stunden / H,
                               "notiz": e.notiz if isinstance(getattr(e, "notiz", None), str) else None})
    st.markdown(f"<div class='ptitle'><b>{logic.WOTAG[d.weekday()]} {d:%d.%m.%Y}</b> "
                f"<span>· KW {d.isocalendar().week:02d}</span></div>", unsafe_allow_html=True)
    if not ss.dd_rows:
        st.caption("Noch nichts geplant. Unten Projekt oder Abwesenheit hinzufügen.")
    for r in ss.dd_rows:
        c = st.columns([3.2, 2.2, 2.4, 0.6], vertical_alignment="center")
        if r["pid"] is not None and r["pid"] in projn.index:
            c[0].markdown(f"**{html.escape(projn.at[r['pid'], 'label'])}**")
            c[1].segmented_control("Ort", ["Vor Ort", "Remote"], default=r["ort"] if r["ort"] in ORT_ICON else "Vor Ort",
                                   key=f"dd_o_{r['rid']}", format_func=lambda x: f"{ORT_ICON[x]} {x}",
                                   label_visibility="collapsed")
        else:
            c[0].markdown(f"**{html.escape(r['notiz']) if r.get('notiz') else ART_ICON.get(r['art'], '') + ' ' + r['art']}**")
        c[2].segmented_control("Umfang", list(UMFANG), default=_umfang_label(r["tage"]), key=f"dd_u_{r['rid']}",
                               format_func=lambda x: f"{x} Tag", label_visibility="collapsed")
        c[3].button("✕", key=f"dd_x_{r['rid']}", on_click=_dd_remove, args=(r["rid"],), type="tertiary",
                    help="Entfernen")
    # Hinzufügen
    opt_map = {}
    for pid in plist:
        opt_map[projn.at[pid, "label"]] = ("Vor Ort", pid)
    for pid in db.projects()["id"]:
        if int(pid) not in plist:
            opt_map[projn.at[int(pid), "label"]] = ("Vor Ort", int(pid))
    for a in list(logic.ABWESEND_ARTEN) + list(logic.INTERN_ARTEN):
        opt_map[f"{ART_ICON.get(a, '')} {a}"] = (a, None)
    st.selectbox("＋ Hinzufügen", list(opt_map), index=None, key="dd_add", placeholder="Projekt, Urlaub, Intern …",
                 on_change=_dd_add, args=(opt_map,))
    summe = sum(UMFANG[ss.get(f"dd_u_{r['rid']}") or _umfang_label(r["tage"])] for r in ss.dd_rows)
    if summe > 1.0001:
        st.markdown(f"<span style='color:{WP['alert']};font-weight:700'>⚠ {bt(summe)} Tage an einem Tag "
                    "(überbucht)</span>", unsafe_allow_html=True)
    elif ss.dd_rows:
        st.caption(f"Summe: {bt(summe)} Tag")
    st.divider()
    wd = [d - dt.timedelta(days=d.weekday()) + dt.timedelta(days=i) for i in range(5)]
    andere = [x for x in wd if x != d and not logic.feiertag_name(x)]
    auch = st.pills("Gleich auch für diese Tage übernehmen", andere, selection_mode="multi", key="dd_auch",
                    format_func=lambda x: f"{logic.WOTAG[x.weekday()]} {x:%d.%m.}")
    b1, b2 = st.columns(2)
    if b1.button("Speichern", type="primary", width="stretch", key="dd_save"):
        tage = [d] + list(auch or [])
        rows = []
        for t in tage:
            for r in ss.dd_rows:
                umf = UMFANG[ss.get(f"dd_u_{r['rid']}") or _umfang_label(r["tage"])]
                if r["pid"] is not None:
                    art = ss.get(f"dd_o_{r['rid']}") or r["ort"] or "Vor Ort"
                    rows.append((cid, t, art, r["pid"], umf * H))
                    db.add_member(r["pid"], cid)
                else:
                    rows.append((cid, t, r["art"], None, umf * H, r.get("notiz")))
        db.clear_days(cid, tage)
        if rows:
            db.upsert_entries(rows, ss.user)
        for k in [k for k in ss.keys() if str(k).startswith("dd_")]:
            ss.pop(k, None)
        ss["mb_msg"] = f"✓ {len(tage)} Tag{'e' if len(tage) > 1 else ''} gespeichert."
        st.rerun()
    if b2.button("Abbrechen", width="stretch", key="dd_cancel"):
        for k in [k for k in ss.keys() if str(k).startswith("dd_")]:
            ss.pop(k, None)
        st.rerun()
    tage_l = [d] + list(auch or [])
    if st.button(f"🗑 {'Tag' if len(tage_l) == 1 else f'{len(tage_l)} Tage'} komplett leeren", type="tertiary",
                 key="dd_clear", help="Löscht alle Einträge an diesem Tag (und an den oben markierten Tagen)"):
        n = db.clear_days(cid, tage_l)
        for k in [k for k in ss.keys() if str(k).startswith("dd_")]:
            ss.pop(k, None)
        ss["mb_msg"] = f"🗑 {n} Einträge gelöscht."
        st.rerun()


def zeitraum_leeren(cid: int, plist: list[int], projn: pd.DataFrame) -> None:
    """Mehrere Tage auf einmal löschen: Zeitraum, optional nur ein Projekt."""
    today = heute()
    ss = st.session_state
    with st.expander("🗑 Planung löschen (Zeitraum, optional nur ein Projekt)"):
        c1, c2, c3 = st.columns([1, 1, 2.2])
        von = c1.date_input("Von", today, format="DD.MM.YYYY", key="zl_von")
        bis = c2.date_input("Bis", today + dt.timedelta(weeks=4), format="DD.MM.YYYY", key="zl_bis")
        opts = {"Alles (Projekte, Urlaub, Intern)": None} | {projn.at[p, "label"]: p for p in plist}
        was = c3.selectbox("Was", list(opts), key="zl_was")
        pid = opts[was]
        e = db.entries(von, bis, consultant_id=cid, project_id=pid) if von <= bis else pd.DataFrame()
        n_del = len(e)
        ok = st.checkbox(f"Ja, {n_del} Einträge löschen", key="zl_ok", disabled=n_del == 0)
        if st.button("🗑 Löschen", disabled=not ok, key="zl_go"):
            n = db.delete_project_days(cid, pid, von, bis)
            ss.pop("zl_ok", None)
            ss["mb_msg"] = f"🗑 {n} Einträge gelöscht."
            st.rerun()


def tage_kalender_html(weeks: list[dt.date], ent: pd.DataFrame, projn: pd.DataFrame, logos: dict,
                       today: dt.date, fact: pd.DataFrame) -> str:
    """Kalender zum Ansehen: Zeilen = Wochen, Spalten = Mo bis Fr, je Tag Logo, Kunde, Ort."""
    used = sorted({int(p) for p in ent["project_id"].dropna()}) if not ent.empty else []
    css = "".join(f".lg{p}{{background-image:url('{logos[p]}')}}" for p in used if p in logos)
    by_day = {d: g for d, g in ent.groupby("tag")} if not ent.empty else {}
    if not fact.empty:
        fact = fact.assign(ws=fact["tag"].map(logic.monday))
        g = fact.groupby("ws")[["kapazitaet", "projekt", "frei"]].sum()
    else:
        g = pd.DataFrame()
    out = [f"<style>{css}</style><div class='legend2' style='margin-bottom:4px'>"
           "<span style='background:#E4E4F6;border-left:3px solid #000055'>🏢 Vor Ort</span>"
           "<span style='border:1px dashed #6688BB'>🏠 Remote</span>"
           "<span style='background:#EBEBEB'>🏛️ Intern</span>"
           "<span style='background:repeating-linear-gradient(45deg,#F4F4F4,#F4F4F4 6px,#EAEAEA 6px,#EAEAEA 12px)'>"
           "🌴 Abwesend</span><span style='box-shadow:inset 0 0 0 2px #B3261E'>mehr als 8 h</span></div>",
           "<table class='cal'><tr><th class='kw'></th>"
           + "".join(f"<th>{w}</th>" for w in logic.WOTAG[:5]) + "<th class='u'>Auslastung</th></tr>"]
    for ws in weeks:
        wd = week_days(ws)
        out.append(f"<tr><td class='kw'><b>KW {ws.isocalendar().week:02d}</b>"
                   f"<span>{wd[0]:%d.%m.} bis {wd[-1]:%d.%m.}</span></td>")
        for d in wd:
            cls = ["d"]
            if d == today:
                cls.append("today")
            if d < today:
                cls.append("past")
            hol = logic.feiertag_name(d)
            if hol:
                out.append(f"<td class='hol'><div class='dn'>{d:%d.%m.}</div>{html.escape(hol)}</td>")
                continue
            es = by_day.get(d)
            tot = float(es["stunden"].sum()) if es is not None else 0.0
            if tot > H + 1e-9:
                cls.append("over")
            chips = []
            if es is not None:
                for e in es.sort_values(["project_id", "art"]).itertuples():
                    teil = f"<small>{hh(e.stunden)}</small>" if abs(e.stunden - H) > 1e-9 else ""
                    if pd.notna(e.project_id) and int(e.project_id) in projn.index:
                        pid = int(e.project_id)
                        p = projn.loc[pid]
                        c = "onsite" if e.art == "Vor Ort" else "remote"
                        chips.append(f"<div class='chip {c}' title='{html.escape(p['label'])} · {e.art} · "
                                     f"{hh(e.stunden)}'><i class='lg{pid}'></i><b>{html.escape(p['label'])}</b>"
                                     f"{teil}</div>")
                    else:
                        c = "absent" if e.art in logic.ABWESEND_ARTEN else "intern"
                        chips.append(f"<div class='chip {c}' title='{html.escape(art_label(e))}'>"
                             f"<b>{html.escape(art_label(e))}</b>{teil}</div>")
            body = "".join(chips) or "<div class='frei'>frei</div>"
            out.append(f"<td class='{' '.join(cls)}'><div class='dn'>{d:%d.%m.}</div>{body}</td>")
        if ws in g.index and g.at[ws, "kapazitaet"] > 0:
            u = g.at[ws, "projekt"] / g.at[ws, "kapazitaet"]
            frei = g.at[ws, "frei"] / H
            col = WP["alert"] if u > 1.0001 else WP["navy"]
            out.append(f"<td class='u'><b style='color:{col}'>{u:.0%}</b><div class='ubar'><div style='width:"
                       f"{min(u, 1) * 100:.0f}%;background:{col}'></div></div>"
                       f"<span>{bt(frei)} {'Tag' if abs(frei - 1) < 1e-9 else 'Tage'} frei</span></td>")
        else:
            out.append("<td class='u'><span>abwesend</span></td>")
        out.append("</tr>")
    out.append("</table>")
    return "".join(out)


def _opt_proj(pid: int, ort: str, projn: pd.DataFrame) -> str:
    return f"{ORT_ICON[ort]} {projn.at[pid, 'label']}"


def _opt_art(art: str) -> str:
    return f"{ART_ICON.get(art, '')} {art}"


EDIT_WOCHEN = 2  # Wochen je Block im Tages-Editor (passt ohne Querscrollen)


def _te_wert(icon: str, tage: float) -> str:
    """Zellwert im Tages-Editor, z. B. '🏢 1', '🏠 ½', '🌴 1'."""
    t = {1.0: "1", 0.5: "½"}.get(round(tage, 2), bt(tage))
    return f"{icon} {t}"


def _te_parse(v: str) -> tuple[str, float] | None:
    """'🏢 ½' -> ('🏢', 0.5); leer -> None"""
    if v is None or not str(v).strip():
        return None
    icon, _, t = str(v).strip().partition(" ")
    t = t.strip()
    tage = 0.5 if t == "½" else _num(t)
    return (icon, tage) if tage > 0 else None


def tage_editor(cid: int, weeks: list[dt.date], ent: pd.DataFrame, plist: list[int], projn: pd.DataFrame) -> None:
    """Tage bearbeiten im gewohnten Format: Zeilen = Projekte und Abwesenheiten, Spalten = Arbeitstage mit Datum.
    Mehrere Projekte am selben Tag möglich (z. B. je ½)."""
    ss = st.session_state
    bloecke = [weeks[i:i + EDIT_WOCHEN] for i in range(0, len(weeks), EDIT_WOCHEN)]
    blab = {i: f"KW {b[0].isocalendar().week:02d}" + (f" bis {b[-1].isocalendar().week:02d}" if len(b) > 1 else "")
                 + f" · ab {b[0]:%d.%m.}" for i, b in enumerate(bloecke)}
    if ss.get("mb_te_blk") not in blab:
        ss.mb_te_blk = 0
    st.pills("Zeitraum", list(blab), format_func=blab.get, key="mb_te_blk", selection_mode="single",
             label_visibility="collapsed")
    weeks = bloecke[ss.mb_te_blk if ss.mb_te_blk is not None else 0]
    days = [ws + dt.timedelta(days=i) for ws in weeks for i in range(5)]
    e = ent[ent["tag"].isin(days)].copy() if not ent.empty else ent
    if not e.empty:
        e["key"] = [f"p{int(p)}" if pd.notna(p) else a for p, a in zip(e["project_id"], e["art"])]
    # Zeilen: meine Projekte, dann Abwesenheiten / Intern
    keys = [f"p{p}" for p in plist]
    sonst = ["Urlaub", "Intern"] + [a for a in ("Akquise", "Krank", "Elternzeit")
                                    if not e.empty and (e["art"] == a).any()]
    col = {d: f"{logic.WOTAG[d.weekday()]} {d.day}.{d.month}." for d in days}
    icon_of_art = {a: ART_ICON.get(a, "•") for a in sonst}
    cell = {}
    if not e.empty:
        for r in e.itertuples():
            ic = ORT_ICON.get(r.art, "🏢") if r.key.startswith("p") else icon_of_art.get(r.art, "•")
            cell[(r.key, r.tag)] = _te_wert(ic, r.stunden / H)
    rows = []
    for kk in keys + sonst:
        lab = projn.at[int(kk[1:]), "label"] if kk.startswith("p") else f"{ART_ICON.get(kk, '')} {kk}"
        r = {"key": kk, "Projekt / Tätigkeit": lab}
        for d in days:
            r[col[d]] = "🎉" if logic.feiertag_name(d) else cell.get((kk, d), LEER)
        rows.append(r)
    original = pd.DataFrame(rows)
    p_opts = [LEER, "🏢 1", "🏢 ½", "🏠 1", "🏠 ½"]
    a_opts = [LEER] + [_te_wert(ic, t) for ic in dict.fromkeys(icon_of_art.values()) for t in (1.0, 0.5)]
    opts = p_opts + [o for o in a_opts if o not in p_opts] + ["🎉"]
    for c in col.values():
        opts += [v for v in original[c] if v not in opts]
    cfg = {"key": None,
           "Projekt / Tätigkeit": st.column_config.TextColumn("Projekt / Tätigkeit", disabled=True, width="medium",
                                                              pinned=True)}
    for d, c in col.items():
        hol = logic.feiertag_name(d)
        cfg[c] = st.column_config.SelectboxColumn(
            c, options=opts, width=82,
            help=f"KW {d.isocalendar().week:02d}" + (f" · Feiertag: {hol}" if hol else ""))
    st.caption("Je Tag und Projekt wählen: **🏢 vor Ort** oder **🏠 remote**, **1** = ganzer Tag, **½** = halber Tag. "
               "Mehrere Projekte am selben Tag sind möglich (z. B. 🏢 ½ und 🏠 ½). 🎉 = Feiertag. "
               "Zeitraum oben wählen.")
    key = f"mb_te2_{cid}_{weeks[0]}"
    edited = st.data_editor(original, column_config=cfg, hide_index=True, width="stretch", num_rows="fixed",
                            key=key, height=38 + 35 * len(original))
    # Tagessumme (aus der aktuellen Eingabe, noch vor dem Speichern)
    summe = {}
    for d, c in col.items():
        if logic.feiertag_name(d):
            summe[c] = "Feiertag"
            continue
        t = sum((_te_parse(v) or ("", 0.0))[1] for v in edited[c])
        summe[c] = "" if t == 0 else (f"⚠ {bt(t)}" if t > 1.0001 else bt(t))
    sdf = pd.DataFrame([{"Projekt / Tätigkeit": "Σ Tage je Tag", **summe}])
    scfg = {c: st.column_config.TextColumn(c, width=82) for c in col.values()}
    scfg["Projekt / Tätigkeit"] = st.column_config.TextColumn("", width="medium", pinned=True)
    st.dataframe(sdf.style.map(lambda v: f"color:{WP['alert']};font-weight:700" if str(v).startswith("⚠")
                               else f"color:{WP['navy']};font-weight:600"),
                 width="stretch", hide_index=True, column_config=scfg)
    if st.button("Änderungen speichern", type="primary", key="mb_te_save"):
        ups, dels, skip = [], [], 0
        for i in range(len(original)):
            kk = original.at[i, "key"]
            pid = int(kk[1:]) if kk.startswith("p") else None
            for d, c in col.items():
                old, new = original.at[i, c], edited.at[i, c]
                if new is None:
                    new = LEER
                if old == new:
                    continue
                if logic.feiertag_name(d):
                    skip += 1
                    continue
                p = _te_parse(new)
                slot = kk
                if p is None:
                    dels.append((d, slot))
                elif pid is not None:
                    ort = "Remote" if p[0] == "🏠" else "Vor Ort"
                    ups.append((cid, d, ort, pid, p[1] * H))
                else:
                    ups.append((cid, d, kk, None, p[1] * H))
        if ups:
            db.upsert_entries(ups, ss.user)
        if dels:
            db.delete_slots(cid, dels)
        n = len(ups) + len(dels)
        ss.pop(key, None)
        ss["mb_msg"] = (f"✓ {n} Einträge gespeichert." if n else "Keine Änderungen.") + \
                       (" Feiertage bleiben frei." if skip else "")
        st.rerun()


def wochen_editor(cid: int, cons: pd.DataFrame, weeks: list[dt.date], ent: pd.DataFrame, plist: list[int],
                  ms: pd.DataFrame, projn: pd.DataFrame) -> None:
    """Wochenmatrix: je Projekt bzw. Abwesenheit Tage pro Woche, die App verteilt auf freie Tage."""
    ss = st.session_state
    user = ss.user
    start, end = weeks[0], weeks[-1] + dt.timedelta(days=4)
    e = ent.copy()
    if not e.empty:
        e["key"] = [f"p{int(p)}" if pd.notna(p) else a for p, a in zip(e["project_id"], e["art"])]
        e["ws"] = e["tag"].map(logic.monday)
    wk_days = e.groupby(["key", "ws"])["stunden"].sum().div(H).to_dict() if not e.empty else {}
    keys = [f"p{p}" for p in plist]
    sonst = ["Urlaub", "Intern"] + [a for a in ("Krank", "Akquise", "Elternzeit")
                                    if not e.empty and (e["art"] == a).any()]
    rows = []
    for kk in keys + sonst:
        if kk.startswith("p"):
            pid = int(kk[1:])
            b = ms.loc[pid] if pid in ms.index else None
            arts = e.loc[e["key"] == kk, "art"] if not e.empty else pd.Series(dtype=str)
            ort = arts.mode().iat[0] if len(arts) else "Vor Ort"
            row = {"key": kk, "Projekt": projn.at[pid, "label"], "Ort": ort,
                   "Offen": "" if b is None else bt(b["offen"])}
        else:
            row = {"key": kk, "Projekt": _opt_art(kk), "Ort": LEER, "Offen": ""}
        for ws in weeks:
            v = wk_days.get((kk, ws), 0.0)
            row[_kw_col(ws)] = bt(v) if v else LEER
        rows.append(row)
    original = pd.DataFrame(rows)
    kw_cols = [_kw_col(w) for w in weeks]
    tage_opts = [LEER] + [bt(x / 2) for x in range(1, 11)]
    for c_ in kw_cols:
        tage_opts += [v for v in original[c_] if v.strip() and v not in tage_opts]
    cfg = {"key": None,
           "Projekt": st.column_config.TextColumn("Projekt / Tätigkeit", disabled=True, width="medium"),
           "Ort": st.column_config.SelectboxColumn("Ort", options=["Vor Ort", "Remote", LEER], width="small",
                                                   help="Ändern = alle Tage des Projekts in den angezeigten Wochen "
                                                        "umstellen. Gilt auch für neu eingetragene Tage."),
           "Offen": st.column_config.TextColumn("Offen BT", disabled=True, width="small")}
    for c_ in kw_cols:
        cfg[c_] = st.column_config.SelectboxColumn(c_, options=tage_opts, width="small")
    st.caption("Tage pro Woche wählen (0,5 = halber Tag), leeren = keine Tage. Die App verteilt sie auf freie "
               "Arbeitstage der Woche, Feiertage ausgenommen. Genaue Tage: Ansicht **📆 Tage**.")
    key = f"mb_ed_{cid}_{start}_{len(weeks)}"
    edited = st.data_editor(original, column_config=cfg, hide_index=True, width="stretch", num_rows="fixed",
                            key=key, height=38 + 35 * max(len(original), 1))
    fact = logic.build_fact(start, end, cons[cons["id"] == cid], ent)
    if not fact.empty:
        fact["ws"] = fact["tag"].map(logic.monday)
        g = fact.groupby("ws")[["kapazitaet", "projekt"]].sum()
        summ = pd.DataFrame([{_kw_col(w): (g.at[w, "projekt"] / g.at[w, "kapazitaet"]
                                           if w in g.index and g.at[w, "kapazitaet"] else None) for w in weeks}],
                            index=["Auslastung"])
        st.dataframe(styled_pct(summ), width="stretch")
    if st.button("Änderungen speichern", type="primary", key="mb_save"):
        n = 0
        for i, r in edited.iterrows():
            o = original.iloc[i]
            kk = o["key"]
            pid = int(kk[1:]) if kk.startswith("p") else None
            ort = r["Ort"] if pid is not None and r["Ort"] in ("Vor Ort", "Remote") else "Vor Ort"
            art = ort if pid is not None else kk
            for w in weeks:
                c_ = _kw_col(w)
                if abs(_num(o[c_]) - _num(r[c_])) > 1e-9:
                    n += db.set_week_target(cid, w, art, pid, _num(r[c_]), ort, user)
            if pid is not None and o["Ort"] != r["Ort"] and r["Ort"] in ("Vor Ort", "Remote"):
                n += db.set_ort(cid, pid, start, end, r["Ort"], user)
        ss.pop(key, None)
        ss["mb_msg"] = f"✓ Gespeichert ({n} Tage angepasst)." if n else "Keine Änderungen."
        st.rerun()


# ====================================================================== Reiter: Projekt anlegen
def page_projekt() -> None:
    today = heute()
    ss = st.session_state
    user = ss.user
    header("Projekt anlegen", "Kunde mit Logo, Team und Tage je Mitarbeiter in einem Durchgang")
    cons = db.consultants()
    names = cons["name"].tolist()
    plab = person_label(cons)
    idmap = dict(zip(cons["name"], cons["id"]))
    mode = st.segmented_control("Was möchtest du tun?", ["＋ Neues Projekt", "✏️ Bestehendes bearbeiten"],
                                default="＋ Neues Projekt", key="pa_mode", label_visibility="collapsed")
    if ss.get("pa_msg"):
        st.success(ss.pop("pa_msg"))
    edit = mode == "✏️ Bestehendes bearbeiten"
    pre, pid = None, None
    if edit:
        allp = db.projects(active_only=False).sort_values(["aktiv", "label"], ascending=[False, True])
        labels = [l_ + ("" if a else " (abgeschlossen)") for l_, a in zip(allp["label"], allp["aktiv"])]
        wahl = st.selectbox("Projekt wählen", labels, index=None, key="pa_sel", placeholder="Projekt suchen")
        if not wahl:
            return
        pre = allp.iloc[labels.index(wahl)]
        pid = int(pre["id"])
    k = f"{pid or 'neu'}"

    # ---------------- 1 | Kunde & Logo
    sub("1 | Kunde und Logo")
    if edit:
        cl = db.clients(active_only=False).set_index("id").loc[int(pre["client_id"])]
        c1, c2 = st.columns([3, 2])
        kname = c1.text_input("Kunde", cl["name"], key=f"pa_kname_{k}")
        dom = c2.text_input("Website (für automatisches Logo)", _s(cl["domain"]), key=f"pa_dom_{k}")
        cid_k, new_name, has_logo = int(pre["client_id"]), None, isinstance(cl["logo"], str) and bool(cl["logo"])
        cur_logo = logo(cl["name"], dom, cl["logo"])
    else:
        cid_k, new_name, dom = client_picker(f"pa{k}")
        kname = new_name
        has_logo = False
        cur_logo = None
        if cid_k:
            cl = db.clients(active_only=False).set_index("id").loc[cid_k]
            has_logo = isinstance(cl["logo"], str) and bool(cl["logo"])
            cur_logo = logo(cl["name"], cl["domain"], cl["logo"])
        elif new_name:
            cur_logo = logo(new_name, dom, None)
    lc1, lc2 = st.columns([1, 5], vertical_alignment="center")
    if cur_logo:
        lc1.image(cur_logo, width=64)
    up = lc2.file_uploader("Kundenlogo hochladen (PNG, JPG oder SVG)", type=["png", "jpg", "jpeg", "svg"],
                           key=f"pa_logo_{k}")
    if (cid_k or new_name) and not has_logo and not up:
        lc2.caption("💡 Noch kein eigenes Logo hinterlegt. Bitte hochladen, damit es in Kalender und Übersicht "
                    "erscheint. Ohne Upload wird das Logo über die Website gesucht.")

    # ---------------- 2 | Projekt & Budget
    sub("2 | Projekt")
    c1, c2 = st.columns([3, 2])
    pname = c1.text_input("Projektname", pre["name"] if edit else "", key=f"pa_name_{k}",
                          placeholder="z. B. OpEx 2.0")
    leit_default = (cons.loc[cons["id"] == pre["leiter_id"], "name"].iloc[0]
                    if edit and pd.notna(pre["leiter_id"]) and (cons["id"] == pre["leiter_id"]).any()
                    else user if user in names else names[0])
    leiter = c2.selectbox("Projektleitung", names, index=names.index(leit_default), key=f"pa_leit_{k}",
                          format_func=lambda n: plab.get(n, n))
    c3, c4 = st.columns(2)
    start = c3.date_input("Start", pre["start"] if edit and pd.notna(pre["start"]) else today,
                          format="DD.MM.YYYY", key=f"pa_start_{k}")
    ende = c4.date_input("Ende", pre["ende"] if edit and pd.notna(pre["ende"]) else today + dt.timedelta(days=90),
                         format="DD.MM.YYYY", key=f"pa_ende_{k}")

    # ---------------- 3 | Team & Tage je Mitarbeiter
    sub("3 | Team: Tage je Mitarbeiter")
    stats = {}
    if edit:
        ms = logic.member_stats(db.entries(project_id=pid), db.members(pid), today)
        cmap = dict(zip(db.consultants(active_only=False)["id"], db.consultants(active_only=False)["name"]))
        for m in ms.itertuples():
            n_ = cmap.get(int(m.consultant_id))
            if n_:
                stats[n_] = (float(m.budget), float(m.gebucht + m.geplant), float(m.offen))
    vorauswahl = [n_ for n_ in stats if n_ in names] if edit else ([user] if user in names else [])
    wer = st.pills("Wer ist im Projekt?", names, selection_mode="multi", default=vorauswahl, key=f"pa_wer_{k}",
                   format_func=lambda n: plab.get(n, n))
    team_rows = []
    if not wer:
        st.caption("Oben die Mitarbeiter anklicken, die im Projekt sind.")
    for n_ in wer:
        b0, v0, o0 = stats.get(n_, (0.0, 0.0, 0.0))
        with st.container(border=True):
            c1, c2, c3 = st.columns([2.2, 1.4, 2.4], vertical_alignment="center")
            c1.markdown(f"**{html.escape(plab.get(n_, n_))}**")
            tage = c2.number_input("Tage", min_value=0.0, step=0.5, value=b0, key=f"pa_bt_{k}_{n_}",
                                   label_visibility="collapsed", format="%.1f")
            if edit and n_ in stats:
                offen = tage - v0
                farbe = WP["alert"] if offen < -0.01 else WP["navy"]
                c3.markdown(f"<span style='color:{WP['mid']};font-size:.85rem'>verplant <b>{bt(v0)}</b> · offen "
                            f"<b style='color:{farbe}'>{bt(offen)}</b> BT</span>", unsafe_allow_html=True)
            else:
                c3.markdown(f"<span style='color:{WP['mid']};font-size:.85rem'>Beratertage für "
                            f"{html.escape(n_)}</span>", unsafe_allow_html=True)
        team_rows.append((n_, float(tage or 0)))
    team = pd.DataFrame(team_rows, columns=["Mitarbeiter", "Budget BT"])
    budget = float(team["Budget BT"].sum()) if not team.empty else 0.0
    st.markdown(f"<div class='kpis'><div><b>{bt(budget)} BT</b><span>Projektbudget gesamt "
                f"(Summe der Mitarbeiter)</span></div></div>", unsafe_allow_html=True)

    # ---------------- 4 | Direkt einplanen (optional)
    sub("4 | Direkt in die Wochen einplanen (optional)")
    plan = st.toggle("Budgettage der Berater automatisch auf die Wochen verteilen", key=f"pa_plan_{k}",
                     help="Verteilt je Berater die noch offenen Tage ab Projektstart. Feinschliff danach in "
                          "**Mein Bereich**.")
    if plan:
        p1, p2 = st.columns(2)
        pw = p1.segmented_control("Tage pro Woche", [1, 2, 3, 4, 5], default=4, key=f"pa_pw_{k}")
        ort = p2.segmented_control("Ort", ["Vor Ort", "Remote"], default="Vor Ort", key=f"pa_ort_{k}",
                                   format_func=lambda x: f"{ORT_ICON[x]} {x}")

    ok = bool((cid_k or new_name) and pname.strip() and start <= ende and len(team))
    b1, b2, _ = st.columns([1.6, 1.6, 4])
    if b1.button("Änderungen speichern" if edit else "Projekt anlegen", type="primary", disabled=not ok,
                 key=f"pa_go_{k}"):
        if new_name:
            cid_k = db.create_client(new_name, dom)
        if edit:
            db.update_client(cid_k, kname, dom)
            db.update_project_core(pid, cid_k, pname, budget)
        else:
            pid = db.ensure_project(cid_k, pname)
            db.update_project_core(pid, cid_k, pname, budget)
        db.update_project_meta(pid, start=start, ende=ende, leiter_id=idmap.get(leiter), aktiv=True)
        db.set_members(pid, [(idmap[r["Mitarbeiter"]], float(r["Budget BT"] or 0)) for _, r in team.iterrows()
                             if r["Mitarbeiter"] in idmap])
        if up is not None:
            mime = "image/svg+xml" if up.name.lower().endswith(".svg") else up.type
            db.set_logo(cid_k, f"data:{mime};base64,{base64.b64encode(up.getvalue()).decode()}")
        n = 0
        if plan:
            ms2 = logic.member_stats(db.entries(project_id=pid), db.members(pid), today)
            for m in ms2.itertuples():
                if m.offen > 0.01:
                    n += len(db.distribute_open(int(m.consultant_id), pid, m.offen, max(start, today),
                                                float(pw or 4), ort or "Vor Ort", user))
        for kk in list(ss.keys()):
            if str(kk).startswith("pa") and kk not in ("pa_mode", "pa_sel"):
                ss.pop(kk, None)
        ss["pa_msg"] = (f"✓ Projekt **{pname.strip()}** {'gespeichert' if edit else 'angelegt'}"
                        + (f", in {n} Wochen eingeplant." if n else ".")
                        + " Planung der Tage: Reiter **Mein Bereich**.")
        st.rerun()
    if edit and b2.button("Projekt abschließen" if pre["aktiv"] else "Wieder aktivieren", key=f"pa_close_{k}"):
        db.update_project_meta(pid, aktiv=not bool(pre["aktiv"]))
        ss["pa_msg"] = "Status geändert."
        st.rerun()


# ====================================================================== Reiter: Gesamtübersicht
def _go_step(order: list[int], step: int) -> None:
    cur = st.session_state.get("go_pid")
    i = order.index(cur) if cur in order else -1
    st.session_state.go_pid = order[(i + step) % len(order)]


def _pk_shift(n: int) -> None:
    st.session_state.pk_off = 0 if n == 0 else st.session_state.get("pk_off", 0) + n


def projekte_durchgehen() -> None:
    today = heute()
    ss = st.session_state
    proj = db.projects()
    if proj.empty:
        st.info("Noch keine Projekte.")
        return
    ent_all = db.entries(project_only=True)
    mem = db.members()
    stats = logic.project_stats(proj, ent_all, mem, today).set_index("id")
    fut = set(ent_all.loc[ent_all["tag"] >= today, "project_id"].dropna().astype(int)) if not ent_all.empty else set()
    laufend = [p for p in stats.index if p in fut or stats.at[p, "offen"] > 0.01
               or (pd.notna(stats.at[p, "ende"]) and stats.at[p, "ende"] >= today)]
    alle = ss.get("go_all", False)
    order = stats[stats.index.isin(stats.index if alle else laufend)].sort_values(["kunde", "name"]).index.tolist()
    if not order:
        st.info("Keine laufenden Projekte.")
        st.toggle("Vergangene Projekte zeigen", key="go_all")
        return
    if ss.get("go_pid") not in order:
        ss.go_pid = order[0]
    labels = dict(zip(stats.index, stats["label"]))
    n = st.columns([1.3, 1.3, 4, 1.8], vertical_alignment="center")
    n[0].button("◀ Vorheriges", on_click=_go_step, args=(order, -1), width="stretch", key="go_prev")
    n[1].button("Nächstes ▶", on_click=_go_step, args=(order, 1), type="primary", width="stretch", key="go_next")
    n[2].selectbox("Projekt", order, format_func=lambda p: labels[p], key="go_pid", label_visibility="collapsed")
    n[3].toggle("Vergangene zeigen", key="go_all", help="Standard: nur laufende Projekte")
    st.caption(f"Projekt {order.index(ss.go_pid) + 1} von {len(order)}")
    pid = ss.go_pid
    r = stats.loc[pid]
    cons_all = db.consultants(active_only=False)
    plab = person_label(cons_all)
    lmap = dict(zip(cons_all["id"], cons_all["name"]))

    with st.container(border=True):
        a, b, c = st.columns([0.6, 4, 6], vertical_alignment="center")
        a.image(logo(r["kunde"], r["domain"], r["logo"]), width=48)
        leit = plab.get(lmap.get(int(r["leiter_id"])), "") if pd.notna(r.get("leiter_id")) else ""
        zeit = (f"{r['start']:%d.%m.%Y} bis {r['ende']:%d.%m.%Y}"
                if pd.notna(r["start"]) and pd.notna(r["ende"]) else "ohne Zeitraum")
        b.markdown(f"<div class='ptitle'><b>{html.escape(r['label'])}</b><br><span>{zeit}"
                   f"{' · Leitung: ' + html.escape(leit) if leit else ''}</span></div>", unsafe_allow_html=True)
        c.markdown(kpi_html(r["budget"], r["gebucht"], r["geplant"], r["offen"]), unsafe_allow_html=True)
        if abs(r["nicht_verteilt"]) > 0.01:
            st.caption(f"Projektbudget {bt(r['budget'])} BT, auf Berater verteilt {bt(r['verteilt'])} BT.")

    ms = logic.member_stats(ent_all, mem, today)
    ms = ms[ms["project_id"] == pid]
    if ms.empty:
        st.caption("Noch kein Team zugeordnet.")
    else:
        team = cons_all[cons_all["id"].isin([int(c) for c in ms["consultant_id"]])]
        msi = ms.set_index("consultant_id")
        offen = {}
        for c_id in team["id"]:
            m = msi.loc[c_id]
            col = WP["alert"] if m["offen"] < -0.01 else WP["navy"]
            offen[int(c_id)] = (f"<b style='color:{col}'>{bt(m['offen'])}</b><br>"
                                f"<span class='vn'>von {bt(m['budget'])}</span>")
        ss.setdefault("pk_off", 0)
        n = st.columns([1, 1, 1, 6], vertical_alignment="center")
        n[0].button("◀ Woche", on_click=_pk_shift, args=(-1,), width="stretch", key="pk_prev")
        n[1].button("Heute", on_click=_pk_shift, args=(0,), width="stretch", key="pk_today")
        n[2].button("Woche ▶", on_click=_pk_shift, args=(1,), width="stretch", key="pk_next")
        n[3].caption("Team dieses Projekts, nächste 3 Wochen. Dieses Projekt ist hervorgehoben, andere Einsätze "
                     "blass. Rechts: offene Budgettage im Projekt.")
        start = logic.monday(today) + dt.timedelta(weeks=ss.pk_off)
        st.markdown(kurzfrist_html(start, 3, team, groesse="gross", fokus_pid=pid,
                                   extra=("Offen BT", offen)), unsafe_allow_html=True)
    with st.expander("Alle Projekte im Überblick (Budget, Ist, geplant, offen)"):
        auswertung_projekte()
    projektgroessen(laufend if not alle else list(stats.index))


def projektgroessen(pids: list[int]) -> None:
    """Projektgrößen: Volumen je Projekt (Budget-Tage × Tagessatz), aufgeteilt in Ist und Rest, mit Logos."""
    today = heute()
    sub("Projektgrößen: Volumen je Projekt")
    cons = db.consultants(active_only=False)
    ts = dict(zip(cons["id"], pd.to_numeric(cons["tagessatz_eff"], errors="coerce").fillna(0.0)))
    if sum(ts.values()) <= 0:
        st.caption("Für die Projektgrößen in € bitte im Reiter **Team** Tagessätze je Position hinterlegen.")
        return
    proj = db.projects(active_only=False).set_index("id")
    mem = db.members()
    ms = logic.member_stats(db.entries(project_only=True), mem, today)
    ms = ms[ms["project_id"].isin(pids)]
    if ms.empty:
        st.caption("Keine Daten.")
        return
    ms = ms.assign(satz=ms["consultant_id"].map(ts).fillna(0.0))
    ms = ms.assign(volumen=ms[["budget", "gebucht"]].max(axis=1).add(ms["geplant"].where(ms["budget"] <= 0, 0))
                   * ms["satz"], ist=ms["gebucht"] * ms["satz"])
    names = {r.id: (_s(r.vollname) or r.name) for r in cons.itertuples()}
    g = ms.groupby("project_id").agg(volumen=("volumen", "sum"), ist=("ist", "sum")).reset_index()
    g = g[g["volumen"] > 0].sort_values("volumen")
    if g.empty:
        st.caption("Noch keine Budgets hinterlegt.")
        return
    split = {pid: "<br>".join(f"{names.get(r.consultant_id, '')}: {r.volumen / 1000:,.0f} T€".replace(",", ".")
                              for r in ms[ms["project_id"] == pid].sort_values("volumen", ascending=False)
                              .itertuples() if r.volumen > 0) for pid in g["project_id"]}
    lab = [proj.at[p, "label"] for p in g["project_id"]]
    rest = (g["volumen"] - g["ist"]).clip(lower=0)
    fig = go.Figure([
        go.Bar(y=lab, x=g["ist"], name="Ist (VAL bzw. bis heute)", orientation="h", marker_color=WP["navy"],
               customdata=[split[p] for p in g["project_id"]],
               hovertemplate="<b>%{y}</b><br>Ist: %{x:,.0f} €<br>%{customdata}<extra></extra>"),
        go.Bar(y=lab, x=rest, name="Rest bis Budget", orientation="h", marker_color=WP["light"],
               customdata=[split[p] for p in g["project_id"]],
               hovertemplate="<b>%{y}</b><br>Rest: %{x:,.0f} €<br>%{customdata}<extra></extra>"),
        go.Scatter(y=lab, x=g["volumen"], mode="text", showlegend=False, hoverinfo="skip",
                   text=[f"  {v / 1000:,.0f} T€".replace(",", ".") for v in g["volumen"]], textposition="middle right",
                   textfont=dict(color=WP["navy"], size=12))])
    logos = client_logo_map(proj.reset_index())
    for i, p in enumerate(g["project_id"]):
        fig.add_layout_image(dict(source=logos.get(int(p)), xref="paper", yref="y", x=-0.005, y=lab[i],
                                  sizex=0.05, sizey=0.8, xanchor="right", yanchor="middle", layer="above"))
    fig.update_layout(barmode="stack", bargap=0.3, separators=",.", legend=dict(x=1, xanchor="right", y=1.06))
    fig.update_xaxes(tickformat=",.0f", ticksuffix=" €", range=[0, float(g["volumen"].max()) * 1.18])
    fig.update_yaxes(ticklabelstandoff=34)
    fig = plotly_layout(fig, 90 + 38 * len(g))
    fig.update_layout(margin=dict(l=60, r=10, t=30, b=10), legend=dict(x=1, xanchor="right", y=1.06))
    st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})
    st.caption("Volumen = Budget-Tage je Mitarbeiter × Tagessatz seiner Position (ohne Budget: Ist plus Plan). "
               "Dunkel = bereits erbracht (VAL bzw. bis heute), hell = Rest bis Budget. Maus über den Balken: "
               "Aufteilung je Mitarbeiter.")


def export_block() -> None:
    today = heute()
    ss = st.session_state
    sub("Aktueller Stand als Excel")
    st.caption("Immer der Live-Stand aus der Datenbank. Blatt **Übersicht**: wer hat welche Projekte, Auslastung, "
               "welche Projekte wie verplant sind. Blatt **Planungsgüte**: Plan vs. VAL je Monat. Blatt **Tage**: Mitarbeiter × Projekt × Arbeitstag im gewohnten "
               "Format. Blatt **Wochen**: Summen je KW. Blatt **Projekte**: Budget, gebucht, geplant, offen.")
    c1, c2, c3, c4 = st.columns([1.3, 1.3, 1.6, 2.2], vertical_alignment="bottom")
    von = c1.date_input("Tage von", dt.date(today.year, 1, 1), format="DD.MM.YYYY", key="xl_von")
    bis = c2.date_input("bis", logic.monday(today) + dt.timedelta(weeks=WOCHEN_HORIZONT, days=-3),
                        format="DD.MM.YYYY", key="xl_bis")
    if c3.button("📊 Excel erstellen", type="primary", disabled=von > bis, key="xl_go", width="stretch"):
        with st.spinner("Excel wird erstellt ..."):
            cons = db.consultants()
            ss["xl_bytes"] = export.build_excel(cons, db.projects(active_only=False).set_index("id"),
                                                db.entries(), db.members(), von, bis, today, EXPORT_TITEL,
                                                val=db.val())
            ss["xl_name"] = f"Ressourcenplanung_Stand_{today:%Y%m%d}.xlsx"
    if ss.get("xl_bytes"):
        c4.download_button("⬇️ Excel herunterladen", ss["xl_bytes"], file_name=ss["xl_name"], width="stretch",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="xl_dl")



def _kf_shift(n: int) -> None:
    st.session_state.kf_off = 0 if n == 0 else st.session_state.get("kf_off", 0) + n


def kurzfrist_html(start: dt.date, n_w: int, cons: pd.DataFrame, groesse: str = "mittel",
                   fokus_pid: int | None = None, extra: tuple[str, dict] | None = None) -> str:
    """Kurzfrist: mehrere Wochen nebeneinander, je Tag nur Logo und Stunden, Farbe = Ort, je Woche Auslastung."""
    today = heute()
    weeks = [start + dt.timedelta(weeks=i) for i in range(n_w)]
    end = weeks[-1] + dt.timedelta(days=4)
    ent = db.entries(start, end)
    projn = db.projects(active_only=False).set_index("id")
    logos = client_logo_map(projn.reset_index())
    used = sorted({int(p) for p in ent["project_id"].dropna()}) if not ent.empty else []
    css = "".join(f".lg{p}{{background-image:url('{logos[p]}')}}" for p in used if p in logos)
    by = {k: g for k, g in ent.groupby(["consultant_id", "tag"])} if not ent.empty else {}
    fact = logic.build_fact(start, end, cons, ent)
    if not fact.empty:
        fact = fact.assign(ws=fact["tag"].map(logic.monday))
    h1 = "<tr><th class='nm' rowspan='2'></th>"
    h2 = "<tr>"
    for ws in weeks:
        h1 += f"<th colspan='5' class='wk'>KW {ws.isocalendar().week:02d} · {ws:%d.%m.} bis {ws + dt.timedelta(days=4):%d.%m.}</th>"
        for d in week_days(ws):
            h2 += f"<th class='{'td' if d == today else ''}'>{logic.WOTAG[d.weekday()]} {d.day}.</th>"
    h1 += f"<th class='au' rowspan='2'>Ø Ausl.<br>{n_w} Wo.</th>"
    if extra:
        h1 += f"<th class='ex' rowspan='2'>{html.escape(extra[0])}</th>"
    h1 += "</tr>"
    h2 += "</tr>"
    rows = ""
    ncol = 2 + 5 * n_w + (1 if extra else 0)
    gruppiert = len(bereiche_von(cons)) > 1
    letzt = None
    for c in cons.itertuples():
        b = getattr(c, "bereich", "") or ""
        if gruppiert and b != letzt:
            rows += (f"<tr class='grp'><td colspan='{ncol}' style='color:{ber_farbe(b)};"
                     f"border-left:4px solid {ber_farbe(b)}'>{html.escape(b)}</td></tr>")
            letzt = b
        rows += (f"<tr><td class='nm' title='{html.escape(_s(c.vollname) or c.name)}'>"
                 f"<b>{html.escape(_vorname_label(c))}</b></td>")
        for ws in weeks:
            for d in week_days(ws):
                tdc = " td" if d == today else ""
                hol = logic.feiertag_name(d)
                if hol:
                    rows += f"<td class='ho{tdc}' title='{html.escape(hol)}'>Feiertag</td>"
                    continue
                g = by.get((c.id, d))
                if g is None:
                    rows += f"<td class='fr{tdc}'></td>"
                    continue
                tot = float(g["stunden"].sum())
                chips = ""
                for e in g.sort_values("stunden", ascending=False).itertuples():
                    hrs = f"{e.stunden:g}h".replace(".", ",")
                    if pd.notna(e.project_id) and int(e.project_id) in projn.index:
                        pid = int(e.project_id)
                        k = ("on" if e.art == "Vor Ort" else "re") + (" bl" if fokus_pid and pid != fokus_pid else "")
                        chips += (f"<span class='k {k}' title='{html.escape(projn.at[pid, 'label'])} · {e.art} · "
                                  f"{hrs}'><i class='lg{pid}'></i>{hrs}</span>")
                    else:
                        k = "ab" if e.art in logic.ABWESEND_ARTEN else "in"
                        ic = "📌" if isinstance(getattr(e, "notiz", None), str) and e.notiz else ART_ICON.get(e.art, "")
                        chips += f"<span class='k {k}' title='{html.escape(art_label(e))} · {hrs}'>{ic}{hrs}</span>"
                rows += f"<td class='{'ov' if tot > H + 1e-9 else ''}{tdc}'>{chips}</td>"
        fw = fact[fact["consultant_id"] == c.id] if not fact.empty else fact
        u = logic.kpi(fw)
        rows += f"<td class='au' style='{util_css(u)}'>{fmt_pct(u) if u is not None else 'abw.'}</td>"
        if extra:
            rows += f"<td class='ex'>{extra[1].get(int(c.id), '')}</td>"
        rows += "</tr>"
    legend = (f"<div class='legend2'><span style='background:{WP['lilac']};border-left:3px solid {WP['navy']}'>Vor Ort"
              f"</span><span style='border:1px dashed {WP['mid']}'>Remote</span>"
              f"<span style='background:{WP['xlgrey']}'>Intern, Termin</span>"
              "<span style='background:repeating-linear-gradient(45deg,#F4F4F4,#F4F4F4 6px,#EAEAEA 6px,#EAEAEA 12px)'>"
              f"Abwesend</span><span style='box-shadow:inset 0 0 0 2px {WP['alert']}'>mehr als 8 h</span></div>")
    return (f"<style>{css}</style><div class='kfw'><table class='kf {groesse}'>{h1}{h2}{rows}</table></div>"
            f"{legend}")


def auslastung_gesamt() -> None:
    """Gesamtübersicht Auslastung: Kurzfrist (3 Wochen), Mittelfrist (Matrix), Verlauf, Umsatz."""
    today = heute()
    ss = st.session_state
    cons = db.consultants()
    if cons.empty:
        st.info("Noch keine Mitarbeiter angelegt.")
        return

    # ---------------- 1 | Kurzfrist
    sub("Kurzfrist: die nächsten 3 Wochen")
    ss.setdefault("kf_off", 0)
    n = st.columns([1, 1, 1, 6])
    n[0].button("◀ Woche", on_click=_kf_shift, args=(-1,), width="stretch", key="kf_prev")
    n[1].button("Heute", on_click=_kf_shift, args=(0,), width="stretch", key="kf_today")
    n[2].button("Woche ▶", on_click=_kf_shift, args=(1,), width="stretch", key="kf_next")
    start = logic.monday(today) + dt.timedelta(weeks=ss.kf_off)
    st.markdown(kurzfrist_html(start, 3, cons), unsafe_allow_html=True)

    # ---------------- 2 | Mittelfrist
    sub("Mittelfrist: Auslastung je Mitarbeiter")
    c1, c2 = st.columns([1.6, 4], vertical_alignment="bottom")
    monate = c1.slider("Monate voraus", 1, 12, 3, key="mf_mon")
    c2.caption("Bis zu drei Monate je Kalenderwoche, danach je Monat. Zeile anklicken: Verlauf darunter.")
    m_start = logic.monday(today)
    m_end = logic.month_end(today, monate - 1)
    fact = logic.build_fact(m_start, m_end, cons, db.entries(m_start, m_end))
    if fact.empty:
        st.info("Keine Daten im Zeitraum.")
        return
    team = logic.aggregate(fact, "Mittelfrist", today, by_consultant=False)
    pers = logic.aggregate(fact, "Mittelfrist", today, by_consultant=True)
    periods = team["p_label"].tolist()
    plab = person_label(cons)
    avg = {nm: logic.kpi(fact[fact["name"] == nm]) for nm in cons["name"]}
    order = list(cons["name"])
    sel = ss.get("aw_sel") if ss.get("aw_sel") in order else None
    pv = pers.pivot(index="name", columns="p_label", values="auslastung").reindex(index=order, columns=periods)
    pv.insert(0, "Ø", [avg[x] for x in order])
    ber = dict(zip(cons["name"], cons["bereich"]))
    gruppiert = len(bereiche_von(cons)) > 1
    if gruppiert:
        pv.index = pd.MultiIndex.from_tuples([(BER_KURZ.get(ber[x], ber[x]), plab.get(x, x)) for x in order],
                                             names=["Bereich", "Mitarbeiter"])
    else:
        pv.index = [plab.get(x, x) for x in order]
        pv.index.name = "Mitarbeiter"
    sel_lab = plab.get(sel, sel) if sel else None
    sty = styled_pct(pv).apply(
        lambda r: ["font-weight:800; text-decoration:underline"
                   if (r.name[-1] if isinstance(r.name, tuple) else r.name) == sel_lab else ""] * len(r), axis=1)
    key = f"aw_df_{ss.get('aw_ver', 0)}_{monate}"
    st.dataframe(sty, width="stretch", height=38 + 35 * len(pv), key=key, on_select=lambda: _aw_pick(key, order),
                 selection_mode="single-row")

    # ---------------- 3 | Verlauf
    if sel:
        d = pers[pers["name"] == sel]
        linien = [(plab.get(sel, sel), d["p_label"], d["auslastung"], WP["navy"], ziel_von(cons).get(sel, ZIEL))]
        name = plab.get(sel, sel)
    elif gruppiert:
        linien = []
        for b in bereiche_von(cons):
            cb = cons[cons["bereich"] == b]
            tb = logic.aggregate(fact[fact["consultant_id"].isin(cb["id"])], "Mittelfrist", today, by_consultant=False)
            linien.append((b, tb["p_label"], tb["auslastung"], ber_farbe(b), team_ziel(cb)))
        name = " und ".join(bereiche_von(cons))
    else:
        linien = [("Team gesamt", team["p_label"], team["auslastung"], WP["navy"], team_ziel(cons))]
        name = "Team gesamt"
    hc = st.columns([5, 1.2], vertical_alignment="center")
    hc[0].markdown(f"**Verlauf: {html.escape(name)}**" + ("" if sel else
                   " <span style='color:#999'>(Zeile in der Tabelle anklicken für einen Mitarbeiter)</span>"),
                   unsafe_allow_html=True)
    hc[1].button("Ganzes Team", on_click=_aw_all, disabled=sel is None, width="stretch", key="aw_all_btn")
    fig = go.Figure()
    for i, (nm, x, y, farbe, _) in enumerate(linien):
        fig.add_trace(go.Scatter(x=x, y=y, mode="lines+markers+text", name=nm, text=[fmt_pct(v) for v in y],
                                 textposition="top center" if i == 0 else "bottom center",
                                 line=dict(color=farbe, width=3), marker=dict(size=8),
                                 textfont=dict(color=farbe, size=11),
                                 hovertemplate=f"{html.escape(nm)} %{{x}}: %{{y:.0%}}<extra></extra>"))
    fig.add_hline(y=1.0, line_color=WP["alert"], line_width=1, line_dash="dot")
    zt = linien[0][4]
    fig.add_hline(y=zt, line_dash="dash", line_color=WP["grey"], annotation_text=f"Ziel {zt:.0%}",
                  annotation_font_color=WP["grey"])
    fig.update_xaxes(categoryorder="array", categoryarray=periods)
    fig.update_yaxes(rangemode="tozero")
    fig.update_layout(showlegend=len(linien) > 1, legend=dict(orientation="h", x=0, xanchor="left", y=-0.12, yanchor="top"))
    fig = plotly_layout(fig, 330, pct=True)
    fig.update_layout(legend=dict(orientation="h", x=0, xanchor="left", y=-0.14, yanchor="top"), margin=dict(b=40))
    st.plotly_chart(fig, width="stretch", config={"displayModeBar": False})

    # ---------------- 4 | Umsatz
    sub("Umsatz Team je Monat")
    umsatz_auswertung(nur_team=True)


def page_gesamt() -> None:
    header("Gesamtübersicht", "Wer ist wo, wie ausgelastet, wie stehen die Projekte")
    view = st.segmented_control("Ansicht", ["📈 Auslastung", "📁 Projekte", "🎯 Planungsgüte", "🤝 Akquise"],
                                default="📈 Auslastung", key="go_view2", label_visibility="collapsed")
    if view == "📁 Projekte":
        projekte_durchgehen()
    elif view == "🎯 Planungsgüte":
        planungsguete_view()
    elif view == "🤝 Akquise":
        akquise_uebersicht()
    else:
        auslastung_gesamt()
    st.write("")
    export_block()


# ====================================================================== Akquise (in der Gesamtübersicht)
def teur(x) -> str:
    """Tausend Euro mit deutschem Tausenderpunkt, z. B. 1.250."""
    try:
        return f"{float(x):,.0f}".replace(",", ".")
    except (TypeError, ValueError):
        return ""


AQ_COLS = {"kunde": "Kunde", "thema": "Thema", "stufe": "Stufe", "val_tage": "Volumen (TEUR)",
           "mt_tage": "Manntage", "wahrscheinlichkeit": "Wahrsch. (%)", "von": "Start", "bis": "Ende"}


def _aq_sig(ak: pd.DataFrame) -> tuple:
    return tuple(zip(ak["id"].astype(int), ak["updated_at"].astype(str))) if not ak.empty else ()


def _aq_num(v):
    try:
        return None if v is None or pd.isna(v) else float(v)
    except (TypeError, ValueError):
        return None


def _aq_date(v):
    if v is None or (not isinstance(v, (dt.date, dt.datetime, str)) and pd.isna(v)):
        return None
    try:
        d = pd.to_datetime(v, errors="coerce")
    except Exception:
        return None
    return None if pd.isna(d) else d.date()


def _aq_zeile(r: dict, kz: list[str]) -> dict:
    """Editorzeile in ein Akquise-Thema übersetzen (einheitliche Typen für Vergleich und Speichern)."""
    stufe = r.get("Stufe")
    ber = r.get("Bereich")
    return {"kunde": _s(r.get("Kunde")).strip(), "thema": _s(r.get("Thema")).strip(),
            "bereich": ber if ber in db.BEREICHE else db.BEREICHE[0],
            "stufe": stufe if stufe in db.AKQUISE_STUFEN else "Lead",
            "val_tage": _aq_num(r.get(AQ_COLS["val_tage"])), "mt_tage": _aq_num(r.get(AQ_COLS["mt_tage"])),
            "wahrscheinlichkeit": _aq_num(r.get(AQ_COLS["wahrscheinlichkeit"])),
            "von": _aq_date(r.get("Start")), "bis": _aq_date(r.get("Ende")),
            "team": [k for k in kz if bool(r.get(k)) and r.get(k) == r.get(k)]}


def _aq_basis(ak: pd.DataFrame, kz: list[str]) -> pd.DataFrame:
    """Startstand des Editors: Angebot vor Opportunity vor Lead, darin nach Start."""
    rang = {s_: i for i, s_ in enumerate(reversed(db.AKQUISE_STUFEN))}
    if not ak.empty:
        brang = {b: i for i, b in enumerate(db.BEREICHE)}
        ak = ak.assign(_b=ak["bereich"].map(brang).fillna(9), _r=ak["stufe"].map(rang).fillna(9)).sort_values(
            ["_b", "_r", "von", "id"], na_position="last")
    rows = []
    for r in ak.itertuples():
        team = set(_s(r.team).split(",")) if _s(r.team) else set()
        rows.append({"id": int(r.id), "Bereich": r.bereich if r.bereich in db.BEREICHE else db.BEREICHE[0],
                     "Kunde": _s(r.kunde), "Thema": _s(r.thema),
                     "Stufe": r.stufe if r.stufe in db.AKQUISE_STUFEN else "Lead",
                     AQ_COLS["val_tage"]: float(r.val_tage) if pd.notna(r.val_tage) else 0.0,
                     AQ_COLS["mt_tage"]: float(r.mt_tage) if pd.notna(r.mt_tage) else 0.0,
                     AQ_COLS["wahrscheinlichkeit"]: float(r.wahrscheinlichkeit) if pd.notna(r.wahrscheinlichkeit)
                     else 0.0,
                     "Start": r.von if pd.notna(r.von) else None, "Ende": r.bis if pd.notna(r.bis) else None,
                     **{k: k in team for k in kz}})
    cols = ["id", "Bereich", "Kunde", "Thema", "Stufe", AQ_COLS["val_tage"], AQ_COLS["mt_tage"], AQ_COLS["wahrscheinlichkeit"],
            "Start", "Ende"] + kz
    df = pd.DataFrame(rows, columns=cols)
    df["id"] = df["id"].astype("Int64")
    for c in ("Bereich", "Kunde", "Thema", "Stufe", "Start", "Ende"):
        df[c] = df[c].astype(object)
    for c in (AQ_COLS["val_tage"], AQ_COLS["mt_tage"], AQ_COLS["wahrscheinlichkeit"]):
        df[c] = df[c].astype(float)
    for k in kz:
        df[k] = df[k].astype(bool)
    return df


def akquise_uebersicht() -> None:
    """Pipeline direkt in der Tabelle pflegen. Jede Änderung wird sofort gespeichert, Kennzahlen rechnen mit."""
    ss = st.session_state
    cons = db.consultants()
    kz = cons["name"].tolist()
    voll = {r.name: f"{_s(r.vollname) or r.name} · {r.bereich}" for r in cons.itertuples()}
    ak = db.akquise()
    sig = _aq_sig(ak)
    # Neu laden, wenn jemand anderes etwas geändert hat oder sich das Team geändert hat
    if ss.get("aq_base") is None or ss.get("aq_sig") != sig or ss.get("aq_kz") != kz:
        base = _aq_basis(ak, kz)
        ss.aq_base = base
        ss.aq_sig = sig
        ss.aq_kz = kz
        ss.aq_saved = {int(r["id"]): _aq_zeile(r, kz) for r in base.to_dict("records")}
        ss.aq_ver = ss.get("aq_ver", 0) + 1
    base = ss.aq_base
    kpi_box = st.container()
    sub("Pipeline und mögliches Staffing")
    cfg = {"Bereich": st.column_config.SelectboxColumn("Bereich", options=db.BEREICHE, default=db.BEREICHE[0],
                                                       required=True, width=90,
                                                       format_func=lambda b: BER_KURZ.get(b, b)),
           "Kunde": st.column_config.TextColumn("Kunde", required=True, width=150),
           "Thema": st.column_config.TextColumn("Thema", width=170),
           "Stufe": st.column_config.SelectboxColumn("Stufe", options=db.AKQUISE_STUFEN, default="Lead",
                                                     required=True, width=100),
           AQ_COLS["val_tage"]: st.column_config.NumberColumn(AQ_COLS["val_tage"], min_value=0, step=10, format="%d",
                                                              default=0, help="In Tausend Euro: 100 = 100.000 €"),
           AQ_COLS["mt_tage"]: st.column_config.NumberColumn(AQ_COLS["mt_tage"], min_value=0, step=0.5, format="%g",
                                                             default=0, width="small"),
           AQ_COLS["wahrscheinlichkeit"]: st.column_config.NumberColumn(AQ_COLS["wahrscheinlichkeit"], min_value=0,
                                                                        max_value=100, step=5, format="%d %%",
                                                                        default=0, width="small"),
           "Start": st.column_config.DateColumn("Start", format="DD.MM.YYYY", width="small"),
           "Ende": st.column_config.DateColumn("Ende", format="DD.MM.YYYY", width="small"),
           **{k: st.column_config.CheckboxColumn(k, help=voll[k], default=False, width=62) for k in kz}}
    ed = st.data_editor(base, key=f"aq_ed_{ss.aq_ver}", num_rows="dynamic", hide_index=True, width="stretch",
                        column_order=[c for c in base.columns if c != "id"], column_config=cfg,
                        height=min(38 + 35 * (len(base) + 2), 700))
    st.caption("Direkt in der Tabelle bearbeiten, jede Änderung wird sofort gespeichert. Neues Thema: unterste Zeile "
               "(＋). Entfernen: Zeile links markieren und auf den Papierkorb klicken. Häkchen = fachlich passend, "
               "unabhängig von der Auslastung. Spaltenkopf anklicken sortiert.")

    # ---------------- Speichern: Änderungen sofort, neue und gelöschte Zeilen mit Neuaufbau
    user = ss.get("user", "")
    neu_aufbau = geaendert = False
    rows = ed.to_dict("records")
    for r in rows:
        z = _aq_zeile(r, kz)
        rid = r.get("id")
        if rid is None or pd.isna(rid):
            if z["kunde"]:
                db.upsert_akquise(z, user)
                neu_aufbau = True
            continue
        rid = int(rid)
        if z["kunde"] and ss.aq_saved.get(rid) != z:
            db.upsert_akquise(z | {"id": rid}, user)
            ss.aq_saved[rid] = z
            geaendert = True
    weg = set(int(i) for i in base["id"].dropna()) - set(int(i) for i in ed["id"].dropna())
    for rid in weg:
        db.archive_akquise(rid, user)
        neu_aufbau = True
    if neu_aufbau:
        ss.aq_base = None
        st.rerun()
    if geaendert:
        ss.aq_base = ed.reset_index(drop=True)  # Startstand nachziehen, damit nichts zurückspringt
    ss.aq_sig = _aq_sig(db.akquise())

    # ---------------- Kennzahlen aus dem aktuellen Tabellenstand
    z = pd.DataFrame([_aq_zeile(r, kz) for r in rows])
    z = z[z["kunde"] != ""] if not z.empty else z
    with kpi_box:
        for b in db.BEREICHE:
            zb = z[z["bereich"] == b] if not z.empty else z
            st.markdown(f"<div style='margin:.4rem 0 .2rem;font-weight:700;color:{ber_farbe(b)}'>"
                        f"{html.escape(b)} <span style='color:{WP['grey']};font-weight:400'>· {len(zb)} "
                        f"{'Thema' if len(zb) == 1 else 'Themen'}</span></div>", unsafe_allow_html=True)
            k = st.columns(4)
            for i, s_ in enumerate(db.AKQUISE_STUFEN):
                t = zb[zb["stufe"] == s_] if not zb.empty else zb
                vol = t["val_tage"].fillna(0).sum() if not t.empty else 0
                k[i].metric(s_, f"{len(t)}", f"{teur(vol)} TEUR", delta_color="off", delta_arrow="off")
            ges = zb["val_tage"].fillna(0).sum() if not zb.empty else 0
            gew = (zb["val_tage"].fillna(0) * zb["wahrscheinlichkeit"].fillna(0) / 100).sum() if not zb.empty else 0
            k[3].metric("Volumen gewichtet (TEUR)", teur(gew), f"von {teur(ges)} TEUR gesamt", delta_color="off",
                        delta_arrow="off", help="Summe Volumen × Wahrscheinlichkeit im Bereich")
        if not z.empty:
            n_per = {k_: sum(k_ in t for t in z["team"]) for k_ in kz}
            if any(n_per.values()):
                st.caption("Nennungen je Kollege: " + " · ".join(f"{k_} {n_per[k_]}" for k_ in kz if n_per[k_]))


# ====================================================================== Reiter: Spielregeln & Montagsrunde
def page_board() -> None:
    """Spielregeln (Commitments für die Planung) und Themen für die Montagsrunde."""
    ss = st.session_state
    header("Team-Board", "Spielregeln für die Planung und Themen für die Montagsrunde")
    if ss.get("bd_msg"):
        st.toast(ss.pop("bd_msg"))
    c1, c2 = st.columns([1, 1], gap="large")
    with c1:
        sub("Spielregeln und Commitments")
        regeln = db.spielregeln()
        if regeln.empty:
            st.caption("Noch keine Spielregeln hinterlegt.")
        for r in regeln.itertuples():
            a, b = st.columns([12, 1], vertical_alignment="center")
            a.markdown(f"<div class='regel'>{html.escape(r.text)}</div>", unsafe_allow_html=True)
            if is_admin() and b.button("✕", key=f"rg_del_{r.id}", type="tertiary", help="Regel entfernen"):
                db.delete_spielregel(int(r.id))
                st.rerun()
        if is_admin():
            with st.form("rg_form", clear_on_submit=True, border=False):
                t = st.text_input("Neue Spielregel", placeholder="z. B. Planungsziel: 4 Tage vor Ort, 1 Tag remote")
                if st.form_submit_button("Hinzufügen") and t.strip():
                    db.add_spielregel(t, ss.user)
                    st.rerun()
        else:
            st.caption("Spielregeln pflegt der Admin. Vorschläge gern als Thema für die Montagsrunde.")
    with c2:
        offen = db.montagsthemen(True)
        sub(f"Themen für die Montagsrunde ({len(offen)} offen)")
        with st.form("mt_form", clear_on_submit=True, border=True):
            t = st.text_area("Neues Thema einbringen", height=80,
                             placeholder="z. B. Wann machen wir die nächste Teamveranstaltung?")
            if st.form_submit_button("Thema einreichen", type="primary") and t.strip():
                db.add_montagsthema(t, ss.user)
                ss["bd_msg"] = "✓ Thema für die Montagsrunde eingereicht."
                st.rerun()
        cons = db.consultants(active_only=False)
        voll = {r.name: _s(r.vollname) or r.name for r in cons.itertuples()}
        if offen.empty:
            st.caption("Keine offenen Themen.")
        for r in offen.itertuples():
            with st.container(border=True):
                st.markdown(f"<div class='mt-t'>{html.escape(r.text)}</div><div class='mt-m'>"
                            f"{html.escape(voll.get(r.von, r.von or ''))} · "
                            f"{pd.Timestamp(r.erstellt):%d.%m.%Y}</div>", unsafe_allow_html=True)
                a, b, _ = st.columns([1.5, 1, 2.5])
                if a.button("✓ erledigt", key=f"mt_ok_{r.id}", width="stretch", help="In der Montagsrunde besprochen"):
                    db.erledige_montagsthema(int(r.id))
                    st.rerun()
                if (is_admin() or r.von == ss.user) and b.button("Löschen", key=f"mt_del_{r.id}", type="tertiary"):
                    db.delete_montagsthema(int(r.id))
                    st.rerun()
        erl = db.montagsthemen(False)
        if not erl.empty:
            with st.expander(f"Besprochene Themen ({len(erl)})"):
                for r in erl.head(30).itertuples():
                    a, b = st.columns([6, 1.2], vertical_alignment="center")
                    a.markdown(f"~~{html.escape(r.text)}~~ <span style='color:{WP['grey']};font-size:.8rem'>· "
                               f"{pd.Timestamp(r.erledigt_am):%d.%m.%Y}</span>", unsafe_allow_html=True)
                    if b.button("wieder öffnen", key=f"mt_re_{r.id}", type="tertiary"):
                        db.erledige_montagsthema(int(r.id), False)
                        st.rerun()


# ====================================================================== Reiter: Verwaltung (Admin)
def _read_export(file) -> dict:
    """Liest einen App-Export (Ressourcenplanung_Stand_*.xlsx) als Wiederherstellung: Team und Tage aus Blatt
    'Tage' (Ort aus der Zellfarbe), Projekte und Budgets aus Blatt 'Projekte', gemeldete VAL aus 'VAL Details'."""
    import openpyxl
    wb = openpyxl.load_workbook(file)
    abs_ = set(logic.ABWESEND_ARTEN) | set(logic.INTERN_ARTEN)
    projekte, budgets, cur = [], [], None
    for r in wb["Projekte"].iter_rows(min_row=4, values_only=True):
        if len(r) < 9:
            continue
        kunde, projekt, berater, budget, start, ende = r[0], r[1], r[2], r[3], r[7], r[8]
        if berater == "Gesamt":
            cur = (str(kunde), str(projekt or kunde))
            projekte.append({"kunde": cur[0], "projekt": cur[1], "website": "",
                             "start": start.date() if start else None, "ende": ende.date() if ende else None,
                             "leitung": "", "budget": float(budget or 0)})
        elif berater and cur:
            budgets.append({"kunde": cur[0], "projekt": cur[1], "kuerzel": str(berater), "budget": float(budget or 0)})
    label = {(p["kunde"] if p["projekt"] == p["kunde"] else f"{p['kunde']} | {p['projekt']}"): (p["kunde"], p["projekt"])
             for p in projekte}
    wt = wb["Tage"]
    days = {c: wt.cell(6, c).value.date() for c in range(6, wt.max_column + 1)
            if isinstance(wt.cell(6, c).value, dt.datetime)}
    team, planung, kz = [], [], None
    for r in range(7, wt.max_row + 1):
        a, b = wt.cell(r, 1).value, wt.cell(r, 2).value
        if a and isinstance(b, str) and b.startswith("Summe ("):
            kz = b[len("Summe ("):-1]
            team.append({"kuerzel": kz, "name": str(a), "rolle": ""})
            continue
        if not b or not kz or a or (b not in abs_ and b not in label):
            continue
        for c, d in days.items():
            v = wt.cell(r, c).value
            if not isinstance(v, (int, float)) or v <= 0:
                continue
            if b in label:
                fill = str(wt.cell(r, c).fill.fgColor.rgb or "")
                planung.append({"kuerzel": kz, "datum": d, "art": "Remote" if fill.endswith("FFFFFF") else "Vor Ort",
                                "kunde": label[b][0], "projekt": label[b][1], "stunden": round(v * H, 2)})
            else:
                planung.append({"kuerzel": kz, "datum": d, "art": b, "kunde": None, "projekt": None,
                                "stunden": round(v * H, 2)})
    if "Stammdaten" in wb.sheetnames:
        rollen = {str(r[0]): (r[2] or "") for r in wb["Stammdaten"].iter_rows(min_row=2, values_only=True) if r[0]}
        for t in team:
            t["rolle"] = rollen.get(t["kuerzel"], "")
    val = []
    if "VAL Details" in wb.sheetnames:
        names = {t["name"]: t["kuerzel"] for t in team}
        mon = {m: i + 1 for i, m in enumerate(MONATE_LANG)}
        for r in wb["VAL Details"].iter_rows(min_row=4, values_only=True):
            if r and r[0] and len(r) >= 8 and r[7] == "gemeldet" and r[2] in label:
                m, y = str(r[0]).split()
                val.append({"kuerzel": names.get(r[1], r[1]), "monat": dt.date(int(y), mon[m], 1),
                            "kunde": label[r[2]][0], "projekt": label[r[2]][1], "tage": float(r[4] or 0)})
    return {"team": team, "projekte": projekte, "budgets": budgets, "planung": planung, "val": val,
            "positionen": []}


def _read_import(file) -> dict:
    sheets = pd.read_excel(file, sheet_name=None)
    if "Tage" in sheets and "Team" not in sheets:  # App-Export als Backup
        file.seek(0)
        return _read_export(file)
    need = {"Team", "Projekte", "Budgets", "Planung"}
    missing = need - set(sheets)
    if missing:
        raise ValueError("Es fehlen die Blätter: " + ", ".join(sorted(missing)))

    def d(v):
        return None if pd.isna(v) else pd.Timestamp(v).date()

    t, p, b, pl = sheets["Team"], sheets["Projekte"], sheets["Budgets"], sheets["Planung"]
    val_rows, pos_rows = [], []
    if "VAL" in sheets:
        val_rows = [{"kuerzel": str(r["Kürzel"]), "monat": d(r["Monat"]), "kunde": str(r["Kunde"]),
                     "projekt": str(r["Projekt"]), "tage": float(r["VAL BT"] or 0)}
                    for _, r in sheets["VAL"].dropna(subset=["Kürzel", "Monat"]).iterrows()]
    if "Positionen" in sheets:
        pos_rows = [{"name": str(r["Position"]),
                     "tagessatz": float(r["Tagessatz"]) if pd.notna(r.get("Tagessatz")) else None,
                     "ziel": float(r["Zielauslastung"]) if pd.notna(r.get("Zielauslastung")) else None}
                    for _, r in sheets["Positionen"].dropna(subset=["Position"]).iterrows()]
    return {
        "val": val_rows, "positionen": pos_rows,
        "team": [{"kuerzel": str(r["Kürzel"]).strip(), "name": _s(r.get("Name")), "rolle": _s(r.get("Rolle"))}
                 for _, r in t.dropna(subset=["Kürzel"]).iterrows()],
        "projekte": [{"kunde": str(r["Kunde"]), "projekt": str(r["Projekt"]), "website": _s(r.get("Website")),
                      "start": d(r.get("Start")), "ende": d(r.get("Ende")),
                      "leitung": _s(r.get("Leitung (Kürzel)")), "budget": float(r.get("Budget gesamt BT") or 0)}
                     for _, r in p.dropna(subset=["Kunde", "Projekt"]).iterrows()],
        "budgets": [{"kunde": str(r["Kunde"]), "projekt": str(r["Projekt"]), "kuerzel": str(r["Kürzel"]),
                     "budget": float(r["Budget BT"] or 0)} for _, r in b.dropna(subset=["Kürzel"]).iterrows()],
        "planung": [{"kuerzel": str(r["Kürzel"]), "datum": d(r["Datum"]), "art": str(r["Art"]),
                     "kunde": _s(r.get("Kunde")) or None, "projekt": _s(r.get("Projekt")) or None,
                     "stunden": float(r["Stunden"])} for _, r in pl.dropna(subset=["Kürzel", "Datum"]).iterrows()],
    }


def import_view() -> None:
    st.markdown("Nur für die **einmalige Erstbefüllung**. Danach ist die App die Datenbasis: alle Eingaben werden "
                "sofort gespeichert, ein erneuter Upload ist nicht nötig.")
    st.caption("Format: Startdaten (Blätter Team, Projekte, Budgets, Planung) **oder** ein Excel-Export der App "
               "(„Aktueller Stand als Excel“) als Wiederherstellung.")
    up = st.file_uploader("Startdaten (xlsx)", type=["xlsx"], key="imp_file")
    if not up:
        return
    try:
        data = _read_import(up)
    except Exception as e:  # noqa: BLE001
        st.error(f"Datei konnte nicht gelesen werden: {e}")
        return
    k = st.columns(4)
    k[0].metric("Mitarbeiter", len(data["team"]))
    k[1].metric("Projekte", len(data["projekte"]))
    k[2].metric("Budgets", len(data["budgets"]))
    k[3].metric("Tageseinträge", len(data["planung"]))
    replace = st.checkbox("Alle bisherigen Planungsdaten löschen und durch diese Datei ersetzen",
                          value=False, key="imp_replace")
    if replace:
        st.error("Achtung: Alles, was das Team seit dem letzten Import in der App eingetragen hat, wird gelöscht.")
    ok = st.text_input("Zur Bestätigung IMPORT eintippen", key="imp_ok")
    if st.button("Importieren", type="primary", disabled=ok != "IMPORT", key="imp_go"):
        with st.spinner("Daten werden übernommen ..."):
            res = db.import_startdaten(data["team"], data["projekte"], data["budgets"], data["planung"],
                                       st.session_state.user, replace=replace, val_rows=data["val"],
                                       positionen_neu=data["positionen"])
        st.success(f"✓ Übernommen: {res['team']} Mitarbeiter, {res['projekte']} Projekte, {res['budgets']} Budgets, "
                   f"{res['eintraege']} Tageseinträge, {res['val']} VAL-Werte"
                   + (f" ({res['uebersprungen']} übersprungen)" if res["uebersprungen"] else "") + ".")
        st.session_state.pop("imp_ok", None)


def feiertage_view() -> None:
    """Admin: firmenweite freie Tage (z. B. Heiligabend, Brückentage). Gelten sofort für alle wie Feiertage."""
    today = heute()
    ss = st.session_state
    st.markdown("Tage, die für **das ganze Team** frei sind, z. B. Heiligabend, Silvester oder Brückentage. Sie wirken "
                "wie gesetzliche Feiertage: nicht planbar, zählen nicht zur Kapazität, erscheinen im Kalender und "
                "im Excel als Feiertag.")
    if ss.get("ft_msg"):
        st.success(ss.pop("ft_msg"))
    with st.container(border=True):
        c1, c2, c3 = st.columns([1, 1, 2])
        von = c1.date_input("Von", today, format="DD.MM.YYYY", key="ft_von")
        bis = c2.date_input("Bis", von, format="DD.MM.YYYY", key="ft_bis",
                            help="Für einen einzelnen Tag Von und Bis gleich lassen.")
        name = c3.text_input("Bezeichnung", key="ft_name", placeholder="z. B. Heiligabend, Brückentag, Betriebsruhe")
        tage = [d.date() for d in pd.date_range(von, bis)] if von <= bis else []
        tage = [d for d in tage if d.weekday() < 5 and not logic.feiertage(d.year).get(d)]
        betroffen = db.entries(min(tage), max(tage)) if tage else pd.DataFrame()
        if not betroffen.empty:
            betroffen = betroffen[betroffen["tag"].isin(tage)]
        n_e = len(betroffen)
        st.caption(f"{len(tage)} Arbeitstag{'e' if len(tage) != 1 else ''} (Wochenenden und gesetzliche Feiertage "
                   "werden übersprungen).")
        ok = True
        if n_e:
            wer = ", ".join(sorted(betroffen["name"].unique()))
            st.warning(f"An diesen Tagen sind bereits **{n_e} Einträge** geplant ({wer}). Sie werden gelöscht.")
            ok = st.checkbox("Ja, Einträge löschen und freie Tage für alle setzen", key="ft_ok")
        if st.button("Für alle eintragen", type="primary", disabled=not tage or not name.strip() or not ok,
                     key="ft_go"):
            n = db.add_firmenfeiertage(tage, name.strip(), ss.user)
            for k in ("ft_name", "ft_ok"):
                ss.pop(k, None)
            ss["ft_msg"] = (f"✓ {len(tage)} freie{'r' if len(tage) == 1 else ''} Tag{'e' if len(tage) != 1 else ''} "
                            f"für alle eingetragen" + (f", {n} Einträge entfernt." if n else "."))
            st.rerun()

    ft = db.firmenfeiertage()
    kommend = {d: n for d, n in ft.items() if d >= today - dt.timedelta(days=7)}
    sub("Eingetragene freie Tage")
    if not kommend:
        st.caption("Noch keine firmenweiten freien Tage eingetragen.")
    for d, n in kommend.items():
        c1, c2 = st.columns([5, 1], vertical_alignment="center")
        c1.markdown(f"**{logic.WOTAG[d.weekday()]} {d:%d.%m.%Y}** · {html.escape(n)}")
        if c2.button("Entfernen", key=f"ft_del_{d}", width="stretch", type="tertiary"):
            db.delete_firmenfeiertag(d)
            st.rerun()
    with st.expander(f"Gesetzliche Feiertage Bayern {today.year} und {today.year + 1} (automatisch)"):
        for y in (today.year, today.year + 1):
            st.markdown("  \n".join(f"{logic.WOTAG[d.weekday()]} {d:%d.%m.%Y} · {n}"
                                     for d, n in sorted(logic.feiertage(y).items())))


def team_termine_view() -> None:
    """Admin: Termine für mehrere oder alle Mitarbeiter einplanen (z. B. Team-Event, Strategietag).
    Wird als 'Intern' mit Namen eingetragen; jeder kann es im eigenen Kalender wieder entfernen."""
    today = heute()
    ss = st.session_state
    st.markdown("Termine für das Team **direkt einplanen**, z. B. Team-Event, Strategietag oder Schulung. Sie stehen "
                "danach bei allen Teilnehmern im Kalender (📌 mit Namen, zählt als Intern). Wer nicht kann, nimmt den "
                "Termin im eigenen Kalender per Klick auf den Tag wieder heraus.")
    if ss.get("tt_msg"):
        st.success(ss.pop("tt_msg"))
    cons = db.consultants()
    lab = {r.name: _vorname_label(r) for r in cons.itertuples()}
    idmap = dict(zip(cons["name"], cons["id"]))
    with st.container(border=True):
        c1, c2, c3 = st.columns([1, 1, 2])
        von = c1.date_input("Von", today + dt.timedelta(days=7), format="DD.MM.YYYY", key="tt_von")
        bis = c2.date_input("Bis", von, format="DD.MM.YYYY", key="tt_bis",
                            help="Für einen einzelnen Tag Von und Bis gleich lassen.")
        name = c3.text_input("Termin", key="tt_name", placeholder="z. B. Team-Event, Strategietag, Schulung")
        c4, c5 = st.columns([1, 3])
        umf = c4.segmented_control("Umfang je Tag", ["½ Tag", "1 Tag"], default="1 Tag", key="tt_umf")
        wer = c5.pills("Teilnehmer", cons["name"].tolist(), selection_mode="multi", default=cons["name"].tolist(),
                       key="tt_wer", format_func=lambda n: lab.get(n, n))
        tage = [d.date() for d in pd.date_range(von, bis)] if von <= bis else []
        tage = [d for d in tage if d.weekday() < 5 and not logic.feiertag_name(d)]
        raus = st.toggle("Projekttage der Teilnehmer an diesen Tagen entfernen", key="tt_raus",
                         help="Aus: Termin kommt zusätzlich dazu (Tag ggf. überbucht). An: geplante Projekttage an "
                              "diesen Tagen werden für die Teilnehmer gelöscht.")
        st.caption(f"{len(tage)} Arbeitstag{'e' if len(tage) != 1 else ''} × {len(wer or [])} Teilnehmer. "
                   "Wochenenden und Feiertage werden übersprungen.")
        if st.button("Termin für alle einplanen", type="primary", key="tt_go",
                     disabled=not tage or not (wer or []) or not name.strip()):
            h = H / 2 if umf == "½ Tag" else H
            n = db.add_team_termin([idmap[w] for w in wer], tage, h, name, ss.user, projekte_raus=raus)
            for k in ("tt_name", "tt_raus"):
                ss.pop(k, None)
            ss["tt_msg"] = f"✓ „{name.strip()}“ eingeplant ({n} Einträge)."
            st.rerun()
    sub("Geplante Team-Termine")
    tt = db.team_termine(today - dt.timedelta(days=7))
    if tt.empty:
        st.caption("Noch keine Team-Termine geplant.")
    for r in tt.itertuples():
        c1, c2 = st.columns([5, 1], vertical_alignment="center")
        c1.markdown(f"**{logic.WOTAG[r.tag.weekday()]} {r.tag:%d.%m.%Y}** · 📌 {html.escape(r.notiz)} · "
                    f"{int(r.n)} Teilnehmer · {hh(r.stunden)}")
        if c2.button("Für alle löschen", key=f"tt_del_{r.tag}_{r.notiz}", type="tertiary", width="stretch"):
            db.delete_team_termin(r.tag, r.notiz)
            st.rerun()


def page_verwaltung() -> None:
    header("Verwaltung", "Freie Tage, Team-Termine, Kunden, Logos und Datenimport")
    t1, t0, t2, t3 = st.tabs(["Freie Tage (alle)", "Team-Termine", "Kunden & Logos", "Import"])
    with t1:
        feiertage_view()
    with t0:
        team_termine_view()
    with t2:
        kunden_logos()
    with t3:
        import_view()
        st.write("")
        with st.expander("Planungsdaten komplett zurücksetzen"):
            st.warning("Löscht **alle Einträge, Projekte und Kunden** unwiderruflich. Das Team bleibt erhalten.")
            ok = st.text_input("Zur Bestätigung LÖSCHEN eintippen", key="reset_ok")
            if st.button("Alles zurücksetzen", disabled=ok != "LÖSCHEN", key="reset_go"):
                db.reset_planning()
                st.session_state.pop("reset_ok", None)
                st.success("Planungsdaten gelöscht.")


# ====================================================================== Main
def main() -> None:
    try:
        init_once()
    except Exception as e:
        st.error("Datenbank nicht erreichbar. Bitte DATABASE_URL in den Secrets prüfen "
                 "(Supabase: Session Pooler Connection String verwenden; pausiertes Projekt wieder starten).")
        st.exception(e)
        st.stop()
    try:
        logic.set_firmen_feiertage(db.firmenfeiertage())
        logic.set_val_ist(db.val(), db.gemeldet_set())
    except Exception:
        pass
    inject_css()
    extra_css()
    if not login_gate():
        st.stop()
    branding()
    if not db.is_pg() and (users() or TEAM_PW):
        st.error("⚠️ **Keine dauerhafte Datenbank verbunden.** Die App läuft mit einem lokalen Zwischenspeicher, "
                 "der bei jedem Neustart gelöscht wird. Alle Eingaben gehen dann verloren. Ursache: In den "
                 "Streamlit-Secrets fehlt `DATABASE_URL` ganz oben (vor dem ersten [Abschnitt]).")
    pages = [st.Page(page_mein_bereich, title="Mein Bereich", icon=":material/person:", default=True),
             st.Page(page_projekt, title="Projekt anlegen", icon=":material/add_circle:"),
             st.Page(page_val, title="VAL buchen", icon=":material/receipt_long:"),
             st.Page(page_gesamt, title="Gesamtübersicht", icon=":material/dashboard:")]
    try:
        n_mt = len(db.montagsthemen(True))
    except Exception:
        n_mt = 0
    pages.insert(0, st.Page(page_board, title=f"Team-Board ({n_mt})" if n_mt else "Team-Board",
                            icon=":material/campaign:"))
    if is_admin():
        pages.append(st.Page(page_team, title="Team", icon=":material/group:"))
        pages.append(st.Page(page_verwaltung, title="Verwaltung", icon=":material/settings:"))
    pg = st.navigation(pages, position="top")
    with st.sidebar:
        st.caption(f"Angemeldet als **{st.session_state.get('user_label', st.session_state.user)}**"
                   f"{' (Admin)' if is_admin() and ADMINS else ''}"
                   f"  \nDatenbank: {'✅ Supabase (dauerhaft)' if db.is_pg() else '⚠️ lokal (nicht dauerhaft)'}")
        if st.button("Abmelden"):
            st.session_state.clear()
            st.rerun()
    pg.run()


main()
